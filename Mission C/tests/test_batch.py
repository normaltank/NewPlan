import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from openpyxl import Workbook
ROOT = Path(__file__).resolve().parents[1]

class BatchContract(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.data = self.folder / '中文 data'
        self.input = self.folder / '原表.xlsx'
        book = Workbook()
        sheet = book.active
        sheet.title = 'Sheet1'
        sheet.append(['ID', '学校', '学院', '师资URL'])
        sheet.append([1, '虚构大学', '光学院', 'https://example.edu/list'])
        sheet.append([1, '虚构大学', '光学院', 'https://example.edu/list2'])
        sheet.append([2, '虚构大学', '理学院', 'https://example.edu/science'])
        sheet = book.create_sheet('Sheet2')
        sheet.append(['学院ID', '个人URL'])
        sheet.append([1, 'https://example.edu/teacher#bio'])
        sheet.append([2, 'https://example.edu/teacher'])
        sheet.append([99, 'https://example.edu/orphan'])
        sheet.append([1, 'bad'])
        book.save(self.input)
    def cli(self, *args, ok=True):
        p = subprocess.run([sys.executable, str(ROOT/'batch.py'), '--data-root', str(self.data), *map(str,args)], capture_output=True, encoding='utf-8', errors='replace')
        if ok: self.assertEqual(p.returncode, 0, p.stderr)
        return json.loads(p.stdout) if ok else p
    def test_preview_import_visible_batch_and_idempotent_content(self):
        preview = self.cli('preview', self.input)
        self.assertFalse(self.data.exists())
        self.assertEqual(len(preview['tasks']), 1)
        self.assertEqual(len(preview['tasks'][0]['relations']), 2)
        self.assertEqual(len(preview['issues']), 2)
        imported = self.cli('import', self.input)
        batch = self.cli('status', imported['id'])
        self.assertEqual(batch['counts']['pending'], 1)
        self.assertEqual(len(batch['tasks'][0]['sources']), 2)
        self.assertEqual(batch['tasks'][0]['relations'][0]['faculty_urls'], ['https://example.edu/list', 'https://example.edu/list2'])
        self.assertEqual(self.cli('import', self.input)['id'], imported['id'])
        self.assertEqual(len(self.cli('list')), 1)
    def cache(self):
        source = self.folder/'page.html'
        source.write_text('<meta charset="utf-8"><main><h1>李明</h1><p>职称：教授</p><p>本人邮箱：liming@example.edu</p><h2>个人简介</h2><p>研究光学。</p></main>', encoding='utf-8')
        import hashlib
        manifest = self.folder/'cache.json'
        manifest.write_text(json.dumps({'version':1,'entries':[{'url':'https://example.edu/teacher', 'source':{'path':'page.html','sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'observed_at':'2026-10-28T01:00:00Z'}}]}), encoding='utf-8')
        return manifest
    def test_run_pause_resume_keeps_bound_artifacts_and_detects_missing(self):
        batch = self.cli('import', self.input)
        ident = batch['id']
        self.cli('pause', ident)
        self.assertEqual(self.cli('run', ident, '--cache-manifest', self.cache())['counts'], {'pending':1})
        self.cli('resume', ident)
        result = self.cli('run', ident, '--cache-manifest', self.cache())
        task = result['tasks'][0]
        self.assertNotIn(task['state'], ['pending','running','failed'])
        artifact = self.data/task['artifact']
        original = {str(p.relative_to(artifact)):p.read_bytes() for p in artifact.rglob('*') if p.is_file()}
        self.cli('run', ident, '--cache-manifest', self.cache())
        self.assertEqual(original, {str(p.relative_to(artifact)):p.read_bytes() for p in artifact.rglob('*') if p.is_file()})
        (artifact/'0001/biography.txt').unlink()
        missing = self.cli('status', ident)
        self.assertEqual(missing['tasks'][0]['state'], 'artifact_missing')
        self.assertEqual(self.cli('run', ident, '--cache-manifest', self.cache())['tasks'][0]['state'], 'artifact_missing')
    def test_queued_cache_survives_input_removal_and_data_root_move(self):
        ident = self.cli('import', self.input)['id']
        self.cli('enqueue', ident, '--cache-manifest', self.cache())
        (self.folder/'cache.json').unlink()
        (self.folder/'page.html').unlink()
        import shutil
        moved = self.folder/'moved 数据'
        shutil.move(self.data, moved)
        self.data = moved
        self.assertTrue(self.cli('status', ident)['queued'])
        done = self.cli('run', ident)
        self.assertFalse(done['queued'])
        self.assertIn(done['tasks'][0]['state'], ['success','review'])
    def test_public_process_pause_concurrency_and_kill_recovery(self):
        from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
        import threading, time
        entered, release = threading.Event(), threading.Event()
        blocked = ['/b']
        requests = []
        class Site(BaseHTTPRequestHandler):
            def log_message(self,*args): pass
            def do_GET(self):
                requests.append(self.path)
                if self.path in blocked:
                    entered.set()
                    release.wait(30)
                body = b'User-agent: *\nAllow: /\n' if self.path == '/robots.txt' else '<meta charset="utf-8"><main><h1>王明</h1><p>职称：教授</p><p>研究光学。</p></main>'.encode()
                try:
                    self.send_response(200); self.send_header('Content-Type','text/html; charset=utf-8'); self.end_headers()
                    self.wfile.write(body)
                except OSError: pass
        server = ThreadingHTTPServer(('127.0.0.1',0),Site)
        threading.Thread(target=server.serve_forever,daemon=True).start()
        self.addCleanup(server.server_close); self.addCleanup(server.shutdown)
        base = f'http://127.0.0.1:{server.server_port}'
        book = Workbook(); sheet=book.active; sheet.title='Sheet1'
        sheet.append(['ID','学校','学院','师资URL']); sheet.append([1,'虚构大学','理学院',base+'/list'])
        sheet=book.create_sheet('Sheet2');sheet.append(['学院ID','个人URL'])
        for suffix in ('/a','/b','/c'): sheet.append([1,base+suffix])
        book.save(self.input)
        ident=self.cli('import',self.input)['id']
        self.cli('enqueue',ident,'--live')
        def start():
            process=subprocess.Popen([sys.executable,str(ROOT/'batch.py'),'--data-root',str(self.data),'run',ident],stdout=subprocess.PIPE,stderr=subprocess.PIPE)
            self.addCleanup(lambda: process.kill() if process.poll() is None else None)
            return process
        worker=start()
        self.assertTrue(entered.wait(30),'second page not started')
        concurrent=self.cli('run',ident,ok=False)
        self.assertNotEqual(concurrent.returncode,0)
        self.cli('pause',ident)
        release.set()
        out,err=worker.communicate(timeout=30)
        self.assertEqual(worker.returncode,0,err)
        paused=json.loads(out)
        self.assertEqual(paused['counts'].get('pending'),1)
        self.assertEqual(requests.count('/a'),1)
        self.assertEqual(requests.count('/b'),1)
        first_artifact=self.data/paused['tasks'][0]['artifact']
        first_json=(first_artifact/'teachers.json').read_bytes()
        blocked[:]=['/c'];entered.clear();release.clear()
        self.cli('resume',ident)
        worker=start()
        self.assertTrue(entered.wait(30),'third page not started')
        worker.kill();worker.communicate(timeout=10);release.set();blocked.clear()
        finished=self.cli('run',ident)
        self.assertNotIn('pending',finished['counts'])
        self.assertNotIn('running',finished['counts'])
        self.assertEqual(requests.count('/a'),1)
        self.assertEqual(requests.count('/b'),1)
        self.assertEqual((first_artifact/'teachers.json').read_bytes(),first_json)
    def test_university_zip_has_contract_order_full_text_and_relative_evidence(self):
        import zipfile, io
        from openpyxl import load_workbook
        manifest=self.cache()
        long='完整研究经历与成果。'*5000
        (self.folder/'page.html').write_text('<meta charset="utf-8"><main><h1>李明</h1><p>本人邮箱：liming@example.edu</p><p>'+long+'</p></main>',encoding='utf-8')
        import hashlib
        value=json.loads(manifest.read_text());value['entries'][0]['source']['sha256']=hashlib.sha256((self.folder/'page.html').read_bytes()).hexdigest()
        value['entries'] += [{'url':'https://example.edu/failed','source':{'http_status':404,'error':{'code':'http_404','message':'实际访问404'},'observed_at':'2026-10-28T01:00:00Z'}}]
        manifest.write_text(json.dumps(value),encoding='utf-8')
        from openpyxl import load_workbook
        book=load_workbook(self.input);book['Sheet2'].append([1,'https://example.edu/failed']);book['Sheet2'].append([1,'https://example.edu/missing']);book.save(self.input);book.close()
        ident=self.cli('import',self.input)['id']
        done=self.cli('run',ident,'--cache-manifest',manifest)
        self.assertEqual(done['state'],'completed_with_gaps')
        exported=self.cli('export',ident)
        self.assertEqual(len(exported['exports']),1)
        archive=self.data/exported['exports'][0]['path']
        with zipfile.ZipFile(archive) as package:
            book=load_workbook(io.BytesIO(package.read('teachers.xlsx')))
            self.assertEqual(book.sheetnames,['有邮箱教师','无邮箱教师','主页失效教师'])
            self.assertEqual(list(book.worksheets[0].values)[0],('序号','姓名','职称','邮箱','学校','学院','采集人','采集日期','简介','主页'))
            numbers=[row[0] for sheet in book for row in list(sheet.values)[1:]]
            self.assertEqual(numbers,list(range(1,len(numbers)+1)))
            self.assertEqual(book['主页失效教师'].max_row,2)
            self.assertEqual(book['无邮箱教师'].max_row,2)
            records=json.loads(package.read('teachers.json'))
            good=next(r for r in records if r['requested_url']=='https://example.edu/teacher')
            self.assertIn(long,package.read(good['biography_file']).decode())
            self.assertIn('光学院',good['input_metadata']['college']);self.assertIn('理学院',good['input_metadata']['college'])
            self.assertIn(good['source_file'],package.namelist())
            self.assertIn(long,package.read('teachers.csv').decode('utf-8-sig'))
            self.assertIn('relations.csv',package.namelist())
            self.assertIn('batch.json',package.namelist())
    def test_offline_demo_produces_download_without_external_input(self):
        demo=self.cli('demo')
        self.assertEqual(demo['batch']['completeness'],'not_verified')
        self.assertTrue((self.data/demo['delivery']['exports'][0]['path']).is_file())
        self.assertEqual(len(demo['batch']['tasks']),2)
