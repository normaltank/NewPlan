import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]


class ServerContract(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.inputs = self.folder / '中文 inputs'
        self.inputs.mkdir()
        self.data = self.folder / '中文 data'
        self.input = self.inputs / '原表.xlsx'
        self.workbook(self.input)
        source = self.folder / 'page.html'
        source.write_text('<meta charset="utf-8"><main><h1>李明</h1><p>姓名：李明</p><p>职称：教授</p><p>本人邮箱：liming@example.edu</p><p>研究光学。</p></main>', encoding='utf-8')
        self.manifest = self.folder / 'manifest.json'
        self.manifest.write_text(json.dumps({'version': 1, 'entries': [{'url': 'https://example.edu/teacher', 'source': {'path': 'page.html', 'sha256': hashlib.sha256(source.read_bytes()).hexdigest()}}]}), encoding='utf-8')

    def workbook(self, path, college='光学院'):
        book = Workbook()
        sheet = book.active
        sheet.title = 'Sheet1'
        sheet.append(['ID', '学校', '学院', '师资URL'])
        sheet.append([1, '虚构大学', college, 'https://example.edu/list'])
        sheet = book.create_sheet('Sheet2')
        sheet.append(['学院ID', '个人URL'])
        sheet.append([1, 'https://example.edu/teacher'])
        book.save(path)
        book.close()

    def cli(self, entry, *args, expected=0):
        result = subprocess.run([sys.executable, str(ROOT / entry), '--data-root', str(self.data), *map(str, args)], capture_output=True, encoding='utf-8', errors='replace', timeout=60)
        self.assertEqual(result.returncode, expected, result.stderr or result.stdout)
        return json.loads(result.stdout)

    def wait_for_logged_work(self, log, count, nested=False):
        import time
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            events = []
            for line in log.read_text(encoding='utf-8').splitlines():
                try:
                    events.append(json.loads(line))
                except json.JSONDecodeError:
                    pass  # A currently written last line is not a complete event yet.
            if sum(len((event['worker'] if nested else event)['processed']) for event in events) >= count:
                return
            time.sleep(0.05)
        self.fail('worker did not flush its completed public report')

    def test_scan_new_file_does_not_reprocess_unchanged_or_renamed_content(self):
        scanned = self.cli('server.py', 'scan', self.inputs, '--cache-manifest', self.manifest, '--run')
        self.assertEqual(len(scanned['imported']), 1)
        ident = scanned['imported'][0]['id']
        before = self.cli('batch.py', 'status', ident)
        self.assertEqual(before['quality_counts'], {'personal_candidate': 1})
        self.assertNotIn('pending', before['counts'])
        artifact = self.data / before['tasks'][0]['artifact']
        original = {p.relative_to(artifact).as_posix(): p.read_bytes() for p in artifact.rglob('*') if p.is_file()}
        self.assertEqual(self.cli('server.py', 'scan', self.inputs, '--cache-manifest', self.manifest, '--run')['imported'], [])
        shutil.copyfile(self.input, self.inputs / '改名副本.xlsx')
        self.workbook(self.inputs / '新增.xlsx', '理学院')
        third = self.cli('server.py', 'scan', self.inputs, '--cache-manifest', self.manifest, '--run')
        self.assertEqual(len(third['imported']), 1)
        self.assertEqual(len(self.cli('batch.py', 'list')), 2)
        self.assertEqual(original, {p.relative_to(artifact).as_posix(): p.read_bytes() for p in artifact.rglob('*') if p.is_file()})
        self.workbook(self.input, '新学院')
        changed = self.cli('server.py', 'scan', self.inputs)
        self.assertEqual(len(changed['imported']), 1)
        self.assertFalse(changed['imported'][0]['queued'])
        self.assertEqual(len(self.cli('batch.py', 'list')), 3)
        exported = self.cli('batch.py', 'export', ident)
        with zipfile.ZipFile(self.data / exported['exports'][0]['path']) as package:
            self.assertIn('teachers.xlsx', package.namelist())
            self.assertEqual(json.loads(package.read('teachers.json'))[0]['name'], '李明')


    def test_scan_defaults_import_only_and_reports_bad_file_with_exit_status(self):
        (self.inputs / '损坏.xlsx').write_bytes(b'not an Excel workbook')
        scanned = self.cli('server.py', 'scan', self.inputs, expected=1)
        self.assertEqual(len(scanned['imported']), 2)
        self.assertTrue(any(batch['issues'] for batch in scanned['imported']))
        batches = self.cli('batch.py', 'list')
        self.assertTrue(all(not batch['queued'] for batch in batches))
        self.assertFalse((self.data / 'results').exists())
        result = self.cli('server.py', 'worker', '--once')
        self.assertEqual(result['processed'], [])
        self.assertFalse((self.data / 'results').exists())
        invalid = subprocess.run([sys.executable, str(ROOT / 'server.py'), '--data-root', str(self.data), 'scan', str(self.inputs), '--allow-paid'], capture_output=True, timeout=20)
        self.assertEqual(invalid.returncode, 2)

    def test_continuous_worker_respects_shared_queue_pause_and_saved_cache(self):
        import time
        ident = self.cli('batch.py', 'import', self.input)['id']
        self.cli('batch.py', 'enqueue', ident, '--cache-manifest', self.manifest)
        self.cli('batch.py', 'pause', ident)
        self.workbook(self.inputs / '未排队.xlsx', '另学院')
        other = self.cli('batch.py', 'import', self.inputs / '未排队.xlsx')['id']
        self.manifest.unlink()
        (self.folder / 'page.html').unlink()
        log = self.folder / 'worker.jsonl'
        with log.open('w', encoding='utf-8') as out:
            process = subprocess.Popen([sys.executable, str(ROOT / 'server.py'), '--data-root', str(self.data), 'worker', '--interval', '0.1'], stdout=out, stderr=subprocess.DEVNULL)
            try:
                time.sleep(0.8)
                self.assertIsNone(process.poll(), 'continuous worker stopped while queue is paused')
                self.assertEqual(self.cli('batch.py', 'status', ident)['counts'], {'pending': 1})
                self.cli('batch.py', 'resume', ident)
                deadline = time.monotonic() + 30
                while time.monotonic() < deadline:
                    done = self.cli('batch.py', 'status', ident)
                    if not done['queued']:
                        break
                    time.sleep(0.1)
                self.assertFalse(done['queued'])
                self.assertEqual(done['quality_counts'], {'personal_candidate': 1})
                self.assertEqual(self.cli('batch.py', 'status', other)['counts'], {'pending': 1})
                self.wait_for_logged_work(log, 1)
            finally:
                process.terminate()
                process.wait(timeout=15)
        events = [json.loads(line) for line in log.read_text(encoding='utf-8').splitlines()]
        self.assertTrue(any(event['processed'] for event in events))

    def test_watch_discovers_one_new_content_without_reexecuting_previous(self):
        import time
        log = self.folder / 'watch.jsonl'
        with log.open('w', encoding='utf-8') as out:
            process = subprocess.Popen([sys.executable, str(ROOT / 'server.py'), '--data-root', str(self.data), 'watch', str(self.inputs), '--cache-manifest', str(self.manifest), '--interval', '0.1'], stdout=out, stderr=subprocess.DEVNULL)
            try:
                def completed(count):
                    deadline = time.monotonic() + 30
                    batches = []
                    while time.monotonic() < deadline:
                        if process.poll() is not None:
                            self.fail('watch stopped before processing inputs')
                        if self.data.exists():
                            batches = self.cli('batch.py', 'list')
                            if len(batches) == count and all(not batch['queued'] and not batch['counts'].get('pending') for batch in batches):
                                return batches
                        time.sleep(0.1)
                    self.fail(f'watch did not finish {count} batches: {batches}')
                first = completed(1)[0]
                artifact = self.data / first['tasks'][0]['artifact']
                original = (artifact / 'binding.json').read_bytes()
                ready = self.folder / '写完后移入.xlsx'
                self.workbook(ready, '理学院')
                ready.replace(self.inputs / '新文件.xlsx')
                completed(2)
                self.wait_for_logged_work(log, 2, nested=True)
                self.assertEqual((artifact / 'binding.json').read_bytes(), original)
            finally:
                process.terminate()
                process.wait(timeout=15)
        events = [json.loads(line) for line in log.read_text(encoding='utf-8').splitlines()]
        self.assertEqual(sum(len(event['imported']) for event in events), 2)
        self.assertEqual(sum(len(event['worker']['processed']) for event in events), 2)

    def test_worker_process_lock_and_restart_preserve_completed_artifact(self):
        from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
        import threading
        entered, release = threading.Event(), threading.Event()
        requests = []
        blocked = ['/b']
        class Site(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass
            def do_GET(self):
                requests.append(self.path)
                if self.path in blocked:
                    entered.set()
                    release.wait(30)
                body = b'User-agent: *\nAllow: /\n' if self.path == '/robots.txt' else '<meta charset="utf-8"><main><h1>李明</h1><p>姓名：李明</p><p>职称：教授</p><p>研究光学。</p></main>'.encode()
                try:
                    self.send_response(200)
                    self.send_header('Content-Type', 'text/html; charset=utf-8')
                    self.end_headers()
                    self.wfile.write(body)
                except OSError:
                    pass
        server = ThreadingHTTPServer(('127.0.0.1', 0), Site)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        self.addCleanup(server.server_close)
        self.addCleanup(server.shutdown)
        base = f'http://127.0.0.1:{server.server_port}'
        book = Workbook()
        sheet = book.active
        sheet.title = 'Sheet1'
        sheet.append(['ID', '学校', '学院', '师资URL'])
        sheet.append([1, '虚构大学', '理学院', base + '/list'])
        sheet = book.create_sheet('Sheet2')
        sheet.append(['学院ID', '个人URL'])
        sheet.append([1, base + '/a'])
        sheet.append([1, base + '/b'])
        book.save(self.input)
        book.close()
        ident = self.cli('server.py', 'scan', self.inputs, '--live')['imported'][0]['id']
        process = subprocess.Popen([sys.executable, str(ROOT / 'server.py'), '--data-root', str(self.data), 'worker', '--once'], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        try:
            self.assertTrue(entered.wait(30), 'second page did not start')
            first = self.cli('batch.py', 'status', ident)['tasks'][0]
            artifact = self.data / first['artifact'] / 'binding.json'
            original = artifact.read_bytes()
            concurrent = self.cli('server.py', 'worker', '--once')
            self.assertEqual(concurrent['busy'], [ident])
            self.assertEqual(requests.count('/a'), 1)
            self.assertEqual(requests.count('/b'), 1)
        finally:
            process.kill()
            process.wait(timeout=15)
            blocked.clear()
            release.set()
        finished = self.cli('server.py', 'worker', '--once')['processed'][0]
        self.assertFalse(finished['queued'])
        self.assertNotIn('pending', finished['counts'])
        self.assertNotIn('running', finished['counts'])
        self.assertEqual(requests.count('/a'), 1)
        self.assertEqual(artifact.read_bytes(), original)

    def test_platform_launcher_runs_relocated_package_with_external_data_root(self):
        import os
        package = self.folder / '移动 Mission C'
        shutil.copytree(ROOT, package, ignore=shutil.ignore_patterns('.venv', '__pycache__', 'data', 'outputs'))
        arguments = ['--data-root', str(self.data), 'scan', str(self.inputs), '--cache-manifest', str(self.manifest), '--run']
        env = dict(os.environ, MISSION_C_PYTHON=sys.executable, PYTHONUTF8='1')
        if os.name == 'nt':
            command = subprocess.list2cmdline([str(package / 'server.bat'), *arguments])
            result = subprocess.run(command, shell=True, env=env, cwd=self.folder, capture_output=True, encoding='utf-8', errors='replace', timeout=60)
        else:
            result = subprocess.run([str(package / 'server.sh'), *arguments], env=env, cwd=self.folder, capture_output=True, encoding='utf-8', errors='replace', timeout=60)
        self.assertEqual(result.returncode, 0, result.stderr or result.stdout)
        report = json.loads(result.stdout)
        self.assertEqual(len(report['imported']), 1)
        self.assertEqual(report['worker']['processed'][0]['quality_counts'], {'personal_candidate': 1})
        self.assertFalse((package / 'data').exists())

    def test_scan_does_not_enqueue_extra_inputs_from_existing_multifile_batch(self):
        outside = self.folder / '目录外.xlsx'
        self.workbook(outside, '目录外学院')
        ident = self.cli('batch.py', 'import', self.input, outside)['id']
        scanned = self.cli('server.py', 'scan', self.inputs, '--cache-manifest', self.manifest, '--run')
        self.assertEqual(scanned['imported'], [])
        self.assertEqual(scanned['worker']['processed'], [])
        self.assertEqual(self.cli('batch.py', 'status', ident)['counts'], {'pending': 1})
        self.assertFalse(self.cli('batch.py', 'status', ident)['queued'])
        self.assertTrue(scanned['unchanged'][0]['not_queued_reason'])

if __name__ == '__main__':
    unittest.main()
