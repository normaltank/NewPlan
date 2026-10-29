"""Server handoff through its public CLI and unpacked source package."""
from contextlib import redirect_stderr, redirect_stdout
import gc
import io
import json
import os
from pathlib import Path
import runpy
import shutil
import signal
import stat
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

import httpx
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from mission_c.batch import BatchService
ENTRY = ROOT / 'server_distribution' / 'run.py'


class ServerDistribution(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.addCleanup(gc.collect)
        self.folder = Path(temporary.name) / '中文 server space'
        self.folder.mkdir()
        self.inputs = self.folder / '原表 input'
        self.inputs.mkdir()
        self.config = self.folder / 'config.json'
        self.config.write_text(json.dumps({'input_dir': '原表 input', 'data_dir': '持久 data',
                                          'output_dir': '交付 output', 'interval_seconds': 0.01}), encoding='utf-8')

    def workbook(self, name='输入.xlsx', urls=('https://example.edu/one',)):
        path = self.inputs / name
        book = Workbook()
        sheet = book.active
        sheet.title = 'Sheet1'
        sheet.append(['ID', '学校', '学院', '师资URL'])
        sheet.append([1, '虚构大学', '光学院', 'https://example.edu/list'])
        sheet = book.create_sheet('Sheet2')
        sheet.append(['学院ID', '个人URL'])
        for url in urls:
            sheet.append([1, url])
        book.save(path)
        book.close()
        return path

    def cli(self, command, *, entry=ENTRY, config=None, expected=0):
        out, err = io.StringIO(), io.StringIO()
        with patch.object(sys, 'argv', [str(entry), '--config', str(config or self.config), command]), redirect_stdout(out), redirect_stderr(err):
            try:
                runpy.run_path(str(entry), run_name='__main__')
            except SystemExit as result:
                code = result.code
            else:
                code = 0
        self.assertEqual(code, expected, out.getvalue() + err.getvalue())
        return [json.loads(line) for line in out.getvalue().splitlines()], err.getvalue()

    def test_preview_and_demo_are_offline_without_private_credentials(self):
        self.workbook()
        with patch.object(httpx.HTTPTransport, 'handle_request', side_effect=AssertionError('unexpected network')):
            preview, _ = self.cli('preview')
            self.assertEqual(preview[0]['tasks'], 1)
            self.assertFalse((self.folder / '持久 data').exists())
            service = BatchService(self.folder / '持久 data')
            existing = service.import_files([self.inputs / '输入.xlsx'])
            before = service.enqueue(existing['id'], live=True)
            demo, _ = self.cli('demo')
        exported = demo[-1]['deliveries'][0]
        self.assertEqual(service.list_batches(), [before])
        self.assertTrue(Path(demo[-1]['output_dir']).samefile(self.folder / '交付 output' / 'offline-demo'))
        workbook = Path(demo[-1]['output_dir']) / exported['excel']
        book = load_workbook(workbook)
        self.assertEqual(book.sheetnames, ['有邮箱教师', '无邮箱教师', '主页失效教师'])
        self.assertEqual(book.worksheets[0].cell(2, 2).value, '李明')
        book.close()
        with zipfile.ZipFile(Path(demo[-1]['output_dir']) / exported['zip']) as archive:
            self.assertIn('teachers.xlsx', archive.namelist())
            self.assertTrue(any(n.endswith('.txt') for n in archive.namelist()))

    def private_entry(self):
        entry = self.folder / 'run.py'
        shutil.copyfile(ENTRY, entry)
        shutil.copytree(ROOT / 'mission_c', self.folder / 'mission_c', ignore=shutil.ignore_patterns('__pycache__', 'web.py', '*static'))
        private = self.folder / '.private'
        private.mkdir()
        (private / 'deployment.json').write_text(json.dumps({'profile': 'deepseek-private-v1', 'key_file': 'credential.txt'}))
        self.key = 'fake-server-portable-key'
        (private / 'credential.txt').write_text(self.key)
        return entry

    def external(self, calls, outcome=200):
        def respond(transport, request):
            if request.url.host == 'api.deepseek.com':
                self.assertEqual(request.headers['Authorization'], 'Bearer ' + self.key)
                calls.append(request.url.path)
                if outcome == 'unknown':
                    raise httpx.ReadTimeout('synthetic unknown outcome')
                if outcome != 200:
                    return httpx.Response(outcome, json={'error': 'synthetic rejection'})
                if request.url.path.endswith('/fetch'):
                    return httpx.Response(200, json={'title': '李明', 'content': '姓名：李明\n职称：教授\n完整个人经历。'})
                url = json.loads(json.loads(request.content)['messages'][1]['content'])['url']
                return httpx.Response(200, json={'choices': [{'finish_reason': 'stop', 'message': {'content': json.dumps({'url': url, 'name': {'value': '李明', 'evidence': '姓名：李明'}, 'title': {'value': '教授', 'evidence': '职称：教授'}, 'email': {'value': 'liming@example.edu', 'evidence': '本人邮箱：liming@example.edu'}, 'body_selector': 'main'})}}]})
            self.assertEqual(request.url.host, 'example.edu')
            if request.url.path == '/robots.txt':
                return httpx.Response(404)
            html = '<main><h1>李明</h1><p>姓名：李明</p><p>职称：教授</p><p>本人邮箱：liming@example.edu</p><p>研究方向：光学。' + '完整个人经历和科研成果。' * 3000 + '</p></main>'
            return httpx.Response(200, content=html.encode(), headers={'content-type': 'text/html; charset=utf-8'})
        return respond

    def test_explicit_once_uses_private_budget_and_auto_exports_without_rescanning_finished_pages(self):
        entry = self.private_entry()
        self.workbook(urls=('https://example.edu/one', 'https://example.edu/one', 'https://example.edu/two'))
        calls = []
        with patch.object(httpx.HTTPTransport, 'handle_request', self.external(calls)):
            first, stderr = self.cli('once', entry=entry)
            second, _ = self.cli('once', entry=entry)
            status, _ = self.cli('status', entry=entry)
        self.assertEqual(calls, ['/v1/tools/fetch', '/chat/completions'] * 2)
        self.assertEqual(first[-1]['batches'][0]['quality_counts'], {'personal_candidate': 2})
        self.assertEqual(second[-1]['imported'], [])
        self.assertEqual(first[-1]['deliveries'], second[-1]['deliveries'])
        batch = status[0]['batches'][0]
        self.assertEqual(sum(batch['counts'].values()), 2)
        self.assertNotIn('pending', batch['counts'])
        from mission_c.batch import BatchService
        persisted = BatchService(self.folder / '持久 data').status(batch['id'])
        self.assertFalse(persisted['queued'])
        archives = list((self.folder / '交付 output').rglob('sources.zip'))
        self.assertEqual(len(archives), 1)
        with zipfile.ZipFile(archives[0]) as archive:
            records = json.loads(archive.read('teachers.json'))
            self.assertEqual(len(records), 2)
            self.assertGreater(len(records[0]['biography']), 32767)
            self.assertTrue(any(n.endswith('.txt') and len(archive.read(n)) > 32767 for n in archive.namelist()))
            self.assertTrue(all(self.key.encode() not in archive.read(n) for n in archive.namelist()))
        self.assertNotIn(self.key, json.dumps(first) + stderr)

    def test_watch_stops_after_current_page_and_once_resumes_same_persistent_budget(self):
        entry = self.private_entry()
        self.workbook(urls=('https://example.edu/one', 'https://example.edu/two'))
        calls = []
        respond = self.external(calls)
        def interrupt_after_first_page(transport, request):
            response = respond(transport, request)
            if request.url.path == '/chat/completions' and len(calls) == 2:
                signal.raise_signal(signal.SIGINT)
            return response
        with patch.object(httpx.HTTPTransport, 'handle_request', interrupt_after_first_page):
            first, _ = self.cli('watch', entry=entry)
        batch = first[-1]['batches'][0]
        self.assertTrue(batch['queued'])
        self.assertEqual(batch['counts'].get('pending'), 1)
        saved = BatchService(self.folder / '持久 data').status(batch['id'])['run_options']
        self.assertEqual(saved['paid_budget_scope'], batch['id'])
        self.assertEqual(saved['paid_max_calls'], 4)
        with patch.object(httpx.HTTPTransport, 'handle_request', self.external(calls)):
            resumed, _ = self.cli('once', entry=entry)
        self.assertEqual(calls, ['/v1/tools/fetch', '/chat/completions'] * 2)
        self.assertFalse(resumed[-1]['batches'][0]['queued'])
        self.assertEqual(len(resumed[-1]['deliveries']), 1)

    def test_legacy_unscoped_queue_is_blocked_without_rebudgeting_or_network(self):
        entry = self.private_entry()
        path = self.workbook()
        service = BatchService(self.folder / '持久 data')
        batch = service.import_files([path])
        service.enqueue(batch['id'], live=True, fetch_mode='supplement', model_mode='assist', allow_paid=True, paid_max_calls=2)
        before = service.status(batch['id'])['run_options']
        with patch.object(httpx.HTTPTransport, 'handle_request', side_effect=AssertionError('unexpected network')):
            result, _ = self.cli('once', entry=entry, expected=1)
        self.assertEqual(result[-1]['blocked'][0]['id'], batch['id'])
        self.assertEqual(service.status(batch['id'])['run_options'], before)
        self.assertEqual(service.status(batch['id'])['counts'], {'pending': 1})

    def test_restart_recovers_owned_import_before_enqueue_without_adopting_unowned_pending_batch(self):
        entry = self.private_entry()
        self.workbook()
        real_open = io.open
        def fail_before_enqueue(path, *args, **kwargs):
            target = Path(path) if isinstance(path, (str, os.PathLike)) else None
            if target and target.parent.name == 'locks' and target.name != 'input-scan.lock':
                raise OSError('synthetic filesystem interruption before first enqueue')
            return real_open(path, *args, **kwargs)
        with patch('io.open', fail_before_enqueue), patch.object(httpx.HTTPTransport, 'handle_request', side_effect=AssertionError('unexpected network')):
            first, _ = self.cli('once', entry=entry, expected=1)
        service = BatchService(self.folder / '持久 data')
        owned = service.list_batches()[0]
        self.assertEqual(owned['counts'], {'pending': 1})
        self.assertIsNone(owned['run_options'])
        # A same-shaped legacy import without this runner's prior intent stays blocked.
        legacy_path = self.workbook('legacy.xlsx', ('https://example.edu/legacy',))
        legacy = service.import_files([legacy_path])
        calls = []
        with patch.object(httpx.HTTPTransport, 'handle_request', self.external(calls)):
            resumed, _ = self.cli('once', entry=entry, expected=1)
            repeated, _ = self.cli('once', entry=entry, expected=1)
        self.assertEqual(calls, ['/v1/tools/fetch', '/chat/completions'])
        self.assertEqual(resumed[-1]['blocked'], [{'id': legacy['id'], 'error': '旧批次没有固定私人预算 scope；保留原队列和预算，请由维护者核对'}])
        self.assertIsNone(service.status(legacy['id'])['run_options'])
        self.assertEqual(service.status(owned['id'])['counts'].get('pending', 0), 0)
        self.assertEqual(resumed[-1]['deliveries'], repeated[-1]['deliveries'])
        self.assertEqual(len(resumed[-1]['deliveries']), 1)

    def test_restart_recovers_before_batch_id_was_saved_even_when_original_input_was_removed(self):
        entry = self.private_entry()
        source = self.workbook()
        real_open = io.open
        writes = []
        def fail_binding_write(path, *args, **kwargs):
            target = Path(path) if isinstance(path, (str, os.PathLike)) else None
            if target and target.parent.name == 'server-import-intents' and target.suffix == '.tmp':
                writes.append(target)
                if len(writes) == 2:
                    raise OSError('synthetic interruption before saving imported batch ID')
            return real_open(path, *args, **kwargs)
        with patch('io.open', fail_binding_write), patch.object(httpx.HTTPTransport, 'handle_request', side_effect=AssertionError('unexpected network')):
            self.cli('once', entry=entry, expected=1)
        service = BatchService(self.folder / '持久 data')
        batch = service.list_batches()[0]
        self.assertEqual(batch['counts'], {'pending': 1})
        self.assertIsNone(batch['run_options'])
        source.unlink()
        calls = []
        with patch.object(httpx.HTTPTransport, 'handle_request', self.external(calls)):
            report, _ = self.cli('once', entry=entry)
        self.assertEqual(report[-1]['recovered'], [batch['id']])
        self.assertEqual(calls, ['/v1/tools/fetch', '/chat/completions'])
        self.assertEqual(len(report[-1]['deliveries']), 1)

    def test_recovery_preserves_free_queue_saved_after_its_snapshot(self):
        entry = self.private_entry()
        self.workbook()
        real_open = io.open
        def interrupted(path, *args, **kwargs):
            target = Path(path) if isinstance(path, (str, os.PathLike)) else None
            if target and target.parent.name == 'locks' and target.name != 'input-scan.lock':
                raise OSError('synthetic interruption before initial enqueue')
            return real_open(path, *args, **kwargs)
        with patch('io.open', interrupted), patch.object(httpx.HTTPTransport, 'handle_request', side_effect=AssertionError('unexpected network')):
            self.cli('once', entry=entry, expected=1)
        service = BatchService(self.folder / '持久 data')
        batch = service.list_batches()[0]
        peer = []
        def save_peer_queue(path, *args, **kwargs):
            target = Path(path) if isinstance(path, (str, os.PathLike)) else None
            if target and target.parent.name == 'server-import-intents' and target.suffix == '.tmp' and not peer:
                peer.append(service.enqueue(batch['id'], live=True, fetch_mode='off', model_mode='off', allow_paid=False, paid_max_calls=0)['run_options'])
                service.pause(batch['id'])
            return real_open(path, *args, **kwargs)
        with patch('io.open', save_peer_queue), patch.object(httpx.HTTPTransport, 'handle_request', side_effect=AssertionError('unexpected network')):
            report, _ = self.cli('once', entry=entry, expected=1)
            repeated, _ = self.cli('once', entry=entry, expected=1)
        self.assertEqual(len(peer), 1)
        after = service.status(batch['id'])
        self.assertEqual(after['run_options'], peer[0])
        self.assertTrue(after['paused'])
        self.assertEqual(after['counts'], {'pending': 1})
        self.assertEqual(report[-1]['recovered'], [])
        self.assertEqual(report[-1]['blocked'][0]['id'], batch['id'])
        self.assertEqual(repeated[-1]['blocked'], report[-1]['blocked'])
        self.assertFalse(list((service.root / 'paid-attempts').glob('budget-*.json')))

    def test_recovery_never_assigns_private_scope_over_related_global_consumption_or_unknown(self):
        entry = self.private_entry()
        self.workbook(urls=('https://example.edu/one',))
        self.workbook('two.xlsx', ('https://example.edu/two',))
        real_open = io.open
        def fail_enqueue(path, *args, **kwargs):
            target = Path(path) if isinstance(path, (str, os.PathLike)) else None
            if target and target.parent.name == 'locks' and target.name != 'input-scan.lock':
                raise OSError('synthetic interruption before enqueue')
            return real_open(path, *args, **kwargs)
        with patch('io.open', fail_enqueue), patch.object(httpx.HTTPTransport, 'handle_request', side_effect=AssertionError('unexpected network')):
            self.cli('once', entry=entry, expected=1)
        service = BatchService(self.folder / '持久 data')
        batches = service.list_batches()
        from mission_c.cli import main as collect
        source = self.folder / 'synthetic-source.html'
        source.write_text('<main><h1>李明</h1><p>姓名：李明</p><p>职称：教授</p><p>研究光学的完整经历。</p></main>', encoding='utf-8')
        calls = []
        for index, outcome in enumerate((401, 'unknown')):
            with patch.object(httpx.HTTPTransport, 'handle_request', self.external(calls, outcome)), redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
                collect(['--url', batches[index]['tasks'][0]['url'], '--html', str(source), '--output', str(self.folder / ('legacy-' + str(index))),
                         '--fetch-mode', 'supplement', '--model-mode', 'assist', '--allow-paid', '--paid-max-calls', '4',
                         '--paid-attempts', str(service.root / 'paid-attempts')], key_loader=lambda: self.key)
        self.assertEqual(calls, ['/v1/tools/fetch', '/chat/completions'] * 2)
        ledger = {p.name: p.read_bytes() for p in (service.root / 'paid-attempts').glob('*.json')}
        with patch.object(httpx.HTTPTransport, 'handle_request', side_effect=AssertionError('unexpected provider resend')):
            report, _ = self.cli('once', entry=entry, expected=1)
        self.assertEqual(len(report[-1]['blocked']), 2)
        self.assertEqual(report[-1]['recovered'], [])
        self.assertEqual(ledger, {p.name: p.read_bytes() for p in (service.root / 'paid-attempts').glob('*.json')})
        for batch in service.list_batches():
            self.assertIsNone(batch['run_options'])
            self.assertEqual(batch['counts'], {'pending': 1})

    def test_new_import_preserves_free_queue_created_after_scanning(self):
        entry = self.private_entry()
        source = self.workbook()
        service = BatchService(self.folder / '持久 data')
        real_replace = os.replace
        peer = []
        def published(source_path, destination):
            result = real_replace(source_path, destination)
            if Path(destination).parent.name == 'server-import-intents' and not peer:
                batch = service.import_files([source])
                peer.append(service.enqueue(batch['id'], live=True, fetch_mode='off', model_mode='off', allow_paid=False, paid_max_calls=0))
            return result
        calls = []
        with patch('os.replace', published), patch.object(httpx.HTTPTransport, 'handle_request', self.external(calls)):
            report, _ = self.cli('once', entry=entry, expected=1)
            repeated, _ = self.cli('once', entry=entry, expected=1)
        self.assertEqual(calls, [])
        self.assertEqual(len(peer), 1)
        after = service.status(peer[0]['id'])
        self.assertEqual(after['run_options'], peer[0]['run_options'])
        self.assertEqual(after['counts'], {'pending': 1})
        self.assertEqual(report[-1]['blocked'][0]['id'], after['id'])
        self.assertEqual(repeated[-1]['blocked'], report[-1]['blocked'])
        self.assertFalse(list((service.root / 'paid-attempts').glob('budget-*.json')))

    def test_first_enqueue_rechecks_peer_queue_at_batch_lock_boundary(self):
        service = BatchService(self.folder / '持久 data')
        batch = service.import_files([self.workbook()])
        real_open = io.open
        peer = []
        def before_lock(path, *args, **kwargs):
            target = Path(path) if isinstance(path, (str, os.PathLike)) else None
            if target == service.root / 'locks' / (batch['id'] + '.lock') and not peer:
                peer.append(None)
                peer[0] = service.enqueue(batch['id'], live=True)['run_options']
                service.pause(batch['id'])
            return real_open(path, *args, **kwargs)
        with patch('io.open', before_lock):
            with self.assertRaises(ValueError):
                service.enqueue(batch['id'], first_only=True, private_budget=True, live=True,
                                fetch_mode='supplement', model_mode='assist', allow_paid=True, paid_max_calls=2)
        after = service.status(batch['id'])
        self.assertEqual(after['run_options'], peer[0])
        self.assertTrue(after['paused'])
        # Existing explicit callers retain their ability to replace options.
        changed = service.enqueue(batch['id'], live=True, fetch_mode='fallback')
        self.assertEqual(changed['run_options']['fetch_mode'], 'fallback')
        self.assertTrue(changed['paused'])

    def test_first_enqueue_keeps_paid_reservation_lock_untouched_and_import_recoverable(self):
        service = BatchService(self.folder / '持久 data')
        batch = service.import_files([self.workbook()])
        attempts = service.root / 'paid-attempts'
        attempts.mkdir()
        paid_lock = attempts / '.paid.lock'
        paid_lock.write_bytes(b'peer reservation')
        with self.assertRaises(ValueError):
            service.enqueue(batch['id'], live=True, first_only=True)
        self.assertEqual(paid_lock.read_bytes(), b'peer reservation')
        self.assertIsNone(service.status(batch['id'])['run_options'])
        paid_lock.unlink()
        queued = service.enqueue(batch['id'], live=True, first_only=True)
        self.assertTrue(queued['queued'])
        self.assertEqual(queued['counts'], {'pending': 1})
        self.assertFalse(paid_lock.exists())

    def test_first_enqueue_preserves_work_completed_before_its_batch_lock(self):
        service = BatchService(self.folder / '持久 data')
        batch = service.import_files([self.workbook()])
        real_open = io.open
        completed = []
        def before_lock(path, *args, **kwargs):
            target = Path(path) if isinstance(path, (str, os.PathLike)) else None
            if target == service.root / 'locks' / (batch['id'] + '.lock') and not completed:
                completed.append(None)
                service.enqueue(batch['id'], live=True)
                completed[0] = service.run(batch['id'])
            return real_open(path, *args, **kwargs)
        calls = []
        with patch('io.open', before_lock), patch.object(httpx.HTTPTransport, 'handle_request', self.external(calls)), redirect_stderr(io.StringIO()):
            with self.assertRaises(ValueError):
                service.enqueue(batch['id'], first_only=True, private_budget=True, live=True,
                                fetch_mode='supplement', model_mode='assist', allow_paid=True, paid_max_calls=2)
        self.assertEqual(calls, [])
        self.assertIsNone(completed[0]['run_options'])
        self.assertNotIn('pending', completed[0]['counts'])
        self.assertEqual(service.status(batch['id']), completed[0])

    def test_missing_private_key_fails_before_import_and_preview_ignores_invalid_private_metadata(self):
        entry = self.private_entry()
        self.workbook()
        (self.folder / '.private' / 'credential.txt').unlink()
        with patch.object(httpx.HTTPTransport, 'handle_request', side_effect=AssertionError('unexpected network')):
            _, error = self.cli('once', entry=entry, expected=2)
            self.assertIn('凭据缺失', error)
            self.assertFalse((self.folder / '持久 data').exists())
            (self.folder / '.private' / 'deployment.json').write_text('not json')
            preview, _ = self.cli('preview', entry=entry)
            self.assertEqual(preview[0]['tasks'], 1)

    def test_provider_unknown_is_visible_preserves_fulltext_and_is_never_reissued(self):
        entry = self.private_entry()
        self.workbook()
        calls = []
        with patch.object(httpx.HTTPTransport, 'handle_request', self.external(calls, 'unknown')):
            first, _ = self.cli('once', entry=entry, expected=1)
            # Different input bytes create a new batch, while URL attempt identity stays global.
            book = load_workbook(self.inputs / '输入.xlsx')
            book['Sheet1'].cell(2, 3, '另学院')
            book.save(self.inputs / '输入.xlsx')
            book.close()
            second, _ = self.cli('once', entry=entry, expected=1)
        self.assertEqual(calls, ['/v1/tools/fetch', '/chat/completions'])
        self.assertEqual(len(second[-1]['batches']), 2)
        self.assertTrue(first[-1]['batches'][0]['assistance_errors'])
        for item in second[-1]['deliveries']:
            with zipfile.ZipFile(self.folder / '交付 output' / item['zip']) as archive:
                records = json.loads(archive.read('teachers.json'))
                self.assertEqual(len(records), 1)
                self.assertGreater(len(records[0]['biography']), 32767)
                self.assertIn('辅助未完成', archive.read('exceptions.csv').decode('utf-8-sig'))

    def test_exact_package_list_modes_and_relocated_launcher_without_web_dependencies(self):
        source = self.folder / 'isolated source'
        source.mkdir()
        for name in ('server_distribution', 'mission_c'):
            shutil.copytree(ROOT / name, source / name, ignore=shutil.ignore_patterns('__pycache__'))
        for name in ('requirements.txt', 'build_server_package.py'):
            shutil.copyfile(ROOT / name, source / name)
        for name in ('mission_c/cache/page.html', 'mission_c/.private/secret.py', 'mission_c/unlisted.py', 'server_distribution/output/result.txt'):
            path = source / name
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text('excluded-fixture-marker')
        private = self.folder / 'fake-deployment.json'
        private.write_text(json.dumps({'profile': 'deepseek-private-v1', 'key_file': 'fake-key.txt'}))
        key = 'fake-server-builder-key'
        (self.folder / 'fake-key.txt').write_text(key)
        modules = ('__init__.py', 'assisted.py', 'batch.py', 'batch_delivery.py', 'batch_quality.py',
                   'cli.py', 'export.py', 'extract.py', 'fetch.py', 'media.py', 'model_assist.py',
                   'ocr_media.py', 'person_names.py', 'private_deployment.py', 'profile_layouts.py',
                   'quality_fields.py', 'render.py')
        files = {'run.py', 'config.json', 'setup.bat', 'setup.sh', 'start.bat', 'start.sh',
                 'README.md', 'FILES.md', 'requirements.txt'} | {'mission_c/' + n for n in modules}
        for is_private in (False, True):
            destination = self.folder / ('PRIVATE-FAKE.zip' if is_private else 'ordinary.zip')
            result = subprocess.run([sys.executable, str(source / 'build_server_package.py'), '--output', str(destination),
                                     *(['--private-config', str(private)] if is_private else [])], capture_output=True, timeout=30)
            self.assertEqual(result.returncode, 0, result.stderr.decode(errors='replace'))
            self.assertNotIn(key.encode(), result.stdout + result.stderr)
            with zipfile.ZipFile(destination) as archive:
                expected = files | ({'.private/deployment.json', '.private/credential.txt', 'PRIVATE.txt'} if is_private else set())
                self.assertEqual(set(archive.namelist()), {'Mission C Server/' + n for n in expected})
                for item in archive.infolist():
                    self.assertEqual(item.create_system, 3)
                    mode = 0o600 if '/.private/' in item.filename else 0o755 if item.filename.endswith('.sh') else 0o644
                    self.assertEqual(stat.S_IMODE(item.external_attr >> 16), mode)
                    self.assertNotIn(b'excluded-fixture-marker', archive.read(item))
                    if '/.private/credential.txt' not in item.filename:
                        self.assertNotIn(key.encode(), archive.read(item))
                moved = self.folder / ('搬移 private' if is_private else '搬移 ordinary')
                archive.extractall(moved)
            package = moved / 'Mission C Server'
            (package / 'input').mkdir()
            shutil.copyfile(self.workbook(), package / 'input' / '原表.xlsx')
            environment = dict(os.environ, MISSION_C_PYTHON=sys.executable, PYTHONUTF8='1', DEEPSEEK_API_KEY='', JINA_API_KEY='')
            launcher = package / ('start.bat' if os.name == 'nt' else 'start.sh')
            if os.name != 'nt':
                launcher.chmod(0o755)  # zipfile extraction itself does not apply Unix mode bits.
            for command in ('preview', 'demo'):
                argv = [str(launcher), command]
                result = subprocess.run(subprocess.list2cmdline(argv) if os.name == 'nt' else argv, shell=os.name == 'nt',
                                        cwd=self.folder, env=environment, capture_output=True, encoding='utf-8', timeout=30)
                self.assertEqual(result.returncode, 0, result.stderr or result.stdout)
                self.assertNotIn(key, result.stdout + result.stderr)
            self.assertEqual(len(list((package / 'output').rglob('sources.zip'))), 1)
            if is_private:
                self.key = key
                calls = []
                with patch.object(httpx.HTTPTransport, 'handle_request', self.external(calls)):
                    result, _ = self.cli('once', entry=package / 'run.py', config=package / 'config.json')
                self.assertEqual(calls, ['/v1/tools/fetch', '/chat/completions'])
                delivery = result[-1]['deliveries'][0]
                self.assertTrue((package / 'output' / delivery['excel']).is_file())
                with zipfile.ZipFile(package / 'output' / delivery['zip']) as archive:
                    self.assertEqual(len(json.loads(archive.read('teachers.json'))), 1)
                    self.assertTrue(all(key.encode() not in archive.read(n) for n in archive.namelist()))


if __name__ == '__main__':
    unittest.main()
