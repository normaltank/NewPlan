"""Real browser delivery flow, public HTTP state and downloaded artifacts only."""
import csv
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
import zipfile

import httpx
from openpyxl import load_workbook
from test_web_browser import WebServer


@unittest.skipUnless(os.environ.get('MISSION_C_BROWSER_TESTS') == '1', 'opt-in real browser')
class DeliveryBrowser(unittest.TestCase):
    def test_upload_preflight_demo_start_reopen_and_real_downloads(self):
        from playwright.sync_api import sync_playwright
        temporary_dir=tempfile.TemporaryDirectory()
        self.addCleanup(temporary_dir.cleanup)
        if temporary_dir:
            temporary=temporary_dir.name
            folder=Path(temporary)
            server=WebServer(folder)
            self.addCleanup(server.close)
            evidence=Path(os.environ.get('MISSION_C_BROWSER_EVIDENCE',temporary))
            evidence.mkdir(parents=True,exist_ok=True)
            events,errors=[],[]
            with sync_playwright() as play:
                channel=os.environ.get('MISSION_C_BROWSER_CHANNEL','msedge')
                browser=play.chromium.launch(**({'channel':channel} if channel!='chromium' else {}),headless=True)
                context=browser.new_context(viewport={'width':1040,'height':1000})
                context.on('request',lambda req:events.append({'method':req.method,'url':req.url}))
                context.on('page',lambda page:page.on('pageerror',lambda error:errors.append(str(error))))
                page=context.new_page()
                page.goto(server.url)
                # The public example download creates synthetic original-format input.
                sample=httpx.get(server.url+'/api/sample.xlsx',trust_env=False).content
                first,second=folder/'原格式甲.xlsx',folder/'原格式乙.xlsx'
                first.write_bytes(sample)
                second.write_bytes(sample)
                page.locator('#files').set_input_files([str(first),str(second)])
                page.get_by_text('已预检 2 个文件，共 3 个待处理条目。',exact=True).wait_for()
                uploaded=page.url.split('batch=')[1]
                before=httpx.get(server.url+'/api/batches/'+uploaded,trust_env=False).json()
                self.assertEqual(before['preflight']['duplicate_rows'],5)
                self.assertEqual(before['processed'],0)
                self.assertFalse(before['queued'])
                self.assertIn('原格式甲.xlsx',page.locator('#file-names').inner_text())
                page.screenshot(path=str(evidence/'simple-01-upload.png'),full_page=True)
                page.get_by_role('button',name='离线试用 / 查看示例').click()
                page.get_by_text('离线示例已准备，共 3 个待处理条目。',exact=True).wait_for()
                ident=page.url.split('batch=')[1]
                page.get_by_role('button',name='开始离线试用',exact=True).click()
                selected=page.url
                page.close()
                page=context.new_page()
                page.goto(selected)
                page.get_by_text('处理完成',exact=True).wait_for(timeout=30000)
                page.get_by_text('已处理 3 / 3 条',exact=True).wait_for()
                self.assertEqual(page.locator('#progress-percent').inner_text(),'100%')
                self.assertIn('待核或失败 2 条',page.locator('#status-text').inner_text())
                completed=httpx.get(server.url+'/api/batches/'+ident,trust_env=False).json()
                page.reload()
                page.get_by_role('link',name='下载教师 Excel',exact=True).wait_for()
                self.assertEqual(httpx.get(server.url+'/api/batches/'+ident,trust_env=False).json(),completed)
                with page.expect_download() as download:
                    page.get_by_role('link',name='下载教师 Excel',exact=True).click()
                download.value.save_as(str(evidence/'simple-teachers.xlsx'))
                with page.expect_download() as download:
                    page.get_by_role('link',name='下载全文与来源 ZIP',exact=True).click()
                download.value.save_as(str(evidence/'simple-sources.zip'))
                page.screenshot(path=str(evidence/'simple-02-complete.png'),full_page=True)
                page.set_viewport_size({'width':375,'height':812})
                page.screenshot(path=str(evidence/'simple-03-mobile.png'),full_page=True)
                self.assertLessEqual(page.evaluate('document.documentElement.scrollWidth'),375)
                # Reopening the bare URL restores this browser's latest batch.
                page.goto(server.url)
                page.get_by_text('处理完成',exact=True).wait_for()
                self.assertIn(ident,page.url)
                self.assertEqual(errors,[])
                self.assertTrue(all(e['url'].startswith((server.url,'edge://')) for e in events),events)
                (evidence/'simple-browser-events.json').write_text(json.dumps(events,ensure_ascii=False,indent=2),encoding='utf-8')
                browser.close()
                self._audit(evidence)

    def _audit(self,evidence):
        book=load_workbook(evidence/'simple-teachers.xlsx')
        self.assertEqual(book.sheetnames,['有邮箱教师','无邮箱教师','主页失效教师'])
        self.assertEqual(sum(sheet.max_row-1 for sheet in book),1)
        self.assertEqual(book.worksheets[0]['B2'].value,'李明')
        self.assertEqual(book.worksheets[0]['D2'].value,'liming@example.edu')
        book.close()
        with zipfile.ZipFile(evidence/'simple-sources.zip') as archive:
            exceptions=list(csv.DictReader(io.StringIO(archive.read('exceptions.csv').decode('utf-8-sig'))))
            self.assertEqual(len(exceptions),2)
            records=json.loads(archive.read('records.json'))
            teacher=next(r for r in records if r.get('name')=='李明')
            self.assertIn('虚构离线教师正例，研究光学。',archive.read(teacher['biography_file']).decode('utf-8'))
            self.assertIn(b'<main>',archive.read(teacher['source_file']))


if __name__=='__main__':
    unittest.main()
