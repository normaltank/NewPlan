"""Public HTTP contracts, with only synthetic local source pages/cache."""
from collections import Counter
import hashlib
import io
import json
from pathlib import Path
import tempfile
import threading
import time
import unittest
import zipfile

import httpx
from openpyxl import Workbook, load_workbook
from test_web_browser import WebServer


def workbook(path, urls, school='虚构大学'):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Sheet1'
    sheet.append(['ID','学校','学院','师资URL'])
    sheet.append([1,school,'光学院','https://example.edu/list'])
    sheet.append([2,school,'理学院','https://example.edu/science'])
    # A literal text cell can begin with '='; it is not an unevaluated formula.
    sheet['B2'].data_type = 's'
    sheet['B3'].data_type = 's'
    sheet = book.create_sheet('Sheet2')
    sheet.append(['学院ID','个人URL'])
    for ident, url in urls:
        sheet.append([ident,url])
    book.save(path)
    book.close()
    return path


class HTTPFlow(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.folder = Path(temporary.name)

    def server(self, *args):
        server = WebServer(self.folder, *args)
        self.addCleanup(server.close)
        client = httpx.Client(base_url=server.url, headers={'X-Mission-C':'1'}, trust_env=False, timeout=40)
        self.addCleanup(client.close)
        return server, client

    def imported(self, client, path, *others):
        response = client.post('/api/import', files=[('files',(p.name,p.read_bytes())) for p in (path,*others)])
        self.assertEqual(response.status_code,200,response.text)
        return response.json()

    def until(self, client, ident, predicate):
        for _ in range(250):
            response = client.get('/api/batches/'+ident)
            self.assertEqual(response.status_code,200,response.text)
            value = response.json()
            if predicate(value):
                return value
            time.sleep(.1)
        self.fail(f'Batch did not reach expected visible state: {value}')

    def test_multifile_preview_is_not_collection_and_cached_download_preserves_full_text(self):
        page = self.folder/'page.html'
        long = '完整研究经历与成果。'*5000
        page.write_text('<main><h1>李明</h1><p>姓名：李明</p><p>本人邮箱：liming@example.edu</p><p>'+long+'</p></main>',encoding='utf-8')
        manifest = self.folder/'manifest.json'
        manifest.write_text(json.dumps({'version':1,'entries':[{'url':'https://example.edu/li','source':{
            'path':page.name,'sha256':hashlib.sha256(page.read_bytes()).hexdigest(),'observed_at':'2026-10-28T00:00:00Z'}}]}),encoding='utf-8')
        original = workbook(self.folder/'原表.xlsx',[(1,'https://example.edu/li'),(2,'https://example.edu/li#bio'),(99,'https://example.edu/orphan'),(1,'bad')],school='=1+1')
        broken = self.folder/'损坏.xlsx'
        broken.write_bytes(b'not a workbook')
        _, client = self.server('--cache-manifest',str(manifest))
        batch = self.imported(client,original,broken)
        self.assertEqual(batch['preflight']['unique_tasks'],1)
        self.assertEqual(batch['preflight']['issue_count'],3)
        self.assertEqual(batch['preflight']['duplicate_rows'],1)
        ident = batch['id']
        time.sleep(.4)
        self.assertEqual(client.get('/api/batches/'+ident).json()['processed'],0)
        self.assertFalse(batch['queued'])
        self.assertEqual(self.imported(client,original,broken)['id'],ident)
        self.assertEqual(client.get('/api/batches/'+ident+'/issues?limit=1').json()['total'],3)
        started=client.post('/api/batches/'+ident+'/start',json={'cache_id':'local'})
        self.assertEqual(started.status_code,200,started.text)
        done=self.until(client,ident,lambda v:not v['queued'])
        self.assertEqual(done['quality_counts'],{'personal_candidate':1})
        details=client.get('/api/batches/'+ident+'/tasks').json()
        self.assertTrue(details['items'][0]['rawtext_saved'])
        self.assertNotIn(long,client.get('/api/batches/'+ident).text)
        self.assertNotIn('biography',details['items'][0])
        exported=client.post('/api/batches/'+ident+'/exports').json()['exports'][0]
        with zipfile.ZipFile(io.BytesIO(client.get(exported['zip_url']).content)) as package:
            book=load_workbook(io.BytesIO(package.read('teachers.xlsx')))
            self.assertEqual(book.sheetnames,['有邮箱教师','无邮箱教师','主页失效教师'])
            self.assertEqual(book['有邮箱教师']['E2'].value,'=1+1')
            self.assertEqual(book['有邮箱教师']['E2'].data_type,'s')
            self.assertEqual(sum(s.max_row-1 for s in book),1)
            records=json.loads(package.read('records.json'))
            self.assertIn(long,package.read(records[0]['biography_file']).decode('utf-8'))
            self.assertEqual(package.read(records[0]['source_file']),page.read_bytes())
            self.assertIn('exceptions.csv',package.namelist())
        other=client.post('/api/demo').json()['id']
        cross=exported['zip_url'].replace(ident,other)
        self.assertEqual(client.get(cross).status_code,404)

    def test_selecting_downloaded_demo_cache_does_not_create_an_extra_batch(self):
        _,client=self.server()
        sample=client.get('/api/sample.xlsx')
        self.assertEqual(sample.status_code,200)
        original=self.folder/'上传原表.xlsx'
        original.write_bytes(sample.content)
        ident=self.imported(client,original)['id']
        count=len(client.get('/api/batches').json())
        self.assertEqual(count,2)
        # ZIP timestamps can otherwise make an unintended regenerated demo identical.
        time.sleep(2.2)
        self.assertEqual(client.post('/api/batches/'+ident+'/start',json={'cache_id':'demo'}).status_code,200)
        self.until(client,ident,lambda v:not v['queued'])
        self.assertEqual(len(client.get('/api/batches').json()),count)

    def test_pause_after_current_page_restart_and_resume_keeps_finished_evidence(self):
        from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
        calls=Counter()
        entered, release=threading.Event(),threading.Event()
        class Source(BaseHTTPRequestHandler):
            def log_message(self,*args):
                pass
            def do_GET(self):
                if self.path=='/robots.txt':
                    content=b'User-agent: *\nAllow: /\n'
                else:
                    calls[self.path]+=1
                    if self.path=='/b':
                        entered.set()
                        release.wait(25)
                    content='<main><h1>李明</h1><p>姓名：李明</p><p>本人邮箱：liming@example.edu</p><p>虚构离线测试服务器。</p></main>'.encode('utf-8')
                self.send_response(200)
                self.send_header('Content-Type','text/html; charset=utf-8')
                self.send_header('Content-Length',str(len(content)))
                self.end_headers()
                self.wfile.write(content)
        source=ThreadingHTTPServer(('127.0.0.1',0),Source)
        self.addCleanup(source.server_close)
        self.addCleanup(source.shutdown)
        threading.Thread(target=source.serve_forever,daemon=True).start()
        self.addCleanup(release.set)
        base=f'http://127.0.0.1:{source.server_port}'
        original=workbook(self.folder/'慢页.xlsx',[(1,base+'/a'),(1,base+'/b'),(2,base+'/c')])
        server,client=self.server('--allow-live')
        ident=self.imported(client,original)['id']
        self.assertEqual(dict(calls),{})
        response=client.post('/api/batches/'+ident+'/start',json={'mode':'live'})
        self.assertEqual(response.status_code,200,response.text)
        self.assertTrue(entered.wait(25),'second synthetic source page not entered')
        self.assertEqual(client.post('/api/batches/'+ident+'/pause').status_code,200)
        release.set()
        paused=self.until(client,ident,lambda v:v['paused'] and v['running']==0)
        self.assertEqual(paused['processed'],2)
        self.assertTrue(paused['queued'])
        exported=client.post('/api/batches/'+ident+'/exports').json()['exports'][0]
        with zipfile.ZipFile(io.BytesIO(client.get(exported['zip_url']).content)) as package:
            originals={name:package.read(name) for name in package.namelist() if name.startswith('evidence/')}
        server.close()
        restarted,newclient=self.server('--allow-live')
        time.sleep(.5)
        self.assertEqual(newclient.get('/api/batches/'+ident).json()['processed'],2)
        self.assertEqual(dict(calls),{'/a':1,'/b':1})
        self.assertEqual(newclient.post('/api/batches/'+ident+'/resume').status_code,200)
        done=self.until(newclient,ident,lambda v:not v['queued'])
        self.assertEqual(done['processed'],3)
        self.assertEqual(dict(calls),{'/a':1,'/b':1,'/c':1})
        exported=newclient.post('/api/batches/'+ident+'/exports').json()['exports'][0]
        with zipfile.ZipFile(io.BytesIO(newclient.get(exported['zip_url']).content)) as package:
            for name,content in originals.items():
                self.assertEqual(package.read(name),content)

    def test_local_origin_upload_limits_and_server_paid_policy_cannot_be_bypassed(self):
        server,client=self.server()
        self.assertEqual(client.get('/api/config',headers={'Host':'evil.example'}).status_code,403)
        self.assertEqual(client.post('/api/demo',headers={'Origin':'https://evil.example'}).status_code,403)
        self.assertEqual(httpx.post(server.url+'/api/demo',trust_env=False).status_code,403)
        self.assertEqual(client.post('/api/demo',headers={'Sec-Fetch-Site':'cross-site'}).status_code,403)
        original=workbook(self.folder/'输入.xlsx',[(1,'https://example.edu/teacher')])
        for name in ('../escape.xlsx','folder/escape.xlsx','bad.xlsm'):
            response=client.post('/api/import',files={'files':(name,original.read_bytes())})
            self.assertEqual(response.status_code,400,response.text)
        # The multipart parser strips legacy Windows client paths before our app.
        normalized=client.post('/api/import',files={'files':('C:\\escape.xlsx',original.read_bytes())})
        self.assertEqual(normalized.status_code,200,normalized.text)
        sources=client.get('/api/batches/'+normalized.json()['id']+'/tasks').json()['items'][0]['sources']
        self.assertEqual(sources[0]['file'],'escape.xlsx')
        response=client.post('/api/import',files=[('files',(f'{i}.xlsx',b'bad')) for i in range(21)])
        self.assertEqual(response.status_code,400,response.text)
        response=client.post('/api/import',content=b'x'*(22*1024*1024))
        self.assertEqual(response.status_code,413,response.text)
        ident=self.imported(client,original)['id']
        self.assertEqual(client.post('/api/batches/'+ident+'/start',json={'mode':'live'}).status_code,403)
        self.assertEqual(client.post('/api/batches/'+ident+'/start',json={'allow_paid':True,'paid_max_calls':5}).status_code,403)
        self.assertEqual(client.post('/api/batches/'+ident+'/start',json={'cache_manifest':'C:/secret/manifest.json'}).status_code,422)
        self.assertFalse(client.get('/api/batches/'+ident).json()['queued'])
        self.assertEqual(client.get('/api/batches/invalid/downloads/no/zip').status_code,404)
        response=client.get('/')
        self.assertIn("frame-ancestors 'none'",response.headers['Content-Security-Policy'])
        self.assertEqual(response.headers['X-Content-Type-Options'],'nosniff')


if __name__ == '__main__':
    unittest.main()
