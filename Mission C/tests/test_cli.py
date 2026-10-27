"""Public CLI contract. All teacher content below is synthetic."""
import json
import csv
import io
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import os
import shutil
from pathlib import Path
import subprocess
import sys
import tempfile
import threading
import unittest

from openpyxl import load_workbook
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject


ROOT = Path(__file__).resolve().parents[1]


def cv_pdf():
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 50 750 Td (Full curriculum vitae. Last publication in attachment.) Tj 0 -20 Td (Email: liming@example.edu) Tj ET")
    page[NameObject("/Contents")] = content
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def cv_image():
    from PIL import Image, ImageDraw, ImageFont
    canvas = Image.new("RGB", (1100, 220), "white")
    font_path = Path("C:/Windows/Fonts/arial.ttf") if os.name == "nt" else Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    font = ImageFont.truetype(str(font_path), 58)
    ImageDraw.Draw(canvas).text((30, 50), "RESEARCH OPTICS", font=font, fill="black")
    buffer = io.BytesIO()
    canvas.save(buffer, format="PNG")
    return buffer.getvalue()


class FixtureSite(BaseHTTPRequestHandler):
    requests = []

    def do_GET(self):
        type(self).requests.append(self.path)
        if self.path == "/robots.txt":
            data = b"User-agent: *\nDisallow: /private\n"
        elif self.path == "/teacher":
            data = '<meta charset="utf-8"><main><h1>李明</h1><p>职称：讲师</p><p>个人简介：研究光学。</p></main>'.encode()
        elif self.path == "/teacher-media":
            data = '<meta charset="utf-8"><main><h1>李明</h1><p>个人简介</p><a href="/cv.pdf">完整简历</a><img src="/cv.png"></main>'.encode()
        elif self.path == "/dynamic":
            data = '<meta charset="utf-8"><main id="cv"></main><script>document.querySelector("#cv").innerHTML="<h1>王芳</h1><p>研究方向：光学设计。</p>";</script>'.encode()
        elif self.path == "/cv.pdf":
            data = cv_pdf()
        elif self.path == "/cv.png":
            data = cv_image()
        else:
            self.send_response(404)
            self.end_headers()
            return
        self.send_response(200)
        self.send_header("Content-Type", "application/pdf" if self.path == "/cv.pdf" else "text/html; charset=utf-8")
        self.end_headers()
        self.wfile.write(data)

    def log_message(self, *args):
        pass


class MissionCContract(unittest.TestCase):
    def test_separate_profile_sections_are_not_silently_dropped(self):
        with tempfile.TemporaryDirectory() as scratch:
            folder = Path(scratch)
            source = folder / "teacher.html"
            source.write_text('<meta charset="utf-8"><title>周宁 - 示例大学</title><h1>周宁</h1>'
                              '<article><h2>个人简介</h2><p>周宁，教授。</p></article>'
                              '<article><h2>发表论文</h2><p>第二个栏目中的全部论文。</p></article>', encoding="utf-8")
            output = folder / "result"
            run = subprocess.run([sys.executable, str(ROOT / "collect.py"), "--url", "https://example.edu/zhou",
                                  "--html", str(source), "--output", str(output)],
                                 capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(run.returncode, 0, run.stdout + run.stderr)
            record = json.loads((output / "teachers.json").read_text(encoding="utf-8"))[0]
            self.assertIn("第二个栏目中的全部论文。", record["biography"])

    def test_detached_folder_runs_and_name_conflict_does_not_guess(self):
        with tempfile.TemporaryDirectory() as scratch:
            folder = Path(scratch)
            detached = folder / "portable"
            shutil.copytree(ROOT, detached, ignore=shutil.ignore_patterns(".venv", "__pycache__", "outputs"))
            source = folder / "teacher.html"
            source.write_text('<meta charset="utf-8"><main><h1>陈晨</h1><p>职称：教授</p>'
                              '<p>邮箱：chen@example.edu</p><p>完整个人简介。</p></main>', encoding="utf-8")
            output = folder / "result"
            run = subprocess.run([sys.executable, str(detached / "collect.py"), "--url", "https://example.edu/chen",
                                  "--html", str(source), "--name", "李明", "--output", str(output)], cwd=folder,
                                 env={**os.environ, "PYTHONPATH": ""}, capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(run.returncode, 0, run.stdout + run.stderr)
            record = json.loads((output / "teachers.json").read_text(encoding="utf-8"))[0]
            self.assertEqual(record["field_status"]["name"], "conflict")
            self.assertEqual((record["name"], record["title"], record["email"]), ("", "", ""))
            self.assertIn("完整个人简介。", record["biography"])
            original = (output / "teachers.json").read_bytes()
            repeated = subprocess.run([sys.executable, str(detached / "collect.py"), "--url", "https://example.edu/chen",
                                       "--html", str(source), "--output", str(output)], capture_output=True)
            self.assertEqual(repeated.returncode, 2)
            self.assertEqual((output / "teachers.json").read_bytes(), original)

    def test_publication_contact_and_supervisor_rank_are_not_personal_fields(self):
        with tempfile.TemporaryDirectory() as scratch:
            folder = Path(scratch)
            source = folder / "teacher.html"
            source.write_text('<meta charset="utf-8"><main><h1>赵云</h1>'
                              '<p>赵云，博士，硕士生导师。师从李明教授。</p>'
                              '<h2>发表论文</h2><p>Example paper. Email: other@example.edu</p></main>', encoding="utf-8")
            output = folder / "result"
            run = subprocess.run([sys.executable, str(ROOT / "collect.py"), "--url", "https://example.edu/zhao",
                                  "--html", str(source), "--output", str(output)],
                                 capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(run.returncode, 0, run.stdout + run.stderr)
            record = json.loads((output / "teachers.json").read_text(encoding="utf-8"))[0]
            self.assertEqual(record["title"], "")
            self.assertEqual(record["email"], "")
            self.assertIn("other@example.edu", record["biography"])

    @unittest.skipUnless(os.environ.get("MISSION_C_OPTIONAL_TESTS") == "1", "optional browser/OCR integration")
    def test_image_text_is_an_explicit_local_ocr_candidate(self):
        server = ThreadingHTTPServer(("127.0.0.1", 0), FixtureSite)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        self.addCleanup(server.server_close)
        self.addCleanup(server.shutdown)
        with tempfile.TemporaryDirectory() as scratch:
            output = Path(scratch) / "result"
            run = subprocess.run([sys.executable, str(ROOT / "collect.py"), "--ocr",
                                  "--url", f"http://127.0.0.1:{server.server_port}/teacher-media",
                                  "--output", str(output)], capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(run.returncode, 0, run.stdout + run.stderr)
            record = json.loads((output / "teachers.json").read_text(encoding="utf-8"))[0]
            image = next(item for item in record["media"] if item["kind"] == "image")
            self.assertEqual(image["status"], "ocr_candidate")
            self.assertIn("RESEARCH", image["text"])
            self.assertIn("OCR 待复核", record["biography"])

    @unittest.skipUnless(os.environ.get("MISSION_C_OPTIONAL_TESTS") == "1", "optional browser/OCR integration")
    def test_dynamic_page_can_be_explicitly_rendered(self):
        server = ThreadingHTTPServer(("127.0.0.1", 0), FixtureSite)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        self.addCleanup(server.server_close)
        self.addCleanup(server.shutdown)
        with tempfile.TemporaryDirectory() as scratch:
            output = Path(scratch) / "result"
            run = subprocess.run([sys.executable, str(ROOT / "collect.py"), "--render",
                                  "--url", f"http://127.0.0.1:{server.server_port}/dynamic",
                                  "--output", str(output)], capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(run.returncode, 0, run.stdout + run.stderr)
            record = json.loads((output / "teachers.json").read_text(encoding="utf-8"))[0]
            self.assertEqual(record["name"], "王芳")
            self.assertIn("光学设计。", record["research"])
            self.assertTrue((output / record["rendered_file"]).is_file())

    def test_profile_pdf_full_text_is_saved_and_unread_images_are_reported(self):
        server = ThreadingHTTPServer(("127.0.0.1", 0), FixtureSite)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        self.addCleanup(server.server_close)
        self.addCleanup(server.shutdown)
        with tempfile.TemporaryDirectory() as scratch:
            output = Path(scratch) / "result"
            run = subprocess.run([sys.executable, str(ROOT / "collect.py"),
                                  "--url", f"http://127.0.0.1:{server.server_port}/teacher-media",
                                  "--output", str(output)], capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(run.returncode, 0, run.stdout + run.stderr)
            record = json.loads((output / "teachers.json").read_text(encoding="utf-8"))[0]
            self.assertIn("Last publication in attachment.", record["biography"])
            self.assertEqual(record["email"], "liming@example.edu")
            pdf = next(item for item in record["media"] if item["kind"] == "pdf")
            self.assertEqual((output / pdf["source_file"]).read_bytes(), cv_pdf())
            self.assertEqual(pdf["status"], "text_saved")
            image = next(item for item in record["media"] if item["kind"] == "image")
            self.assertEqual(image["status"], "not_read")
            self.assertIn("media_not_fully_read", record["warnings"])

    def test_batch_urls_keep_success_and_report_unavailable_or_disallowed_pages(self):
        server = ThreadingHTTPServer(("127.0.0.1", 0), FixtureSite)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        self.addCleanup(server.server_close)
        self.addCleanup(server.shutdown)
        FixtureSite.requests = []
        base = f"http://127.0.0.1:{server.server_port}"
        with tempfile.TemporaryDirectory() as scratch:
            output = Path(scratch) / "result"
            run = subprocess.run([sys.executable, str(ROOT / "collect.py"),
                                  "--url", base + "/teacher", "--url", base + "/missing",
                                  "--url", base + "/private", "--output", str(output)],
                                 capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(run.returncode, 1, run.stdout + run.stderr)
            rows = json.loads((output / "teachers.json").read_text(encoding="utf-8"))
            self.assertEqual(len(rows), 3)
            self.assertEqual(rows[0]["name"], "李明")
            self.assertEqual(rows[0]["status"], "text_saved")
            self.assertEqual(rows[1]["error"]["code"], "http_404")
            self.assertEqual(rows[2]["error"]["code"], "robots_disallowed")
            self.assertNotIn("/private", FixtureSite.requests)

    def test_saved_homepage_retains_complete_biography_and_evidence(self):
        with tempfile.TemporaryDirectory() as scratch:
            folder = Path(scratch)
            source = folder / "teacher.html"
            tail = "最后一项成果：用于验证长简介没有被截断。"
            source.write_text('''<!doctype html><html><head><meta charset="utf-8">
                <title>张敏 - 示例大学</title></head><body>
                <nav>教师导航</nav><main><h1>张敏</h1><article>
                <p>职称：副教授</p><p>邮箱：zhangmin@example.edu</p>
                <h2>个人简介</h2><p>张敏，博士，硕士生导师。</p>
                <h2>研究方向</h2><p>光电测量与信号处理。</p>
                <h2>发表论文</h2><p>示例论文，2025，第一作者。</p>
                <p>''' + ("完整原文段落。" * 6000) + tail + '''</p></article></main>
                <footer>学院邮箱：office@example.edu</footer></body></html>''', encoding="utf-8")
            output = folder / "result"
            run = subprocess.run([sys.executable, str(ROOT / "collect.py"),
                                  "--url", "https://faculty.example.edu/zhangmin",
                                  "--html", str(source), "--output", str(output)],
                                 capture_output=True, text=True, encoding="utf-8",
                                 env={**os.environ, "PYTHONIOENCODING": "utf-8"})
            self.assertEqual(run.returncode, 0, run.stdout + run.stderr)
            record = json.loads((output / "teachers.json").read_text(encoding="utf-8"))[0]
            self.assertEqual(record["name"], "张敏")
            self.assertEqual(record["title"], "副教授")
            self.assertEqual(record["email"], "zhangmin@example.edu")
            self.assertIn("光电测量与信号处理。", record["research"])
            self.assertIn("示例论文，2025，第一作者。", record["publications"])
            self.assertIn(tail, record["biography"])
            self.assertIn(tail, (output / record["biography_file"]).read_text(encoding="utf-8"))
            self.assertEqual((output / record["source_file"]).read_bytes(), source.read_bytes())
            self.assertEqual(record["source_mode"], "saved_html")
            self.assertTrue(record["observed_at"])
            self.assertEqual(record["requested_url"], "https://faculty.example.edu/zhangmin")
            self.assertTrue(record["field_evidence"]["title"]["evidence"])

    def test_export_keeps_long_text_and_treats_sheet_values_as_data(self):
        with tempfile.TemporaryDirectory() as scratch:
            folder = Path(scratch)
            source = folder / "teacher.html"
            source.write_text('<meta charset="utf-8"><main><h1>王芳</h1><p>职称：教授</p><p>'
                              + "连续原文。" * 7000 + '最后的出版记录。</p></main>', encoding="utf-8")
            output = folder / "result"
            run = subprocess.run([sys.executable, str(ROOT / "collect.py"), "--url", "https://example.edu/wang",
                                  "--html", str(source), "--output", str(output), "--school", "=1+2"],
                                 capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(run.returncode, 0, run.stdout + run.stderr)
            with (output / "teachers.csv").open(encoding="utf-8-sig", newline="") as stream:
                row = next(csv.DictReader(stream))
            self.assertIn("最后的出版记录。", row["简介"])
            workbook = load_workbook(output / "teachers.xlsx")
            self.assertEqual(workbook.sheetnames, ["有邮箱教师", "无邮箱教师", "主页失效教师"])
            worksheet = workbook["无邮箱教师"]
            self.assertEqual(worksheet.cell(2, 5).data_type, "s")
            self.assertEqual(worksheet.cell(2, 5).value, "=1+2")
            self.assertIn("完整原文", worksheet.cell(2, 9).value)
            self.assertIn("0001/biography.txt", worksheet.cell(2, 9).value)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
