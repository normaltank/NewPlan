"""Persistent Excel batches, shared by CLI, local UI and unattended runners."""
from contextlib import contextmanager
import hashlib
import json
from pathlib import Path
import shutil
import sqlite3
import uuid
from urllib.parse import urlsplit, urlunsplit

from openpyxl import load_workbook
from .fetch import valid_url


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def stable_url(value):
    parts = urlsplit(valid_url(str(value).strip()))
    return urlunsplit((parts.scheme.lower(), parts.netloc.lower(), parts.path or '/', parts.query, ''))


def text(value):
    return '' if value is None else str(value).strip()


def identity(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return text(value)


def preview(paths):
    """Read original two-sheet workbooks. Never fetch URLs or write state."""
    files = []
    for item in paths:
        path = Path(item)
        files.extend(sorted(path.glob('*.xlsx')) if path.is_dir() else [path])
    tasks, issues, inputs = {}, [], []
    for path in files:
        if path.name.startswith('~$'):
            continue
        source = {'file': path.name}
        try:
            sha = digest(path)
            inputs.append({'file': path.name, 'sha256': sha, 'path': str(path.resolve())})
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                if not {'Sheet1', 'Sheet2'} <= set(book.sheetnames):
                    raise ValueError('需要 Sheet1 和 Sheet2')
                headers1 = [text(v) for v in next(book['Sheet1'].values)]
                headers2 = [text(v) for v in next(book['Sheet2'].values)]
                if headers1[:4] != ['ID', '学校', '学院', '师资URL'] or headers2[:2] != ['学院ID', '个人URL']:
                    raise ValueError('表头必须为 ID/学校/学院/师资URL 和 学院ID/个人URL')
                colleges = {}
                for row, values in enumerate(book['Sheet1'].values, 1):
                    if row == 1 or not any(v is not None for v in values): continue
                    ident, school, college, url = (list(values) + [None]*4)[:4]
                    rel = {'id': identity(ident), 'school': text(school), 'college': text(college), 'faculty_urls': [], 'sources': []}
                    location = dict(source, sheet='Sheet1', row=row, sha256=sha)
                    if not rel['id'] or not rel['school'] or not rel['college']:
                        issues.append(dict(location, reason='学院ID/学校/学院缺失')); continue
                    if url: rel['faculty_urls'].append(text(url))
                    rel['sources'].append(location)
                    prior = colleges.get(rel['id'])
                    if prior and (prior['school'], prior['college']) != (rel['school'], rel['college']):
                        prior['conflict'] = True
                        issues.append(dict(location, reason='学院ID关联冲突')); continue
                    if prior:
                        prior['faculty_urls'] = list(dict.fromkeys(prior['faculty_urls'] + rel['faculty_urls']))
                        prior['sources'].append(location)
                    else: colleges[rel['id']] = rel
                for row, values in enumerate(book['Sheet2'].values, 1):
                    if row == 1 or not any(v is not None for v in values): continue
                    ident, url = (list(values) + [None]*2)[:2]
                    location = dict(source, sheet='Sheet2', row=row, sha256=sha, college_id=identity(ident), input_url=text(url))
                    rel = colleges.get(identity(ident))
                    if not rel or rel.get('conflict'):
                        issues.append(dict(location, reason='学院关联缺失或冲突')); continue
                    try: url = stable_url(url)
                    except ValueError as error:
                        issues.append(dict(location, reason=str(error))); continue
                    key = hashlib.sha256((rel['school']+'\0'+url).encode()).hexdigest()
                    task = tasks.setdefault(key, {'id': key, 'school': rel['school'], 'url': url, 'relations': [], 'sources': []})
                    task['sources'].append(location)
                    if rel not in task['relations']: task['relations'].append(rel)
            finally: book.close()
        except (OSError, ValueError, KeyError, StopIteration) as error:
            issues.append(dict(source, reason=str(error)))
        except Exception as error:
            # A malformed workbook is isolated; other files remain importable.
            issues.append(dict(source, reason=f'{type(error).__name__}: {error}'))
    return {'inputs': inputs, 'tasks': list(tasks.values()), 'issues': issues, 'completeness': 'not_verified'}


class BatchService:
    def __init__(self, data_root):
        self.root = Path(data_root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        with self.connect() as db:
            db.executescript('''CREATE TABLE IF NOT EXISTS batches(id TEXT PRIMARY KEY, fingerprint TEXT UNIQUE, payload TEXT, paused INTEGER DEFAULT 0);
CREATE TABLE IF NOT EXISTS queue(batch TEXT PRIMARY KEY, options TEXT);
CREATE TABLE IF NOT EXISTS tasks(batch TEXT, id TEXT, payload TEXT, state TEXT DEFAULT 'pending', artifact TEXT, error TEXT, PRIMARY KEY(batch,id));''')

    def connect(self):
        db = sqlite3.connect(self.root/'batches.sqlite3', timeout=30)
        db.row_factory = sqlite3.Row
        return db

    def preview(self, paths):
        return preview(paths)

    def import_files(self, paths):
        value = preview(paths)
        fingerprint = hashlib.sha256(json.dumps(sorted((i['file'], i['sha256']) for i in value['inputs'])).encode()).hexdigest()
        with self.connect() as db:
            db.execute('BEGIN IMMEDIATE')
            existing = db.execute('SELECT id FROM batches WHERE fingerprint=?', (fingerprint,)).fetchone()
            if existing: return self.status(existing['id'])
            batch_id = uuid.uuid4().hex
            for item in value['inputs']:
                target = Path('imports') / batch_id / item['sha256'] / item['file']
                (self.root/target).parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(item.pop('path'), self.root/target)
                item['stored_file'] = target.as_posix()
            db.execute('INSERT INTO batches(id,fingerprint,payload) VALUES(?,?,?)', (batch_id, fingerprint, json.dumps(value, ensure_ascii=False)))
            for task in value['tasks']:
                db.execute('INSERT INTO tasks(batch,id,payload) VALUES(?,?,?)', (batch_id,task['id'],json.dumps(task,ensure_ascii=False)))
        return self.status(batch_id)

    def status(self, batch_id):
        with self.connect() as db:
            row = db.execute('SELECT * FROM batches WHERE id=?', (batch_id,)).fetchone()
            if row is None: raise ValueError('未知批次')
            result = json.loads(row['payload'])
            setting = db.execute('SELECT options FROM queue WHERE batch=?',(batch_id,)).fetchone()
            result.update(id=batch_id, paused=bool(row['paused']), tasks=[], counts={}, quality_counts={}, queued=bool(setting), run_options=json.loads(setting['options']) if setting else None)
            for taskrow in db.execute('SELECT * FROM tasks WHERE batch=? ORDER BY rowid', (batch_id,)):
                task = json.loads(taskrow['payload'])
                state = taskrow['state']
                if state in {'success','review','failed'} and taskrow['artifact']:
                    try: self.bound(taskrow['artifact'])
                    except (ValueError,OSError,KeyError): state = 'artifact_missing'
                task.update(state=state, artifact=taskrow['artifact'], error=taskrow['error'])
                quality = {'category':'not_read','eligible_for_workbook':False,'reasons':['尚未完成或产物不可用']}
                if taskrow['artifact'] and state not in {'pending','running','artifact_missing','program_error'}:
                    try: quality = self.bound(taskrow['artifact']).get('quality',quality)
                    except (ValueError,OSError,KeyError): pass
                task['quality'] = quality
                category = quality['category']
                result['quality_counts'][category] = result['quality_counts'].get(category,0)+1
                result['tasks'].append(task)
                result['counts'][state] = result['counts'].get(state,0)+1
            result['state'] = 'paused' if result['paused'] else 'pending' if any(t['state'] in {'pending','running'} for t in result['tasks']) else 'program_error' if any(t['state'] == 'program_error' for t in result['tasks']) else 'completed_with_gaps' if result['issues'] or any(t['state'] != 'success' for t in result['tasks']) else 'completed'
            return result

    def list_batches(self):
        with self.connect() as db: ids = [r['id'] for r in db.execute('SELECT id FROM batches ORDER BY rowid')]
        return [self.status(ident) for ident in ids]

    def pause(self, batch_id):
        self.status(batch_id)
        with self.connect() as db: db.execute('UPDATE batches SET paused=1 WHERE id=?', (batch_id,))
        return self.status(batch_id)

    def resume(self, batch_id):
        self.status(batch_id)
        with self.connect() as db: db.execute('UPDATE batches SET paused=0 WHERE id=?', (batch_id,))
        return self.status(batch_id)

    @contextmanager
    def worker_lock(self, batch_id):
        """Kernel-held lock: released on process death, never stolen by timeout."""
        path = self.root/'locks'/f'{batch_id}.lock'
        path.parent.mkdir(exist_ok=True)
        stream = path.open('a+b')
        stream.seek(0)
        if not stream.read(1): stream.write(b'0'); stream.flush()
        stream.seek(0)
        try:
            if __import__('os').name == 'nt':
                import msvcrt
                msvcrt.locking(stream.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(stream, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError:
            stream.close()
            raise ValueError('批次正在由另一进程执行')
        try: yield
        finally: stream.close()

    def bound(self, artifact):
        folder = self.root/artifact
        binding = json.loads((folder/'binding.json').read_text(encoding='utf-8'))
        for relative, sha in binding['files'].items():
            path = (folder/relative).resolve()
            if not path.is_relative_to(folder.resolve()) or digest(path) != sha:
                raise ValueError('产物缺失或内容改变')
        return binding

    def run(self, batch_id, *, cache_manifest=None, live=False, max_items=None,
            fetch_mode='off', model_mode='off', allow_paid=False, paid_max_calls=0):
        """Process pending pages with the existing public collector, one commit per page."""
        import contextlib
        import sys
        from .cli import main as collect
        from .assisted import atomic_json
        if cache_manifest or live:
            self.enqueue(batch_id, cache_manifest=cache_manifest, live=live, fetch_mode=fetch_mode,
                         model_mode=model_mode, allow_paid=allow_paid, paid_max_calls=paid_max_calls)
            cache_manifest, live = None, False
        saved = self.status(batch_id)['run_options']
        if not cache_manifest and not live and saved:
            cache_manifest = self.root/saved['cache_manifest'] if saved.get('cache_manifest') else None
            live, fetch_mode, model_mode = saved['live'], saved['fetch_mode'], saved['model_mode']
            allow_paid, paid_max_calls = saved['allow_paid'], saved['paid_max_calls']
        if not live and not cache_manifest:
            raise ValueError('严格缓存模式需要 cache_manifest；在线必须显式 live=True')
        if live and cache_manifest:
            raise ValueError('live 与 cache_manifest 互斥')
        if max_items is not None and max_items < 1: raise ValueError('max_items 必须为正数')
        self.status(batch_id)
        with self.worker_lock(batch_id):
            processed = 0
            for task in self.status(batch_id)['tasks']:
                with self.connect() as db:
                    paused = db.execute('SELECT paused FROM batches WHERE id=?',(batch_id,)).fetchone()['paused']
                if paused: break
                if max_items is not None and processed >= max_items: break
                if task['state'] not in {'pending','running'}: continue
                # A process may have died after writing a binding but before committing.
                if task['state'] == 'running' and task['artifact']:
                    try:
                        binding = self.bound(task['artifact'])
                    except (ValueError,OSError): pass
                    else:
                        with self.connect() as db:
                            db.execute('UPDATE tasks SET state=?,error=? WHERE batch=? AND id=?', (binding['state'],binding['error'],batch_id,task['id']))
                        continue
                relative = Path('results')/batch_id/task['id']/uuid.uuid4().hex
                output = self.root/relative
                output.parent.mkdir(parents=True, exist_ok=True)
                with self.connect() as db:
                    db.execute("UPDATE tasks SET state='running',artifact=?,error=NULL WHERE batch=? AND id=?", (relative.as_posix(),batch_id,task['id']))
                args = ['--url', task['url'], '--school', task['school'], '--college', '; '.join(dict.fromkeys(r['college'] for r in task['relations'])), '--output', str(output), '--fetch-mode', fetch_mode, '--model-mode', model_mode]
                if cache_manifest: args += ['--cache-manifest',str(Path(cache_manifest).resolve())]
                if allow_paid:
                    args += ['--allow-paid','--paid-max-calls',str(paid_max_calls),'--paid-attempts',str(self.root/'paid-attempts')]
                state, error = 'failed', None
                try:
                    with contextlib.redirect_stdout(sys.stderr): code = collect(args)
                    records = json.loads((output/'teachers.json').read_text(encoding='utf-8'))
                    if len(records) != 1 or records[0]['requested_url'] != task['url']:
                        raise ValueError('采集产物与任务不匹配')
                    record = records[0]
                    state = 'success' if record['status'] == 'text_saved' and not record['review_required'] else 'review' if record['biography'] else 'failed'
                    error = json.dumps(record.get('error'),ensure_ascii=False) if record.get('error') else None
                    if code not in (0,1): raise ValueError('采集器异常退出')
                    from .batch_quality import assess
                    quality = assess(record, output)
                    if not quality['eligible_for_workbook'] and state == 'success': state = 'review'
                    binding = {'quality':quality,'state':state,'error':error,'url':task['url'],'files':{p.relative_to(output).as_posix():digest(p) for p in output.rglob('*') if p.is_file()}}
                    atomic_json(output/'binding.json', binding)
                except (Exception, SystemExit) as exception:
                    error = f'{type(exception).__name__}: {exception}'
                    state = 'program_error'
                with self.connect() as db:
                    db.execute('UPDATE tasks SET state=?,error=? WHERE batch=? AND id=?', (state,error,batch_id,task['id']))
                processed += 1
            if not any(t['state'] in {'pending','running'} for t in self.status(batch_id)['tasks']):
                with self.connect() as db: db.execute('DELETE FROM queue WHERE batch=?',(batch_id,))
        return self.status(batch_id)


    def enqueue(self, batch_id, *, cache_manifest=None, live=False, fetch_mode='off',
                model_mode='off', allow_paid=False, paid_max_calls=0):
        """Persist explicit run intent and self-contained offline inputs; never starts work."""
        self.status(batch_id)
        if bool(cache_manifest) == bool(live): raise ValueError('选择缓存清单或显式 live')
        if fetch_mode not in {'off','fallback','supplement'} or model_mode not in {'off','assist'}:
            raise ValueError('无效采集模式')
        if paid_max_calls < 0: raise ValueError('付费预算不能为负数')
        options = dict(live=live, cache_manifest=None, fetch_mode=fetch_mode, model_mode=model_mode,
                       allow_paid=allow_paid, paid_max_calls=paid_max_calls)
        with self.worker_lock(batch_id):
            if cache_manifest:
                from .assisted import CacheManifest
                cache_manifest = Path(cache_manifest).resolve()
                cache = CacheManifest(cache_manifest)
                value = json.loads(cache_manifest.read_text(encoding='utf-8-sig'))
                folder = Path('cache')/batch_id/uuid.uuid4().hex
                (self.root/folder).mkdir(parents=True)
                def snapshot(node):
                    if isinstance(node, dict):
                        if 'path' in node:
                            raw = cache.read(node)
                            target = hashlib.sha256(raw).hexdigest()+Path(node['path']).suffix
                            (self.root/folder/target).write_bytes(raw)
                            node['path'] = target
                        for item in node.values(): snapshot(item)
                    elif isinstance(node,list):
                        for item in node: snapshot(item)
                snapshot(value)
                manifest = folder/'manifest.json'
                (self.root/manifest).write_text(json.dumps(value,ensure_ascii=False),encoding='utf-8')
                options['cache_manifest'] = manifest.as_posix()
            with self.connect() as db:
                db.execute('INSERT OR REPLACE INTO queue VALUES(?,?)',(batch_id,json.dumps(options)))
        return self.status(batch_id)

    def export(self, batch_id, school=None):
        from .batch_delivery import export_universities
        with self.worker_lock(batch_id):
            return export_universities(self,batch_id,school)

    def demo(self):
        """A self-contained synthetic import/queue/run/download demonstration."""
        from openpyxl import Workbook
        folder = self.root/'demo'/uuid.uuid4().hex
        folder.mkdir(parents=True)
        book = Workbook()
        sheet=book.active;sheet.title='Sheet1'
        sheet.append(['ID','学校','学院','师资URL'])
        sheet.append([1,'虚构示例大学','光学院','https://example.edu/list'])
        sheet.append([2,'虚构示例大学','理学院','https://example.edu/science'])
        sheet=book.create_sheet('Sheet2');sheet.append(['学院ID','个人URL'])
        sheet.append([1,'https://example.edu/demo']);sheet.append([2,'https://example.edu/demo'])
        sheet.append([1,'https://example.edu/unavailable'])
        book.save(folder/'示例.xlsx');book.close()
        source=folder/'page.html'
        source.write_text('<meta charset="utf-8"><main><h1>示例教师</h1><p>职称：教授</p><p>本人邮箱：demo@example.edu</p><h2>完整经历</h2><p>虚构离线材料，仅供功能演示。</p></main>',encoding='utf-8')
        manifest={'version':1,'entries':[{'url':'https://example.edu/demo','source':{'path':'page.html','sha256':digest(source),'observed_at':'2026-10-28T00:00:00Z'}},{'url':'https://example.edu/unavailable','source':{'http_status':404,'error':{'code':'http_404','message':'虚构的实际访问失败'},'observed_at':'2026-10-28T00:00:00Z'}}]}
        (folder/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False),encoding='utf-8')
        batch=self.import_files([folder/'示例.xlsx'])
        self.enqueue(batch['id'],cache_manifest=folder/'manifest.json')
        result=self.run(batch['id'])
        return {'batch':result,'delivery':self.export(batch['id'])}
