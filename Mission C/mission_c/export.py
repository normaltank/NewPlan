"""Human-readable exports; JSON and TXT are the unabridged record."""
import csv
import json
import re

from openpyxl import Workbook
from openpyxl.comments import Comment

HEADERS = ["序号", "姓名", "职称", "邮箱", "学校", "学院", "采集人", "采集日期", "简介", "主页"]
EXTRA = ["研究方向", "发表信息", "状态", "字段状态", "待复核事项", "完整简介文件", "原件文件"]


def row_values(index, record):
    metadata = record["input_metadata"]
    return [index, record["name"], record["title"], record["email"], metadata["school"],
            metadata["college"], metadata["collector"], record["observed_at"][:10],
            record["biography"], record["requested_url"]]


def csv_value(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def excel_value(value):
    # XML 1.0 cannot store these controls. Unmodified values remain in JSON/TXT.
    return re.sub(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]", "", value) if isinstance(value, str) else value


def export_tables(output, records):
    with (output / "teachers.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(HEADERS + EXTRA)
        for index, record in enumerate(records, 1):
            extra = [record["research"], record["publications"], record["status"],
                     json.dumps(record.get("field_status", {}), ensure_ascii=False),
                     "; ".join(record["warnings"]) + ("; " + record["error"]["code"] if record.get("error") else ""),
                     record["biography_file"], record["source_file"]]
            writer.writerow([csv_value(value) for value in row_values(index, record) + extra])
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title in ("有邮箱教师", "无邮箱教师", "主页失效教师"):
        sheet = workbook.create_sheet(title)
        sheet.append(HEADERS)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1:J1"
    for index, record in enumerate(records, 1):
        code = record.get("error", {}).get("code", "")
        unavailable = code.startswith("http_") or code in {"network_error", "redirect_limit"}
        title = "主页失效教师" if unavailable else ("有邮箱教师" if record["email"] else "无邮箱教师")
        sheet = workbook[title]
        values = row_values(index, record)
        if len(record["biography"].encode("utf-16-le")) // 2 > 32700:
            prefix = record["biography"].encode("utf-16-le")[:64000].decode("utf-16-le", errors="ignore")
            values[8] = prefix + f"\n[单元格仅展示前段；完整原文：{record['biography_file']}；也见 teachers.json / teachers.csv]"
        sheet.append([excel_value(value) for value in values])
        row = sheet.max_row
        for cell in sheet[row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
        sheet.cell(row, 2).comment = Comment(
            "本表为待复核个人页结果。学校/学院为调用者输入，未经网站核验。\n"
            + json.dumps({"status": record["status"], "warnings": record["warnings"],
                          "fields": record.get("field_status"), "error": record.get("error")}, ensure_ascii=False),
            "Mission C")
    workbook.save(output / "teachers.xlsx")
    workbook.close()
