"""Human-facing tables, separate from the complete evidence/audit JSON."""
from identity_export import consolidate, merge_rows, MERGE_HEADERS
import csv
import io
import zipfile
import hashlib
import re
import json
from datetime import datetime
from urllib.parse import urlparse
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

TEACHER_HEADERS=['序号','姓名','职称','邮箱','学校','学院','采集人','采集日期','简介','主页']
SOURCE_HEADERS=['大学','学院名','学院链接','师资链接']
STATUS_HEADERS=['大学','学院名','运行状态','学院完整性','师资入口情况','运行ID']
NOTE_HEADERS=['学校','学院','姓名','交付标注','未读字段 / 缺失原因','采用材料日期','采用主页','保留材料','复核路径']
STATES={'queued':'等待采集','running':'采集中','finished':'运行结束','failed':'运行失败','interrupted':'已中断'}
COMPLETENESS={'unverified':'完整性未验证','verified':'人工确认完整','partial':'部分交付','no_faculty':'人工确认无师资'}
TEACHER_SHEETS=('有邮箱教师','无邮箱教师','主页失效教师')


def homepage_checks(cat,person):
    """Only classify actual recorded access outcomes; missing CV is not a dead URL."""
    result=[]
    for url in person['review'].get('fields',{}).get('profile_url','').splitlines():
        if not url.strip():continue
        key=canonical(url)
        pages=[p for p in cat['pages'] if canonical(p['url'])==key]
        readable=any(200<=p.get('status',0)<400 and p.get('body_status')!='permission_denied' for p in pages)
        failures=[i for i in cat['issues'] if canonical(i['url'])==key and i['issue_kind']=='request_failed']
        denied=any(p.get('body_status')=='permission_denied' or p.get('status',0)>=400 for p in pages)
        state='available' if readable else ('failed' if failures or denied else 'unknown')
        result.append({'url':url,'state':state,'issue_ids':[i['id'] for i in failures]})
    return result


def teacher_sheet(person):
    states=[x['state'] for m in person.get('members',[person]) for x in m.get('homepage_checks',[])]
    if 'failed' in states and 'available' not in states:return TEACHER_SHEETS[2]
    return TEACHER_SHEETS[0] if person.get('email','').strip() else TEACHER_SHEETS[1]


def collected_date(job):
    # An export/review date must never masquerade as the acquisition date.
    if not job.get('started'):
        return ''
    instant=datetime.fromisoformat(job['started'].replace('Z','+00:00'))
    return instant.astimezone(ZoneInfo('Asia/Singapore')).date().isoformat()


def accepted_records(cat):
    return [{**p['review'].get('fields',{}),'university':cat['job']['target']['university'],
             'college':cat['job']['target']['college'],
             'sources':list(dict.fromkeys([*p['sources'],*(a['roster']['source_url'] for a in p['review'].get('roster_associations',[]))])),
             'roster_associations':p['review'].get('roster_associations',[]),
             'delivery_notes':p['review'].get('delivery_notes',[]),'material_status':p['review'].get('material_status'),
             'material_versions':p['review'].get('material_versions'),
             'homepage_checks':homepage_checks(cat,p),
             'person_id':p['id'],'observed_name':p['name'],'name_normalization':p['review'].get('name_normalization'),'reviewer':p['review']['actor'],
             'review_evidence':p['review']['evidence'],'run_id':cat['job']['id'],
             'acceptance':p['review']['status'],'policy':p['review'].get('policy','human'),
             'field_evidence':p['review'].get('selection',{}),'collector':'自动采集',
             'collected_date':collected_date(cat['job'])}
            for p in cat['people'] if p['review']['status'] in ('accepted','auto_accepted')]


def canonical(url):
    return url.split('#')[0].replace('http://','https://',1).rstrip('/')


def source_links(cat):
    """Only return observed qualified sources, never raw model suggestions/input URLs."""
    rejected={canonical(i['url']) for i in cat['issues'] if i['issue_kind']=='source_not_faculty_roster'}
    accepted_sources={canonical(url) for p in cat['people'] if p['review']['status'] in ('accepted','auto_accepted') for url in p['sources']}
    next_pages={canonical(url) for r in cat['rosters'] for url in r.get('next',[])}
    candidates=[r['url'] for r in cat['rosters']
                if r.get('source_qualification')=='faculty_roster' or canonical(r['url']) in accepted_sources]
    # Includes saved table/document sources that were not registered as a roster page.
    candidates += [u for p in cat['people'] if p['review']['status'] in ('accepted','auto_accepted') for u in p['sources']]
    result=[];seen=set();host=urlparse(cat['job']['target']['url']).hostname
    for url in candidates:
        key=canonical(url);parsed=urlparse(url)
        if (key not in seen and key not in rejected and key not in next_pages
                and parsed.scheme in ('http','https') and parsed.hostname==host):
            result.append(url);seen.add(key)
    return result


def tables(catalogs):
    teachers=[];sources=[];statuses=[]
    merged=consolidate([p for c in catalogs for p in accepted_records(c)])
    merged.sort(key=lambda p:TEACHER_SHEETS.index(teacher_sheet(p)))
    for p in merged:
        teachers.append([len(teachers)+1,p.get('name',''),p.get('title',''),p.get('email',''),
                         p['university'],p['college'],'自动采集',p['collected_date'],p.get('biography',''),p.get('profile_url','')])
    for cat in catalogs:
        target=cat['job']['target']
        links=source_links(cat)
        for url in links or ['']:
            sources.append([target['university'],target['college'],target['url'],url])
        note=(f'已依据本次材料核实 {len(links)} 个入口；不代表覆盖全部师资'
              if links else '尚未核实师资入口；空值不表示无师资')
        if target.get('entries'):
            missing=[u for u in target['entries'] if canonical(u) not in {canonical(x) for x in links}]
            if missing:note+='；给定入口未列入已核实表：'+'；'.join(missing)
        state=cat['college_review']['status']
        statuses.append([target['university'],target['college'],STATES.get(cat['job']['state'],cat['job']['state']),
                         COMPLETENESS.get(state,state),note,cat['job']['id']])
    notes=[]
    missing_labels={'access_failed':'访问或识别失败','not_read':'尚未读取','not_stated_in_source':'本次材料未列出','conflict':'字段值冲突'}
    field_labels={'title':'职称','email':'邮箱','biography':'简介','profile_url':'主页'}
    for person in merged:
        for member in person['members']:
            if not member.get('delivery_notes'):
                continue
            versions=member.get('material_versions') or {}
            selected=next((v for v in versions.get('materials',[]) if v['person_id']==versions.get('selected_id')), {})
            notes.append([person['university'],member['college'],person['name'],
                          '；'.join(n['message'] for n in member['delivery_notes']),
                          '；'.join(field_labels.get(k,k)+'：'+missing_labels.get(v,v) for k,v in member.get('missing',{}).items()),
                          (selected.get('date') or {}).get('date',''),member.get('profile_url',''),
                          '\n'.join(v['url']+' '+((v.get('date') or {}).get('date','日期未核实')) for v in versions.get('materials',[])),
                          '/review/'+member['run_id']+'?item='+member['person_id']])
    return {'teachers':teachers,'sources':sources,'statuses':statuses,'merges':merge_rows(merged),'records':merged,'notes':notes}


def csv_bytes(headers,rows):
    buf=io.StringIO(newline='');writer=csv.writer(buf);writer.writerow(headers)
    for row in rows:
        # CSV has no cell type; protect Excel's formula interpretation.
        writer.writerow(["'"+v if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')) else v for v in row])
    return ('\ufeff'+buf.getvalue()).encode('utf-8')


def workbook_bytes(data,overflow_refs=None):
    wb=Workbook();wb.remove(wb.active)
    groups={name:[] for name in TEACHER_SHEETS}
    for row,person in zip(data['teachers'],data['records'],strict=True):groups[teacher_sheet(person)].append((row,person))
    sheets=[(name,TEACHER_HEADERS,[x[0] for x in groups[name]],[9,14,24,30,25,30,14,16,85,55]) for name in TEACHER_SHEETS]
    for title,headers,rows,widths in sheets:
        sheet=wb.create_sheet(title);sheet.freeze_panes='A2';sheet.sheet_view.showGridLines=False
        for ri,row in enumerate([headers]+rows,1):
            for ci,value in enumerate(row,1):
                reference=None
                if isinstance(value,str) and len(value)>32767:
                    reference=(overflow_refs or {}).get((row[0],ci))
                    if not reference:raise ValueError('存在超过 Excel 单元格上限的简介，请下载 CSV 保留全文')
                    value=f'全文共 {len(value):,} 字符，超过 Excel 单元格上限。点击打开完整文本：{reference}（请先解压整个交付包）'
                if isinstance(value,str):value=ILLEGAL_CHARACTERS_RE.sub('',value)
                cell=sheet.cell(ri,ci,value)
                if reference:
                    cell.hyperlink=reference
                    cell.comment=Comment('完整原文未截断，保存在同目录附件及教师完整 CSV 中。','自动采集')
                if isinstance(value,str):cell.data_type='s' # Source strings are never Excel formulas.
                cell.font=Font(name='Arial',size=11,bold=ri==1,color='FFFFFF' if ri==1 else '203A39')
                cell.fill=PatternFill('solid',fgColor='176D5D' if ri==1 else ('F0F6F2' if ri%2==0 else 'FFFFFF'))
                cell.alignment=Alignment(vertical='top',wrap_text=True)
            sheet.row_dimensions[ri].height=27 if ri==1 else 75
        for ci,width in enumerate(widths,1):sheet.column_dimensions[get_column_letter(ci)].width=width
        sheet.auto_filter.ref=sheet.dimensions
        summary='三个表互斥；访问失败且无已确认可用主页者进入第三表。未提供主页或尚未读取不等于失效。\n'
        summary+='\n'.join(' / '.join(str(v) for v in row[:5]) for row in data['statuses'])
        sheet.cell(1,1).comment=Comment(summary,'自动采集')
        for index,(_,person) in enumerate(groups[title],2):
            notes=list(dict.fromkeys(n['message'] for m in person.get('members',[]) for n in m.get('delivery_notes',[])))
            notes+=['失效链接：'+c['url'] for m in person['members'] for c in m.get('homepage_checks',[]) if c['state']=='failed']
            notes+=['来源：'+u for u in person.get('sources',[])]
            if len(person['members'])>1:notes.append('已合并 '+str(len(person['members']))+' 份材料；学院关联：'+person['college'])
            for field,values in person.get('field_conflicts',{}).items():
                notes.append(('职称' if field=='title' else '邮箱')+'待核：'+'；'.join(v['college']+' = '+v['value'] for v in values))
            if notes:sheet.cell(index,10).comment=Comment('\n'.join(notes)+'\n详细来源、标注与合并依据可在工作台下载对应 CSV。','自动采集')
    buf=io.BytesIO();wb.save(buf);return buf.getvalue()


def university_zip(catalogs):
    """One workbook per university, preserving evidence-based identity merging."""
    groups={}
    for cat in catalogs:groups.setdefault(cat['job']['target']['university'],[]).append(cat)
    buf=io.BytesIO();manifest=[]
    def filename(value):return re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',value).strip(' .')[:70] or '未命名'
    with zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED) as archive:
        for university,items in groups.items():
            folder='大学-'+filename(university)+'-'+hashlib.sha256(university.encode()).hexdigest()[:8]
            data=tables(items);refs={};long_fields=[]
            for row in data['teachers']:
                for col,value in enumerate(row,1):
                    if isinstance(value,str) and len(value)>32767:
                        path=f'全文附件/{row[0]:05d}-{filename(row[1])}-{TEACHER_HEADERS[col-1]}.txt'
                        archive.writestr(folder+'/'+path,value.encode('utf-8-sig'))
                        refs[(row[0],col)]=path
                        long_fields.append({'name':row[1],'field':TEACHER_HEADERS[col-1],'chars':len(value),'path':path})
            archive.writestr(folder+'/教师列表.xlsx',workbook_bytes(data,refs))
            for name,key,headers in [('教师完整','teachers',TEACHER_HEADERS),('师资入口','sources',SOURCE_HEADERS),
                                     ('交付标注','notes',NOTE_HEADERS),('学院状态','statuses',STATUS_HEADERS),('合并说明','merges',MERGE_HEADERS)]:
                archive.writestr(folder+'/'+name+'.csv',csv_bytes(headers,data[key]))
            manifest.append({'university':university,'colleges':len(items),'teachers':len(data['teachers']),
                             'sheets':{label:sum(teacher_sheet(p)==label for p in data['records']) for label in TEACHER_SHEETS},
                             'long_fields':long_fields,'folder':folder})
        archive.writestr('清单.json',json.dumps(manifest,ensure_ascii=False,indent=2).encode('utf8'))
        archive.writestr('请先阅读.txt',('每个大学一个文件夹和一个教师 Excel，内部按有邮箱、无邮箱、主页失效分三表。\n'
            '请先完整解压。超出 Excel 单元格上限的原文保存在“全文附件”，Excel 对应格提供链接，原文没有截断。\n'
            '“教师完整.csv”另保留所有字段全文；不要导入 Excel 后另存为 XLSX 当作完整原件。\n'
            '来源、合并、交付标注及学院完整性分别保留。运行结束不代表学院完整。\n').encode('utf-8-sig'))
    return buf.getvalue()
