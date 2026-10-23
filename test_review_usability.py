import copy,io,json,tempfile,unittest,zipfile
from pathlib import Path
from unittest.mock import patch
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from test_workbench import fixture,accept
from workbench import make_app
from workbench_export import university_zip,TEACHER_SHEETS
from quality_pipeline import build,assessment_fingerprint,saved_evidence_digest

class UsabilityTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.store,self.key=fixture(Path(self.tmp.name))
        self.person=self.store.catalog(self.key)['people'][0]['id']
        self.store.decide(self.key,self.person,accept())
    def tearDown(self):self.tmp.cleanup()
    def test_universities_separate_and_long_text_preserved(self):
        a=self.store.catalog(self.key);b=copy.deepcopy(a);b['job']['target']['university']='第二大学'
        text='完整研究成果。'*5000;b['people'][0]['review']['fields']['biography']=text
        with zipfile.ZipFile(io.BytesIO(university_zip([a,b]))) as z:
            manifest=json.loads(z.read('清单.json'));self.assertEqual(len(manifest),2)
            for entry in manifest:
                book=load_workbook(io.BytesIO(z.read(entry['folder']+'/教师列表.xlsx')))
                self.assertEqual(book.sheetnames,list(TEACHER_SHEETS));self.assertEqual(sum(s.max_row-1 for s in book),1)
                row=next(s for s in book if s.max_row>1)
                self.assertEqual(row['E2'].value,entry['university'])
                if entry['long_fields']:
                    relative=entry['long_fields'][0]['path']
                    self.assertEqual(z.read(entry['folder']+'/'+relative).decode('utf-8-sig'),text)
                    self.assertEqual(row['I2'].hyperlink.target,relative)
                    self.assertIn('超过 Excel',row['I2'].value)
    def test_review_does_not_load_other_colleges(self):
        with patch.object(self.store,'batch_catalogs',side_effect=AssertionError('unrelated batch read')):
            response=TestClient(make_app(self.store)).get('/api/jobs/'+self.key)
        self.assertEqual(response.status_code,200);self.assertEqual(len(response.json()['people']),1)
    def test_related_materials_remain_available_on_demand(self):
        cat=self.store.catalog(self.key)
        group={'college':'学院甲；学院乙','field_conflicts':{},'members':[
            {'run_id':self.key,'person_id':self.person,'college':'学院甲','biography':'不应出现在关联响应的长正文'},
            {'run_id':'other','person_id':'other','college':'学院乙'}]}
        with patch.object(self.store,'batch_catalogs',return_value=({},[cat])),patch('workbench.tables',return_value={'records':[group]}):
            response=TestClient(make_app(self.store)).get('/api/jobs/'+self.key+'/related')
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json()[0]['members'][0]['person_id'],self.person)
        self.assertNotIn('biography',response.json()[0]['members'][0])
    def test_evidence_does_not_resolve_unrelated_registered_paths(self):
        root=Path(self.store.job(self.key)['run_dir'])
        pages=json.loads((root/'pages.json').read_text(encoding='utf8'))
        pages.append({'url':'https://test.edu.cn/unrelated','directory':'unrelated','kind':'profile'})
        (root/'pages.json').write_text(json.dumps(pages),encoding='utf8')
        original=Path.resolve
        def resolve(path,*args,**kwargs):
            if 'unrelated' in path.parts:raise OSError('unrelated file unavailable')
            return original(path,*args,**kwargs)
        with patch.object(Path,'resolve',resolve):
            client=TestClient(make_app(self.store))
            response=client.get('/api/jobs/'+self.key+'/evidence',params={'path':'page/body-text.txt','text':'true'})
        self.assertEqual(response.status_code,200);self.assertIn('张三',response.text)
        (root/'private.txt').write_text('not registered',encoding='utf8')
        self.assertEqual(client.get('/api/jobs/'+self.key+'/evidence',params={'path':'private.txt'}).status_code,403)
    def test_persisted_digest_skips_unchanged_bytes_but_detects_changes(self):
        import hashlib
        path=Path(self.tmp.name)/'original.html';path.write_text('原文',encoding='utf8');cache={}
        expected,changed=saved_evidence_digest(path,cache,'page');self.assertTrue(changed)
        persisted=json.loads(json.dumps(cache))
        with patch('quality_pipeline.content_digest',side_effect=AssertionError('unchanged original reread')):
            value,changed=saved_evidence_digest(path,persisted,'page')
        self.assertEqual(value,expected);self.assertFalse(changed)
        path.write_text('修改后的原文',encoding='utf8')
        value,changed=saved_evidence_digest(path,persisted,'page')
        self.assertTrue(changed);self.assertEqual(value,hashlib.sha256(path.read_bytes()).digest());self.assertNotEqual(value,expected)
        path.unlink();self.assertEqual(saved_evidence_digest(path,persisted,'page')[0],b'missing')
    def test_monitor_summary_respects_new_manual_decision_and_source_change(self):
        build(self.store,self.key)
        # Valid snapshots avoid building/reading complete biographies.
        with patch.object(self.store,'catalog',side_effect=AssertionError('full catalog read')),patch('quality_pipeline.assessment_fingerprint',side_effect=AssertionError('overview must not scan originals')):
            self.assertEqual(self.store.listing()['jobs'][0]['accepted'],1)
        self.store.decide(self.key,self.person,{**accept(1),'status':'needs_followup'})
        self.assertEqual(self.store.listing()['jobs'][0]['accepted'],0)
        self.assertGreater(self.store.listing()['jobs'][0]['pending'],0)
        root=Path(self.store.job(self.key)['run_dir']);before=assessment_fingerprint(root)
        (root/'page/source.html').write_text('<p>原件发生变化</p>',encoding='utf8')
        self.assertNotEqual(before,assessment_fingerprint(root))
        with patch.object(self.store,'catalog',side_effect=AssertionError('full catalog read')):
            job=self.store.listing()['jobs'][0]
        self.assertEqual(job['statistics_source'],'last_check')
        self.assertEqual(job['accepted'],0)  # New manual decisions still override old counts.
        # Overview counts must not make stale automatic decisions deliverable.
        self.assertFalse(self.store.catalog(self.key)['automation'])
        response=TestClient(make_app(self.store)).get('/api/jobs/'+self.key+'/export')
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json()['accepted'],[])
        build(self.store,self.key)
        self.assertEqual(self.store.listing()['jobs'][0]['statistics_source'],'last_check')

    def test_monitor_without_summary_or_during_collection_reads_live_catalog(self):
        with patch.object(self.store,'catalog',wraps=self.store.catalog) as call:
            self.store.listing();self.assertTrue(call.called)
        build(self.store,self.key)
        self.store.update(self.key,state='running')
        with patch.object(self.store,'catalog',wraps=self.store.catalog) as call:
            job=self.store.listing()['jobs'][0];self.assertTrue(call.called)
        self.assertEqual(job['statistics_source'],'live')

if __name__=='__main__':unittest.main()
