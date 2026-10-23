import json
import tempfile
import unittest
from pathlib import Path

from quality_pipeline import build
from material_delivery import page_date
from test_workbench import fixture
from workbench_export import accepted_records, tables, workbook_bytes


class DeliveryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.store, self.key = fixture(Path(self.temp.name))
        self.root = Path(self.store.job(self.key)['run_dir'])
        self.obs = json.loads((self.root/'observations.json').read_text())
        self.pages = json.loads((self.root/'pages.json').read_text())
        self.issues = []
        self.rosters = [{'url':'https://test.edu.cn/list','branch':'师资队伍','source_qualification':'faculty_roster'}]
        (self.root/'list').mkdir()
        self.pages.append({'url':'https://test.edu.cn/list','title':'师资队伍','directory':'list','kind':'roster'})
        self.body('page','张三，副教授，博士。长期从事材料科学的教学与研究工作。邮箱 zhang@test.edu.cn', '2019-06-06')

    def tearDown(self):
        self.temp.cleanup()

    def body(self, folder, text, day):
        path = self.root/folder; path.mkdir(exist_ok=True)
        (path/'body-text.txt').write_text(text, encoding='utf8')
        (path/'source.html').write_text('<h1>张三</h1>'+('<span class="arti_update">发布时间：'+day+'</span>' if day else '')+'<div class="wp_articlecontent">'+text+'</div>', encoding='utf8')
        for p in self.pages:
            if p['directory']==folder:
                p.update(title='张三',body_chars=len(text),retrieved_at='2026-10-14T10:00:00+08:00')

    def second(self, day='2025-03-28', text=None, url='https://test.edu.cn/zhang2'):
        self.obs.append({**self.obs[0],'profile_url':url})
        self.pages.append({'url':url,'directory':'second','kind':'profile'})
        self.body('second', text if text is not None else '张三，教授，博士。最新研究方向为智能制造技术，主讲先进制造课程。邮箱 zhang@test.edu.cn', day)

    def assess(self):
        anchors=''.join('<p><a href="'+o['profile_url']+'">'+o['name']+'</a>：教授</p>' for o in self.obs if o.get('profile_url'))
        (self.root/'list/source.html').write_text('<title>师资队伍</title><div class="wp_articlecontent">'+anchors+'</div>',encoding='utf8')
        for name,data in [('observations',self.obs),('pages',self.pages),('roster-pages',self.rosters),('issues',self.issues)]:
            (self.root/(name+'.json')).write_text(json.dumps(data),encoding='utf8')
        build(self.store,self.key,force=True)
        return self.store.catalog(self.key)

    def test_failed_homepage_delivers_roster_fields_with_explicit_annotation(self):
        self.body('page','',None)
        self.issues=[{'kind':'request_failed','url':'https://test.edu.cn/zhang','error_type':'ConnectError'}]
        cat=self.assess();person=cat['people'][0]
        self.assertEqual(person['review']['status'],'auto_accepted')
        self.assertEqual(person['review']['material_status'],'profile_access_failed')
        self.assertEqual(person['review']['fields']['title'],'教授')
        self.assertEqual(person['review']['fields']['biography'],'')
        self.assertEqual(person['review']['fields']['missing']['biography'],'access_failed')
        self.assertEqual(cat['issues'][0]['review']['status'],'deferred')
        self.assertEqual(len(accepted_records(cat)),1)
        from openpyxl import load_workbook
        from fastapi.testclient import TestClient
        from workbench import make_app
        import io
        book=load_workbook(io.BytesIO(workbook_bytes(tables([cat]))))
        self.assertEqual(book['主页失效教师'].max_column,10)
        self.assertIn('主页访问失败',book['主页失效教师'].cell(2,10).comment.text)
        with TestClient(make_app(self.store)) as client:
            result=client.get('/api/jobs/'+self.key+'/export?format=csv&table=notes')
            self.assertEqual(result.status_code,200)
        self.assertIn('主页访问失败',result.text)

    def test_unvisited_roster_link_is_not_a_dead_homepage(self):
        self.body('page','',None)
        self.issues=[{'kind':'out_of_host_scope','url':'https://test.edu.cn/zhang'}]
        person=self.assess()['people'][0]
        self.assertEqual(person['review']['material_status'],'profile_not_visited')
        self.assertEqual(person['review']['fields']['missing']['biography'],'not_read')
        self.assertEqual(person['review']['status'],'auto_accepted')

    def test_pending_queue_and_latin_name_keep_identity_without_fake_failure(self):
        self.body('page','',None)
        self.obs[0]['name']='Simone Baldi'
        (self.root/'summary.json').write_text(json.dumps({'pending_urls':['https://test.edu.cn/zhang']}),encoding='utf8')
        person=self.assess()['people'][0]
        self.assertEqual(person['review']['material_status'],'profile_queued')
        self.assertEqual(person['review']['fields']['missing']['biography'],'not_read')
        self.assertEqual(person['review']['status'],'auto_accepted')

    def test_department_breadcrumb_confirms_legacy_faculty_source(self):
        from roster_identity import RosterEvidence
        self.assess();self.rosters[0].pop('source_qualification')
        self.rosters[0]['branch']='流行病学'
        self.pages[-1]['title']='流行病学'
        (self.root/'list/source.html').write_text('<title>流行病学</title><li class="col_path">首页 &gt; 导师风采 &gt; 流行病学</li><a href="/zhang">张三（邮箱：zhang@test.edu.cn）</a>',encoding='utf8')
        cat=self.store.catalog(self.key,apply_auto=False);cat['rosters']=self.rosters;cat['pages']=self.pages
        self.assertIsNotNone(RosterEvidence(cat,self.root).evidence(cat['people'][0],True))
        (self.root/'list/source.html').write_text('<title>流行病学</title><nav>导师风采</nav><a href="/zhang">张三</a>',encoding='utf8')
        self.assertIsNone(RosterEvidence(cat,self.root).evidence(cat['people'][0],True))

    def test_unknown_or_old_source_does_not_allow_empty_record(self):
        self.body('page','',None);self.issues=[{'kind':'request_failed','url':'https://test.edu.cn/zhang'}]
        self.rosters[0]['source_qualification']='unknown'
        self.assertEqual(self.assess()['people'][0]['review']['status'],'pending')
        self.rosters[0]['source_qualification']='faculty_roster';self.obs[0]['historical_roster']=True
        self.assertEqual(self.assess()['people'][0]['review']['status'],'pending')

    def test_failed_attachment_does_not_drop_known_biography(self):
        self.issues=[{'kind':'image_ocr_failed','url':'https://test.edu.cn/a.jpg','source_page':'https://test.edu.cn/zhang'}]
        p=self.assess()['people'][0]
        self.assertEqual(p['review']['status'],'auto_accepted')
        self.assertEqual(p['review']['material_status'],'attachment_access_failed')
        self.assertIn('长期从事',p['review']['fields']['biography'])

    def test_verified_identity_uses_latest_explicit_page_date(self):
        self.second();cat=self.assess();accepted=[p for p in cat['people'] if p['review']['status']=='auto_accepted']
        self.assertEqual(len(accepted),1)
        self.assertEqual(accepted[0]['profile_url'],'https://test.edu.cn/zhang2')
        self.assertEqual(accepted[0]['review']['fields']['title'],'教授')
        self.assertEqual(accepted[0]['review']['material_versions']['rule'],'latest_dated_material')
        self.assertEqual(len(accepted[0]['review']['material_versions']['materials']),2)
        self.assertIn('副教授',(self.root/'page/body-text.txt').read_text(encoding='utf8'))

    def test_same_name_with_different_identity_is_not_resolved_by_date(self):
        self.second(text='张三，讲师。研究方向为病理学与公共卫生。邮箱 other@test.edu.cn')
        self.assertTrue(all(p['review']['status']=='pending' for p in self.assess()['people']))

    def test_prefixed_roster_labels_preserve_identity_and_latest_material(self):
        self.second();self.assess()
        path=self.root/'list/source.html'
        path.write_text(path.read_text('utf8').replace('>张三</a>','>导师简介-张三</a>'),encoding='utf8')
        build(self.store,self.key,force=True)
        people=self.store.catalog(self.key)['people']
        accepted=[p for p in people if p['review']['status']=='auto_accepted']
        self.assertEqual(len(accepted),1)
        self.assertEqual(accepted[0]['profile_url'],'https://test.edu.cn/zhang2')
        self.assertEqual(sum(p['review']['status']=='duplicate' for p in people),1)

    def test_shared_mail_is_not_identity_evidence(self):
        self.second();self.obs.append({**self.obs[0],'name':'李四','profile_url':'https://test.edu.cn/li'})
        self.pages.append({'url':'https://test.edu.cn/li','directory':'third','title':'李四','body_chars':50})
        (self.root/'third').mkdir();(self.root/'third/body-text.txt').write_text('李四，教授，博士。长期开展数学研究和相关教学工作。邮箱 zhang@test.edu.cn',encoding='utf8')
        self.assertTrue(all(p['review']['status']=='pending' for p in self.assess()['people'] if p['name']=='张三'))

    def test_missing_tied_and_future_dates_do_not_pick_newest(self):
        self.second(day=None)
        for day in (None,'2019-06-06','2099-01-01'):
            self.body('second','张三，教授。最新材料包括学术与教学经历。邮箱 zhang@test.edu.cn',day)
            self.assertTrue(all(p['review']['status']=='pending' for p in self.assess()['people']))

    def test_publication_date_excludes_cv_dates_and_url_dates(self):
        (self.root/'page/source.html').write_text('<footer>2026-10-14</footer><div class="wp_articlecontent"><time datetime="2026-10-13">工作经历</time></div>',encoding='utf8')
        self.assertIsNone(page_date(self.root,self.pages[0]))

    def test_same_cms_article_aliases_do_not_create_two_failed_people(self):
        first='https://test.edu.cn/2019/0514/c1a123/page.htm';second='https://test.edu.cn/_s60/2019/0514/c1a123/page.psp'
        self.obs[0]['profile_url']=first;self.pages[0]['url']=first;self.body('page','',None)
        self.second(text='',day=None,url=second)
        self.issues=[{'kind':'request_failed','url':u} for u in (first,second)]
        people=self.assess()['people'];self.assertEqual(sum(p['review']['status']=='auto_accepted' for p in people),1)
        self.assertEqual(sum(p['review']['status']=='duplicate' for p in people),1)
        self.assertTrue(all(i['review']['status']=='deferred' for i in self.store.catalog(self.key)['issues']))

    def test_manual_followup_is_not_replaced(self):
        self.second();cat=self.assess();p=cat['people'][0]
        self.store.decide(self.key,p['id'],{'revision':0,'actor':'测试员','status':'needs_followup','reason':'有不同人线索','evidence':'官方材料'})
        p=next(x for x in self.assess()['people'] if x['id']==p['id'])
        self.assertEqual(p['review']['status'],'needs_followup')


if __name__=='__main__':
    unittest.main()
