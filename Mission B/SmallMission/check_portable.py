"""Check the unzipped Windows app with no Python or browser on PATH."""
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import io
import base64
from datetime import datetime
import json
import os
from pathlib import Path
import subprocess
import threading
import time
from urllib.parse import urlsplit, parse_qs
import zipfile

import httpx
from openpyxl import load_workbook
from runtime_paths import VERSION
from test_multisource import paired_file

ROOT=Path(__file__).resolve().parent


class Source(BaseHTTPRequestHandler):
    def log_message(self,*args):pass
    def do_GET(self):
        self.server.requests.append(self.path)
        browser='HeadlessChrome/' in self.headers.get('User-Agent','')
        person = 'b' if '?page=2' in self.path else 'a'
        html = ('<title>打包验收教师名单</title><div class="roster"><a href="https://target.invalid/' + person + '">姓名</a></div>'
                '<div class="pb_sys_common"><a href="/roster">首页</a><a href="/roster?page=2">尾页</a></div>') if browser else 'HTTP fallback required'
        raw=html.encode('utf-8')
        self.send_response(200 if browser else 503)
        self.send_header('Content-Type','text/html;charset=utf-8');self.send_header('Content-Length',str(len(raw)));self.end_headers();self.wfile.write(raw)


def main():
    test=ROOT/'outputs'/('portable validation 中文 '+datetime.now().strftime('%Y%m%d-%H%M%S'))
    test.mkdir(parents=True,exist_ok=True)
    archive=ROOT/'release'/f'MissionB-Windows-x64-v{VERSION}.zip'
    with zipfile.ZipFile(archive) as z:z.extractall(test)
    app=test/f'MissionB-Windows-x64-v{VERSION}'
    config=json.loads((app/'source_plans.json').read_text(encoding='utf-8'))
    source=ThreadingHTTPServer(('127.0.0.1',0),Source);source.requests=[]
    thread=threading.Thread(target=source.serve_forever,daemon=True);thread.start()
    roster=f'http://127.0.0.1:{source.server_port}/roster'
    config['plans'][roster]={'selector':'.roster','title_contains':'打包验收教师名单'}
    (app/'source_plans.json').write_text(json.dumps(config,ensure_ascii=False),encoding='utf-8')
    env={k:v for k,v in os.environ.items() if k.upper() not in ('PYTHONHOME','PYTHONPATH','VIRTUAL_ENV','PLAYWRIGHT_BROWSERS_PATH')}
    env['PATH']=str(Path(os.environ['SystemRoot'])/'System32')
    process=subprocess.Popen([str(app/'MissionB.exe'),'--no-browser'],cwd=app,env=env,creationflags=subprocess.CREATE_NO_WINDOW)
    try:
        deadline=time.monotonic()+45
        while not (app/'data'/'instance.json').exists():
            assert process.poll() is None,'Packaged process exited before startup'
            assert time.monotonic()<deadline,'Startup timeout'
            time.sleep(.2)
        info=json.loads((app/'data'/'instance.json').read_text(encoding='utf-8'))
        token=parse_qs(urlsplit(info['url']).fragment)['token'][0]
        with httpx.Client(base_url=f'http://127.0.0.1:{info["port"]}',trust_env=False,timeout=10) as c:
            assert c.post('/api/session',json={'token':token}).status_code==200
            assert c.get('/').status_code==200
            assert c.get('/app.js').status_code==200
            assert c.get('/api/template.xlsx').status_code==200
            upload=test/'双表输入.xlsx'
            paired_file(upload,'测试大学','测试学院',[roster],['https://target.invalid/a'])
            p=c.post('/api/preview',json={'files':[{'name':upload.name,'content':base64.b64encode(upload.read_bytes()).decode()}]}).json()
            assert not p['errors'],p
            response=c.post('/api/jobs',json={'preview_id':p['id']});assert response.status_code==201,response.text
            identifier=response.json()['id'];deadline=time.monotonic()+100
            while True:
                j=c.get('/api/jobs/'+identifier).json()
                if j['status'] not in ('running','stopping'):break
                assert time.monotonic()<deadline,'Task timeout'
                time.sleep(.5)
            assert j['status']=='completed',j
            assert j['results'][0]['complete'],j
            assert j['results'][0]['stats']['missing']==1,j
            output=app/'data'/'runs'/identifier/'output'
            result=json.loads((output/'results.json').read_text(encoding='utf-8'))[0]
            assert len(result['attempts'])==3 and result['attempts'][1]['status']==200 and result['attempts'][2]['status']==200,result['attempts']
            assert result['sources'][1]['discovered_from']==roster,result['sources']
            assert all(p in ('/roster','/roster?page=2') for p in source.requests),source.requests
            book=load_workbook(io.BytesIO(c.get('/api/jobs/'+identifier+'/download/excel').content),read_only=True)
            assert book['修正后四列表'].max_row==3;book.close()
            summary=load_workbook(io.BytesIO(c.get('/api/jobs/'+identifier+'/download/summary').content),read_only=True,data_only=True)
            assert len(summary.sheetnames)==1 and list(summary.active.values)==[('大学名','学院名','是否出错'),('测试大学','测试学院','Yes')]
            summary.close()
            # A second launch must return without starting another server or changing history.
            duplicate=subprocess.Popen([str(app/'MissionB.exe'),'--no-browser'],cwd=app,env=env,creationflags=subprocess.CREATE_NO_WINDOW)
            assert duplicate.wait(timeout=15)==0
            assert c.get('/api/jobs').json()[0]['id']==identifier
            assert c.post('/api/shutdown',json={}).status_code==200
        assert process.wait(timeout=15)==0
        manifest=json.loads((output/'manifest.json').read_text(encoding='utf-8'))
        report={'status':'passed','packaged_executable':True,'no_python_or_browser_on_PATH':True,'path_contains_spaces_and_Chinese':True,'browser_fallback':result['attempts'][1]['status'],'explicit_source_requests':source.requests,'model_api_calls':manifest['model_api_calls'],'profile_fetches':manifest['explicit_profile_fetches'],'duplicate_launch_and_exit':'passed','dependencies':manifest['dependencies']}
        (test/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
        print(json.dumps(report,ensure_ascii=False))
    finally:
        if process.poll() is None:
            process.terminate();process.wait(timeout=10)
        source.shutdown();thread.join();source.server_close()


if __name__=='__main__':main()
