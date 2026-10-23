"""Input interoperability and local app integration checks, using isolated data."""
import asyncio
import base64
import json
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import patch

import httpx
from openpyxl import Workbook, load_workbook

from app import make_server, Store
from audit import Fetcher, parse_roster, reconcile
from input_adapter import parse_file, tasks_from_rows
from reporting import export
from reporting import input_assessment, export_names
from runtime_paths import APP_ROOT
from test_audit import SOURCE, PLAN, task

HEAD = '学校\t学院\t师资url\t教师url\n'
ROW = '测试大学\t测试学院\t' + SOURCE + '\thttps://example.edu.cn/a'


class InputTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=APP_ROOT)
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def parse(self, text, name='input.txt', encoding='utf-8'):
        path = self.root / name
        path.write_bytes(text.encode(encoding))
        return parse_file(path)

    def test_paste_tsv_and_blank_url(self):
        rows, errors, _ = self.parse(HEAD + ROW + '\n测试大学\t测试学院\t' + SOURCE + '\t\n')
        self.assertFalse(errors)
        self.assertEqual(len(rows), 2)
        self.assertIsNone(rows[1]['url'])

    def test_csv_quoted_commas_and_utf16(self):
        for encoding in ('utf-8-sig', 'utf-16', 'gb18030'):
            rows, errors, _ = self.parse('school,college,roster_url,teacher_url\n"Test, University",Test College,' + SOURCE + ',https://example.edu.cn/a', 'data.csv', encoding)
            self.assertFalse(errors)
            self.assertEqual(rows[0]['school'], 'Test, University')

    def test_markdown_links_and_separator(self):
        rows, errors, _ = self.parse('| 大学 | 学院名 | 师资URL | 教师个人URL |\n| --- | --- | --- | --- |\n| 测试大学 | 测试学院 | [师资](' + SOURCE + ') | <https://example.edu.cn/a> |')
        self.assertFalse(errors)
        self.assertEqual(rows[0]['url'], 'https://example.edu.cn/a')

    def test_space_separated_no_header(self):
        rows, errors, _ = self.parse(ROW.replace('\t', '  '))
        self.assertFalse(errors)
        self.assertEqual(rows[0]['college'], '测试学院')

    def test_reordered_headers_and_extra_columns(self):
        rows, errors, _ = self.parse('备注,教师url,学院,学校,师资url\n备注,https://example.edu.cn/a,学院,大学,' + SOURCE, 'input.csv')
        self.assertFalse(errors)
        self.assertEqual(rows[0]['url'], 'https://example.edu.cn/a')

    def test_extra_columns_after_standard_headers(self):
        rows, errors, _ = self.parse(HEAD.rstrip() + '\t备注\n' + ROW + '\t忽略')
        self.assertFalse(errors)
        self.assertEqual(len(rows), 1)

    def test_excel_merged_cells_multiple_sheets_and_hyperlink_labels(self):
        book = Workbook()
        for ws in (book.active, book.create_sheet('第二页')):
            ws.append(HEAD.strip().split('\t'))
            ws.append(['测试大学', '测试学院', SOURCE, '张老师'])
            ws['D2'].hyperlink = 'https://example.edu.cn/a'
            ws.append([None, None, None, 'https://example.edu.cn/b'])
            for col in 'ABC':
                ws.merge_cells(f'{col}2:{col}3')
        path = self.root / 'merged.xlsx'; book.save(path)
        rows, errors, _ = parse_file(path)
        self.assertFalse(errors)
        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[0]['url'], 'https://example.edu.cn/a')
        self.assertEqual(rows[1]['school'], '测试大学')
        self.assertEqual(len(tasks_from_rows(rows)), 1)

    def test_static_hyperlink_formula_supported_other_formula_rejected(self):
        for value, good in [('=HYPERLINK("https://example.edu.cn/a","张老师")', True), ('=CONCAT("https://example.edu.cn/",A2)', False)]:
            book = Workbook(); book.active.append(HEAD.strip().split('\t'))
            book.active.append(['学校名', '学院名', SOURCE, value])
            path = self.root / 'formula.xlsx'; book.save(path)
            rows, errors, _ = parse_file(path)
            self.assertEqual(not errors, good)
            if good:
                self.assertEqual(rows[0]['url'], 'https://example.edu.cn/a')

    def test_conflicting_hyperlink_is_blocked_before_running(self):
        book = Workbook(); ws = book.active; ws.append(HEAD.strip().split('\t'))
        ws.append(['大学名','学院名',SOURCE,'https://example.edu.cn/a']); ws['D2'].hyperlink='https://example.edu.cn/b'
        path = self.root / 'conflict.xlsx'; book.save(path)
        rows, errors, _ = parse_file(path)
        self.assertFalse(rows)
        self.assertIn('超链接目标不同', errors[0])

    def test_bad_row_reports_physical_line(self):
        rows, errors, _ = self.parse(HEAD + ROW + '\n\n只有一列')
        self.assertEqual(len(rows), 1)
        self.assertIn('第4行', errors[0])

    def test_broken_workbook_and_duplicate_headers(self):
        _, errors, _ = self.parse('not a workbook', 'broken.xlsx')
        self.assertTrue(errors)
        _, errors, _ = self.parse('学校\t学院\t师资url\t教师url\t教师个人url\n' + ROW)
        self.assertTrue(errors)

    def test_empty_and_invalid_roster(self):
        for text in ('', HEAD, HEAD + ROW.replace(SOURCE, 'www.example.edu.cn')):
            _, errors, _ = self.parse(text)
            self.assertTrue(errors)

    def test_legacy_xls_fixture(self):
        path = APP_ROOT / 'tests' / 'fixtures' / 'plain.xls'
        rows, errors, _ = parse_file(path)
        self.assertFalse(errors)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]['url'], 'https://example.edu.cn/a')

    def test_same_college_different_sources_does_not_overwrite_export(self):
        results = []
        for i in range(2):
            t = task('https://example.edu.cn/' + str(i)); t['id'] = 'test' + str(i)
            t['roster_url'] = SOURCE + '?group=' + str(i)
            parsed = parse_roster(f'<title>教师</title><div class="roster"><a href="/{i}">甲</a></div>', SOURCE, PLAN)
            results.append({**t, **reconcile(t, parsed), 'source': parsed, 'attempts': [{'method':'测试','captured_at':'测试'}]})
        manifest = {'scope':'测试','configuration_note':'测试','mode':'测试','model_api_calls':0,'explicit_profile_fetches':0,'finished_at':'测试','needs_human':0}
        export(self.root, results, manifest)
        self.assertEqual(len(list((self.root/'corrected').glob('*.xlsx'))), 2)

    def test_original_input_difference_labels_do_not_hide_repairs(self):
        def assessed(**stats):
            return input_assessment({'complete':True,'stats':stats})
        self.assertEqual(assessed(missing=3,extra=2)['title'],'缺少 3 条，多出 2 条')
        self.assertEqual(assessed(missing=3,extra=0)['title'],'缺少 3 条')
        self.assertEqual(assessed(missing=0,extra=2)['title'],'多出 2 条')
        self.assertEqual(assessed(missing=0,extra=0)['title'],'无链接差异')
        for key in ('input_duplicates','protocol_changes','input_placeholders'):
            self.assertEqual(assessed(missing=0,extra=0,**{key:1})['tone'],'changed')
        self.assertEqual(input_assessment({'complete':False,'stats':{}})['title'],'尚未判定')

    def test_export_names_include_college_batch_and_unique_run(self):
        results=[{'id':'first','school':'甲大学','college':'乙学院'}, {'id':'second','school':'甲大学','college':'丙学院'}]
        meta={'finished_at':'2026-10-15T10:00:00+08:00','run_id':'123456abcdef'}
        single,names=export_names(results[:1],meta)
        self.assertTrue(single.startswith('甲大学_乙学院_核验结果_'))
        self.assertTrue(names['first'].startswith('甲大学_乙学院_修正后链接_'))
        self.assertNotEqual(single,export_names(results[:1],{**meta,'run_id':'abcdef123456'})[0])
        self.assertIn('等2学院',export_names(results,meta)[0])


class AppTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=APP_ROOT)
        self.root = Path(self.temp.name)
        self.server = make_server(self.root/'data')
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.client = httpx.Client(base_url=f'http://127.0.0.1:{self.server.server_port}', trust_env=False)

    def tearDown(self):
        worker = self.server.store.worker
        if worker:
            worker.join(timeout=10)
        self.client.close(); self.server.shutdown(); self.thread.join(); self.server.server_close(); self.temp.cleanup()

    def login(self):
        response = self.client.post('/api/session', json={'token': self.server.token})
        self.assertEqual(response.status_code, 200)

    def test_session_and_cross_origin_protection(self):
        self.assertEqual(self.client.get('/api/jobs').status_code,401)
        self.assertEqual(self.client.post('/api/session',json={'token':'wrong'}).status_code,403)
        self.login()
        self.assertEqual(self.client.get('/api/jobs').status_code,200)
        self.assertEqual(self.client.post('/api/preview',json={'text':HEAD+ROW},headers={'Origin':'https://outside.invalid'}).status_code,403)
        self.assertEqual(self.client.get('/api/jobs',headers={'Host':'outside.invalid'}).status_code,403)

    def test_upload_paths_and_stale_preview_blocked(self):
        self.login()
        response = self.client.post('/api/preview',json={'files':[{'name':'../outside.csv','content':base64.b64encode(b'x').decode()}]})
        self.assertEqual(response.status_code,400)
        p = self.client.post('/api/preview',json={'text':HEAD+ROW}).json()
        (self.root/'data'/'previews'/p['id']/'inputs'/'粘贴输入.txt').write_text(HEAD+ROW+'2',encoding='utf-8')
        self.assertEqual(self.client.post('/api/jobs',json={'preview_id':p['id']}).status_code,400)

    def test_full_preview_job_export_history_flow(self):
        self.login()
        config = self.root / 'plans.json'; config.write_text(json.dumps({'scope':'测试','plans':{SOURCE:PLAN}}),encoding='utf-8')
        calls = []
        async def fake_http(fetcher, task):
            calls.append(task['roster_url'])
            html='<title>教师</title><div class="roster"><a href="/a">甲</a><a href="/b">乙</a></div>'
            meta={'requested_url':SOURCE,'final_url':SOURCE,'status':200,'captured_at':'测试','method':'模拟来源'}
            fetcher.save(task['id'],'http',html,meta)
            return html, meta
        with patch('app.plans_path',return_value=config), patch.object(Fetcher,'http',fake_http):
            p=self.client.post('/api/preview',json={'text':HEAD+ROW}).json()
            self.assertFalse(p['errors']);self.assertTrue(p['groups'][0]['supported'])
            response=self.client.post('/api/jobs',json={'preview_id':p['id']});self.assertEqual(response.status_code,201)
            identifier=response.json()['id']
            self.server.store.worker.join(timeout=10)
        j=self.client.get('/api/jobs/'+identifier).json()
        self.assertEqual(j['status'],'completed',j)
        self.assertEqual(j['results'][0]['stats']['missing'],1)
        self.assertEqual(calls,[SOURCE])
        detail=self.client.get('/api/jobs/'+identifier+'/college/'+j['results'][0]['id']).json()
        self.assertEqual(detail['groups'][0]['occurrences'],2)
        self.assertEqual(self.client.get('/api/jobs/'+identifier+'/download/excel').status_code,200)
        self.assertEqual(self.client.get('/api/jobs/'+identifier+'/download/all').status_code,200)
        store=Store(self.root/'data');self.assertEqual(store.get(identifier)['status'],'completed')

    def test_unknown_source_stays_manual_and_keeps_input(self):
        self.login()
        async def fake_http(fetcher, task):
            return '<title>名单</title><a href="/someone">老师</a>',{'status':200,'final_url':SOURCE,'captured_at':'test','method':'测试'}
        with patch.object(Fetcher,'http',fake_http),patch.object(Fetcher,'browser_fetch',fake_http):
            p=self.client.post('/api/preview',json={'text':HEAD+ROW}).json()
            self.assertFalse(p['groups'][0]['supported'])
            j=self.client.post('/api/jobs',json={'preview_id':p['id']}).json()
            self.server.store.worker.join(timeout=10)
        j=self.client.get('/api/jobs/'+j['id']).json()
        self.assertEqual(j['status'],'completed')
        self.assertTrue(j['results'][0]['needs_human'])
        self.assertEqual(list((self.root/'data'/'runs'/j['id']/'output'/'corrected').glob('*.xlsx')),[])


if __name__=='__main__':
    unittest.main(verbosity=2)
