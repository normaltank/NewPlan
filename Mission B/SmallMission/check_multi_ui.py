"""Browser acceptance of dual-sheet uploads and the single batch summary."""
import asyncio
from datetime import datetime
import json
import os
import threading
from unittest.mock import patch
from openpyxl import load_workbook
from playwright.async_api import async_playwright, expect
from app import make_server
from audit import Fetcher
from runtime_paths import APP_ROOT, VERSION
from test_multisource import paired_file


async def main():
    out = APP_ROOT / 'outputs' / ('ui-multisource-' + datetime.now().strftime('%Y%m%d-%H%M%S'))
    out.mkdir(parents=True)
    files, plans = [], {}
    for index in range(4):
        sources = [f'https://example.edu.cn/{index}/{page}' for page in (1, 2)]
        teachers = ['https://example.edu.cn/a', 'https://example.edu.cn/x' if index == 1 else 'https://example.edu.cn/b']
        path = out / f'测试大学_学院{index}.xlsx'
        paired_file(path, '测试大学', f'学院{index}', sources, teachers)
        files.append(path)
        plans.update({u: {'selector': '.roster', 'title_contains': '教师'} for u in sources})
    plans['https://example.edu.cn/3/2'] = {'selector': None}
    config = out / 'plans.json'
    config.write_text(json.dumps({'plans': plans, 'scope': '测试夹具，仅验证功能'}), encoding='utf-8')
    async def fake(fetcher, task):
        url = task['roster_url']
        html = '<title>教师</title><div class="roster"><a href="/' + ('a' if url.endswith('/1') else 'b') + '">姓名</a></div>'
        meta = {'requested_url': url, 'final_url': url, 'status': 503 if url.endswith('/2/2') else 200,
                'method': '本地测试夹具', 'captured_at': 'test'}
        fetcher.save(task['id'], 'http', html, meta)
        return html, meta
    server = make_server(out / 'data')
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    os.environ['PLAYWRIGHT_BROWSERS_PATH'] = str(APP_ROOT / '.portable-browser')
    errors = []
    try:
        with patch('app.plans_path', return_value=config), patch.object(Fetcher, 'http', fake), patch.object(Fetcher, 'browser_fetch', fake):
            async with async_playwright() as pw:
                browser = await pw.chromium.launch(headless=True)
                page = await browser.new_page(viewport={'width': 1440, 'height': 1080})
                page.on('pageerror', lambda e: errors.append(str(e)))
                await page.goto(f'http://127.0.0.1:{server.server_port}/#token={server.token}')
                await expect(page.locator('#version')).to_have_text('VERSION ' + VERSION)
                # Actual supplied inputs are previewed only; their pages are not mocked as successful.
                real_files = [p for p in (APP_ROOT.parent / '学校_学院_2').glob('*.xlsx') if not p.name.startswith('~$')]
                await page.locator('#files').set_input_files(real_files)
                await page.locator('#preview-button').click()
                await page.locator('#preview-groups table').wait_for()
                text = await page.locator('#preview-stats').inner_text()
                assert '983' in text and '37' in text and '10' in text, text
                assert not await page.locator('#start-button').is_disabled()
                await page.screenshot(path=str(out / '01-real-input-preview.png'), full_page=True)
                while await page.locator('#file-list button').count():
                    await page.locator('#file-list button').first.click()
                await page.locator('#files').set_input_files(files)
                await page.locator('#preview-button').click()
                await page.locator('#preview-groups table').wait_for()
                assert await page.locator('#preview-groups tbody tr').count() == 4
                await page.locator('#start-button').click()
                await page.locator('#downloads:not([hidden])').wait_for(timeout=30000)
                await expect(page.locator('#download-summary')).to_be_visible()
                assert await page.locator('#result-table tbody tr').count() == 4
                await page.screenshot(path=str(out / '02-batch-summary.png'), full_page=True)
                async with page.expect_download() as download_info:
                    await page.locator('#download-summary').click()
                download = await download_info.value
                await download.save_as(out / 'downloaded-summary.xlsx')
                book = load_workbook(out / 'downloaded-summary.xlsx', read_only=True, data_only=True)
                assert len(book.sheetnames) == 1
                assert list(book.active.values) == [('大学名', '学院名', '是否出错'), ('测试大学', '学院0', 'No'),
                    ('测试大学', '学院1', 'Yes'), ('测试大学', '学院2', '学院网页无法访问'), ('测试大学', '学院3', '未完成核验')]
                book.close()
                await page.locator('#result-table tbody tr').first.click()
                await page.locator('#college-panel:not([hidden])').wait_for()
                assert '/0/1' in await page.locator('#college-content').inner_text()
                assert '/0/2' in await page.locator('#college-content').inner_text()
                await page.locator('[data-view="history"]').click()
                await page.locator('.history-card').first.click()
                await expect(page.locator('#download-summary')).to_be_visible()
                await page.set_viewport_size({'width': 760, 'height': 1000})
                assert await page.evaluate('document.documentElement.scrollWidth<=window.innerWidth')
                await browser.close()
        assert not errors, errors
        report = {'version': VERSION, 'actual_input_preview': {'files': 10, 'sources': 37, 'teacher_rows': 983},
                  'simulated_four_college_job': 'passed', 'single_sheet_three_column_download': 'passed',
                  'history_and_760px_layout': 'passed', 'browser_errors': errors, 'actual_new_websites_audited': False}
        (out / 'verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps({'output': str(out), **report}, ensure_ascii=False))
    finally:
        if server.store.worker:
            server.store.worker.join(timeout=15)
        server.shutdown()
        thread.join()
        server.server_close()


if __name__ == '__main__':
    asyncio.run(main())
