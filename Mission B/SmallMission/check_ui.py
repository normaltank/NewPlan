"""Developer end-to-end check. Uses a separate server, data directory and browser."""
import asyncio
import json
import os
import shutil
from pathlib import Path
import threading
import time

from openpyxl import load_workbook
from playwright.async_api import async_playwright, expect
from app import make_server
from runtime_paths import APP_ROOT


async def main():
    out=APP_ROOT/'outputs'/'ui-validation-v2.1-20261015'
    out.mkdir(parents=True,exist_ok=True)
    server=make_server(out/'data')
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    os.environ['PLAYWRIGHT_BROWSERS_PATH']=str(APP_ROOT/'.portable-browser')
    errors=[]
    try:
        async with async_playwright() as pw:
            browser=await pw.chromium.launch(headless=True)
            page=await browser.new_page(viewport={'width':1440,'height':1100},device_scale_factor=1)
            page.on('pageerror',lambda e:errors.append(str(e)))
            await page.goto(f'http://127.0.0.1:{server.server_port}/#token={server.token}')
            await expect(page.locator('#version')).to_have_text('VERSION 2.1.0')
            await page.screenshot(path=str(out/'01-import.png'),full_page=True)
            await page.locator('#paste').fill('学校\t学院\t师资url\t教师url\n坏行')
            await page.locator('#preview-button').click()
            await page.locator('#preview-errors .alert').wait_for()
            assert await page.locator('#start-button').is_disabled()
            await page.locator('#paste').fill('')
            files=list((APP_ROOT/'outputs'/'live-20261015-final'/'inputs').glob('*山东理工*.xlsx'))
            assert len(files)==1
            await page.locator('#files').set_input_files(files[0])
            await page.locator('#preview-button').click()
            await page.locator('#preview-groups table').wait_for()
            assert '122' in await page.locator('#preview-stats').inner_text()
            await page.screenshot(path=str(out/'02-preview.png'),full_page=True)
            await page.locator('#start-button').click()
            await page.locator('#view-job:not([hidden])').wait_for()
            await page.locator('#downloads:not([hidden])').wait_for(timeout=180000)
            assert '157' in await page.locator('#result-table').inner_text()
            assert '缺少 35 条' in await page.locator('#input-summary').inner_text()
            assert '缺少 35 条' in await page.locator('#result-table').inner_text()
            await page.locator('#result-table tbody tr').click()
            await page.locator('#college-panel:not([hidden])').wait_for()
            assert '159' in await page.locator('#college-content').inner_text()
            await page.screenshot(path=str(out/'03-results.png'),full_page=True)
            async with page.expect_download() as download_info:
                await page.locator('#download-excel').click()
            download=await download_info.value
            assert download.suggested_filename.startswith('山东理工大学_农业工程与食品科学学院_核验结果_')
            await download.save_as(out/'downloaded.xlsx')
            book=load_workbook(out/'downloaded.xlsx',read_only=True)
            assert book['修正后四列表'].max_row==158
            book.close()
            await page.locator('[data-view="history"]').click()
            await page.locator('.history-card').first.wait_for()
            await page.locator('.history-card').first.click()
            await page.locator('#downloads:not([hidden])').wait_for()
            await page.set_viewport_size({'width':760,'height':1000})
            await page.locator('[data-view="new"]').click()
            assert await page.evaluate('document.documentElement.scrollWidth<=window.innerWidth')
            await page.screenshot(path=str(out/'04-small-screen.png'),full_page=True)
            # Old 2.0 tasks must gain clear labels and readable download names
            # without rewriting their historical workbooks.
            old=APP_ROOT/'outputs'/'live-20261015-final'
            old_results=json.loads((old/'results.json').read_text(encoding='utf-8'))
            legacy_id='a'*32
            legacy_output=out/'data'/'runs'/legacy_id/'output'
            legacy_output.mkdir(parents=True,exist_ok=True)
            for name in ('results.json','manifest.json','MissionB_核验结果.xlsx'):
                shutil.copy2(old/name,legacy_output/name)
            with server.store.lock:
                server.store.jobs[legacy_id]={'id':legacy_id,'created':'2026-10-15T18:00:00+08:00','status':'completed','completed':10,'total':10,'message':'历史记录显示验证','results':old_results}
                server.store.save(server.store.jobs[legacy_id])
            await page.set_viewport_size({'width':1440,'height':1100})
            await page.locator('[data-view="history"]').click()
            await page.locator('.history-card').filter(has_text='10 / 10').click()
            text=await page.locator('#result-table').inner_text()
            assert '缺少 35 条' in text and '多出 2 条' in text and '无链接差异' in text and '尚未判定' in text
            assert '缺少 40 条，多出 3 条' in await page.locator('#input-summary').inner_text()
            await page.screenshot(path=str(out/'05-all-statuses.png'),full_page=True)
            async with page.expect_download() as legacy_download_info:
                await page.locator('#download-excel').click()
            legacy_download=await legacy_download_info.value
            assert legacy_download.suggested_filename.startswith('南京理工大学_能源与动力工程学院_等10学院_核验结果_')
            await browser.close()
        assert not errors,errors
        results=next(j for j in server.store.jobs.values() if j['id']!=legacy_id)
        (out/'validation.json').write_text(json.dumps({'browser_errors':errors,'job_status':results['status'],'stats':results['results'][0]['stats'],'flows':['format errors block run','Excel upload preview','live source fetch','progress and per-college detail','Excel download','history reopen','760px layout']},ensure_ascii=False,indent=2),encoding='utf-8')
        print(json.dumps({'status':'passed','output':str(out),'stats':results['results'][0]['stats']},ensure_ascii=False))
    finally:
        server.shutdown();thread.join();server.server_close()


if __name__=='__main__':
    asyncio.run(main())
