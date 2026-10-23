"""Distribution boundary tests; isolated local data and simulated model/network."""
import base64
import copy
import hashlib
import io
import json
import os
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch
import zipfile
import zlib

import httpx
from openpyxl import load_workbook

from app import atomic_json, brief, now
from lite_app import ARCHIVE_NAME, Journal, LiteStore, make_server, public_rows
from read_lite_log import restore
from reporting import export
from runtime_paths import APP_ROOT
from test_distribution import HEAD, ROW
from test_result_summary import sample_results, sample_manifest, HTML
from test_api_engine import FakeClient, HTML as PAGED_HTML, URL as PAGED_URL

KEY = 'sk-' + '0123456789abcdef' * 2  # Synthetic credential for redaction tests.


def fixtures():
    results = copy.deepcopy(sample_results())
    names = ('计算机学院', '信息学院', '物理学院', '化学学院', '材料学院', '数学学院', '机械学院')
    for result, college in zip(results, names):
        result['college'] = college
    return results


def seed(store, identifier='7' * 32, results=None):
    results = fixtures() if results is None else copy.deepcopy(results)
    directory = store.root / 'runs' / identifier / 'output'
    directory.mkdir(parents=True)
    manifest = sample_manifest()
    export(directory, results, manifest)
    atomic_json(directory / 'results.json', results)
    atomic_json(directory / 'manifest.json', manifest)
    job = dict(id=identifier, created=now(), status='completed', completed=len(results), total=len(results),
               results=[brief(r) for r in results], message='private detail')
    store.jobs[identifier] = job
    store.save(job)
    store.export_simple(job)
    return job


async def simulated_run(args, on_progress, should_cancel, results=None):
    """Inject deterministic results at the remote-work boundary, keep real exports."""
    args.output.mkdir(parents=True)
    results, manifest = fixtures() if results is None else copy.deepcopy(results), sample_manifest()
    for i, result in enumerate(results, 1):
        on_progress('college_finished', {'index': i, 'result': result})
    export(args.output, results, manifest)
    atomic_json(args.output / 'results.json', results)
    atomic_json(args.output / 'manifest.json', manifest)
    evidence = args.output / 'evidence' / 'fixture'
    evidence.mkdir(parents=True)
    (evidence / 'http.html').write_text(HTML, encoding='utf-8')
    api = args.output / 'api'
    api.mkdir()
    atomic_json(api / '01-check-response.json', {'reason': 'model fixture'})
    return args.output, results, manifest


class LiteTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=APP_ROOT)
        self.root = Path(self.temp.name)
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY': ''}):
            self.server = make_server(self.root / 'lite-data')
        self.store = self.server.store
        self.store.api_key = KEY
        self.store.journal.protect(KEY)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.client = httpx.Client(base_url=f'http://127.0.0.1:{self.server.server_port}', timeout=15, trust_env=False)
        response = self.client.post('/api/session', json={'token': self.server.token})
        self.assertEqual(response.status_code, 200)

    def tearDown(self):
        if self.store.worker:
            self.store.worker.join(10)
        self.client.close()
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(5)
        self.temp.cleanup()

    def test_page_shows_three_categories_and_download_contains_only_error_colleges(self):
        job = seed(self.store)
        original = (self.store.root / 'runs' / job['id'] / 'output' / 'results.json').read_bytes()
        state = self.client.get('/api/state').json()['job']
        self.assertEqual(state['counts'], {'正确': 1, '错误': 1, '特殊情况': 5})
        self.assertEqual(state['summary_headers'], ['学校', '学院', '核验结果'])
        self.assertEqual(state['summary_rows'], [row[:3] for row in public_rows(fixtures())])
        self.assertTrue(state['results_ready'])
        self.assertNotIn('reason', json.dumps(state))
        response = self.client.get(f"/api/jobs/{job['id']}/download")
        self.assertEqual(response.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            self.assertEqual(len(archive.namelist()), 1)
            name = archive.namelist()[0]
            self.assertTrue(name.startswith('核验后链接/'))
            self.assertIn('信息学院', name)
            with io.BytesIO(archive.read(name)) as content:
                book = load_workbook(content, read_only=True)
                rows = [[v or '' for v in r] for r in book.active.values]
                self.assertEqual(book.sheetnames, ['核验后链接'])
                self.assertEqual(rows[0], ['学校', '学院', '师资url', '教师url'])
                self.assertEqual({row[1] for row in rows[1:]}, {'信息学院'})
                self.assertEqual({row[3] for row in rows[1:]}, {'https://example.edu.cn/a', 'https://example.edu.cn/b'})
                book.close()
        self.assertEqual(original, (self.store.root / 'runs' / job['id'] / 'output' / 'results.json').read_bytes())
        self.assertIn('学院网页无法访问', original.decode())

    def test_twenty_uploaded_files_are_accepted_and_twenty_one_are_rejected(self):
        content = base64.b64encode((HEAD + ROW).encode()).decode()
        files = [{'name': f'名单{i}.txt', 'content': content} for i in range(21)]
        with patch('app.run', simulated_run):
            response = self.client.post('/api/start', json={'files': files[:20]})
            self.assertEqual(response.status_code, 201, response.text)
            self.store.worker.join(10)
            response = self.client.post('/api/start', json={'files': files})
            self.assertEqual(response.status_code, 400, response.text)
            self.assertIn('最多20个文件', response.json()['error'])
        self.assertEqual(self.client.get('/api/state').json()['max_files'], 20)

    def test_no_error_results_finish_without_download_and_survive_restart(self):
        for i, results in enumerate(([fixtures()[0]], fixtures()[2:], [fixtures()[0], fixtures()[-1]])):
            with self.subTest(categories=[r['college'] for r in results]):
                job = seed(self.store, identifier=str(i + 1) * 32, results=results)
                state = self.client.get(f"/api/jobs/{job['id']}").json()
                self.assertEqual(state['status'], 'completed')
                self.assertTrue(state['results_ready'])
                self.assertFalse(state['download_ready'])
                self.assertEqual(state['summary_rows'], [row[:3] for row in public_rows(results)])
                self.assertEqual(self.client.get(f"/api/jobs/{job['id']}/download").status_code, 400)
        recovered = LiteStore(self.store.root, self.store.journal)
        self.assertTrue(all(recovered.view(j['id'])['results_ready'] for j in recovered.jobs.values()))
        recovered.close()

    def test_error_college_download_keeps_completed_groups_and_excludes_other_colleges(self):
        good, bad, special = fixtures()[:3]
        for result in (good, bad, special):
            result['college'] = '合并学院'
        other = copy.deepcopy(good)
        other.update(id='f' * 12, school='另一个大学')
        job = seed(self.store, results=[good, bad, special, other])
        state = self.client.get(f"/api/jobs/{job['id']}").json()
        self.assertEqual(state['summary_rows'], [['示例大学', '合并学院', '错误'], ['另一个大学', '合并学院', '正确']])
        response = self.client.get(f"/api/jobs/{job['id']}/download")
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            self.assertEqual(len(archive.namelist()), 2)
            self.assertTrue(all('示例大学_合并学院' in name for name in archive.namelist()))

    def test_legacy_download_is_derived_without_overwriting_old_outputs(self):
        job = seed(self.store)
        directory = self.store.root / 'runs' / job['id']
        (directory / ARCHIVE_NAME).unlink()
        legacy = directory / '核验结果.zip'
        with zipfile.ZipFile(legacy, 'w') as archive:
            archive.writestr('使用说明.txt', '旧版内容保留')
            for path in (directory / 'output' / 'corrected').glob('*.xlsx'):
                archive.write(path, '核验后链接/' + path.name)
        before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in directory.rglob('*') if p.is_file()}
        response = self.client.get(f"/api/jobs/{job['id']}/download")
        self.assertEqual(response.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            self.assertEqual(len(archive.namelist()), 1)
            self.assertIn('信息学院', archive.namelist()[0])
        self.assertTrue(all(hashlib.sha256(p.read_bytes()).hexdigest() == digest for p, digest in before.items()))

    def test_advanced_routes_absent_and_local_session_enforced(self):
        job = seed(self.store)
        for path in ('/api/jobs', '/api/settings', '/api/config', '/api/preview', '/MissionB.log',
                     f"/api/jobs/{job['id']}/download/all", f"/api/jobs/{job['id']}/college/1"):
            self.assertEqual(self.client.get(path).status_code, 404, path)
        self.assertEqual(self.client.post('/api/settings', json={'mode':'rules'}).status_code, 404)
        self.assertEqual(self.client.get('/api/state', headers={'Cookie':''}).status_code, 401)
        self.assertEqual(self.client.get('/api/state', headers={'Host':'evil.invalid'}).status_code, 403)
        self.assertEqual(self.client.post('/api/start', json={}, headers={'Origin':'https://evil.invalid'}).status_code, 403)

    def test_start_export_log_and_restore(self):
        with patch('app.run', simulated_run):
            response = self.client.post('/api/start', json={'text': HEAD + ROW, 'api_key': KEY})
            self.assertEqual(response.status_code, 201, response.text)
            self.store.worker.join(10)
        job = self.client.get('/api/state').json()['job']
        self.assertEqual(job['status'], 'completed')
        self.assertTrue(job['download_ready'])
        text = self.store.journal.path.read_text(encoding='utf-8')
        self.assertNotIn(KEY, text)
        self.assertNotIn(self.server.token, text)
        records = [json.loads(line) for line in text.splitlines()]
        self.assertTrue(any(r['event'] == 'run_started' and r['inputs'] for r in records))
        self.assertTrue(any(r.get('path') == 'api/01-check-response.json' for r in records))
        output = self.root / 'restored'
        self.assertEqual(restore(self.store.journal.path, output), job['id'])
        self.assertEqual((output / 'evidence/evidence/fixture/http.html').read_text(encoding='utf-8'), HTML)
        result = json.loads((output / 'evidence/results.json').read_text(encoding='utf-8'))
        self.assertEqual(public_rows(result), public_rows(fixtures()))
        self.assertEqual(result[-1]['reason'], '')

    def test_failure_logs_redacted_traceback_without_exposing_cause(self):
        async def fail(*args, **kwargs):
            raise RuntimeError('internal provider failure ' + KEY)
        with patch('app.run', fail):
            response = self.client.post('/api/start', json={'text': HEAD + ROW, 'api_key': KEY})
            self.assertEqual(response.status_code, 201)
            self.store.worker.join(10)
        job = self.client.get('/api/state').json()['job']
        self.assertEqual(job['status'], 'failed')
        self.assertNotIn('provider', json.dumps(job))
        self.assertFalse(job['download_ready'])
        self.assertEqual(self.client.get(f"/api/jobs/{job['id']}/download").status_code, 400)
        text = self.store.journal.path.read_text(encoding='utf-8')
        self.assertNotIn(KEY, text)
        records = [json.loads(line) for line in text.splitlines()]
        error = next(r for r in records if r.get('path') == 'error.txt')
        raw = zlib.decompress(base64.b64decode(error['content'])).decode()
        self.assertIn('RuntimeError: internal provider failure [REDACTED]', raw)
        self.assertNotIn(KEY, raw)

    def test_missing_key_invalid_input_and_duplicate_start(self):
        self.store.api_key = ''
        self.assertEqual(self.client.post('/api/start', json={'text':HEAD + ROW}).status_code, 400)
        self.assertEqual(self.client.post('/api/start', json={'text':HEAD + ROW, 'api_key':KEY}).status_code, 400)
        self.store.api_key = KEY
        self.assertEqual(self.client.post('/api/start', json={'text':'bad', 'api_key':KEY}).status_code, 400)
        barrier = threading.Event()
        async def delayed(args, on_progress, should_cancel):
            barrier.wait(8)
            if should_cancel():
                raise InterruptedError('fixture cancellation')
            return await simulated_run(args, on_progress, should_cancel)
        with patch('app.run', delayed):
            response = self.client.post('/api/start', json={'text':HEAD + ROW, 'api_key':KEY})
            self.assertEqual(response.status_code, 201)
            self.assertEqual(self.client.post('/api/start', json={'text':HEAD + ROW}).status_code, 400)
            self.assertEqual(self.client.post('/api/shutdown', json={}).status_code, 400)
            identifier = response.json()['id']
            self.assertEqual(self.client.post(f'/api/jobs/{identifier}/cancel', json={}).json()['status'], 'stopping')
            barrier.set()
            self.store.worker.join(10)
        self.assertEqual(self.store.view(identifier)['status'], 'cancelled')

    def test_interrupted_finalization_recovers_without_stuck_page(self):
        for i, status in enumerate(('running', 'finalizing', 'completed')):
            identifier = str(i + 1) * 32
            job = dict(id=identifier, created=now(), status=status, completed=0, total=1, results=[])
            self.store.save(job)
        recovered = LiteStore(self.store.root, self.store.journal)
        self.assertTrue(all(j['status'] == 'interrupted' for j in recovered.jobs.values()))

    def test_preconfigured_credential_stays_on_server(self):
        assets = self.root / 'assets'
        assets.mkdir()
        atomic_json(assets / 'lite_config.json', {'api_key': KEY})
        with patch('lite_app.ASSETS', assets), patch.dict(os.environ, {'DEEPSEEK_API_KEY':''}):
            configured = LiteStore(self.root / 'configured-data', Journal(self.root / 'configured.log'))
        self.assertEqual(configured.api_key, KEY)
        self.assertNotIn(KEY, (self.root / 'configured.log').read_text(encoding='utf-8'))
        self.assertNotIn(KEY, self.client.get('/api/state').text)
        self.assertEqual(self.client.get('/lite_config.json').status_code, 404)
        self.assertEqual(self.client.get('/_internal/lite_config.json').status_code, 404)
        with patch('app.run', simulated_run):
            response = self.client.post('/api/start', json={'text':HEAD + ROW, 'api_key':'do-not-use-this-key', 'mode':'rules'})
            self.assertEqual(response.status_code, 201)
            self.store.worker.join(10)
        self.assertEqual(self.store.api_key, KEY)
        self.assertEqual(self.store.mode, 'deepseek')

    def test_log_redacts_multiple_keys_and_evidence(self):
        journal = self.store.journal
        journal.protect(KEY)
        journal.protect('another-credential-123')
        journal.write('example', text=KEY + ' another-credential-123 Bearer fake-token https://x/?api_key=hidden')
        text = journal.path.read_text(encoding='utf-8')
        for secret in (KEY, 'another-credential-123', 'fake-token', 'api_key=hidden'):
            self.assertNotIn(secret, text)
        for line in text.splitlines():
            json.loads(line)

    def test_real_pipeline_preserves_pagination_dedup_and_equal_count_difference(self):
        requests = []
        async def local_fetch(fetcher, task):
            url = task['roster_url']
            requests.append(url)
            self.assertIn(url, (PAGED_URL, 'https://example.edu.cn/staff/1.htm'))
            html = PAGED_HTML if url == PAGED_URL else '<title>教师名单</title><div class="roster"><a href="/c">丙</a><div class="pb_sys_common"><span>2/2</span></div></div>'
            meta = dict(requested_url=url, final_url=url, method='HTTP', status=200, captured_at=now())
            fetcher.save(task['id'], 'http', html, meta)
            return html, meta
        text = HEAD + '\n'.join(f'隔离大学\t{college}\t{PAGED_URL}\thttps://example.edu.cn/{path}'
            for college, last in (('正确学院','c'), ('差异学院','d')) for path in ('index.htm','a','a','b',last))
        with patch('audit.Fetcher.http', local_fetch), patch('api_engine.DeepSeekJSON', lambda *a, **k: FakeClient()):
            response = self.client.post('/api/start', json={'text':text, 'api_key':KEY})
            self.assertEqual(response.status_code, 201, response.text)
            self.store.worker.join(20)
        job = self.client.get('/api/state').json()['job']
        self.assertEqual(job['status'], 'completed')
        self.assertEqual(job['counts'], {'正确':1,'错误':1,'特殊情况':0})
        self.assertEqual(set(requests), {PAGED_URL, 'https://example.edu.cn/staff/1.htm'})
        out = self.store.root / 'runs' / job['id'] / 'output'
        manifest = json.loads((out / 'manifest.json').read_text(encoding='utf-8'))
        self.assertEqual(manifest['explicit_profile_fetches'], 0)
        results = json.loads((out / 'results.json').read_text(encoding='utf-8'))
        self.assertEqual([r['stats']['source_unique'] for r in results], [4,4])
        self.assertEqual([r['stats']['input_unique'] for r in results], [4,4])
        self.assertEqual([r['stats']['input_duplicates'] for r in results], [1,1])


if __name__ == '__main__':
    unittest.main()
