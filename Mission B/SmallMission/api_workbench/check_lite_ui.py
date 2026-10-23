"""Real browser and unpacked Windows executable checks, with isolated fixtures."""
import argparse
import asyncio
from collections import Counter
from datetime import datetime
import hashlib
import io
import json
import os
from pathlib import Path
import subprocess
import threading
import time
from unittest.mock import patch
import zipfile

import httpx
from openpyxl import load_workbook
from playwright.async_api import async_playwright, expect

from lite_app import Journal, LiteStore, make_server, LITE_VERSION
from lite_mail import MailReporter, load_config
from read_lite_log import restore
from reporting import SUMMARY_HEADERS
from runtime_paths import APP_ROOT
from test_lite import KEY, seed, simulated_run, fixtures
from test_distribution import HEAD, ROW


def check_download(path):
    with zipfile.ZipFile(path) as archive:
        assert len([n for n in archive.namelist() if n.startswith('核验后链接/')]) == 2
        book = load_workbook(io.BytesIO(archive.read('核验汇总.xlsx')), read_only=True)
        values = [[v if v is not None else '' for v in row] for row in book.active.values]
        assert book.sheetnames == ['核验汇总'] and values[0] == SUMMARY_HEADERS
        assert [r[2] for r in values[1:]] == ['正确','错误'] + ['特殊情况'] * 5
        assert all(r[3] == '' for r in values[1:])
        book.close()
        (path.parent / '核验汇总_实际导出样例.xlsx').write_bytes(archive.read('核验汇总.xlsx'))
    return dict(Counter(r[2] for r in values[1:]))


async def exercise_ui(base, token, out, server=None):
    errors = []
    async with async_playwright() as pw:
        browser = await pw.chromium.launch(headless=True)
        page = await browser.new_page(viewport={'width':1080, 'height':940}, device_scale_factor=1)
        page.on('pageerror', lambda error: errors.append(str(error)))
        await page.goto(base + '/#token=' + token)
        await expect(page.locator('#start')).to_be_enabled()
        assert await page.locator('input[type=password]').count() == 0
        for text in ('API','DeepSeek','模型','历史记录','逐条证据','特殊情况说明'):
            assert text not in await page.locator('body').inner_text(), text
        assert await page.evaluate('localStorage.length + sessionStorage.length') == 0
        if server:
            await page.screenshot(path=str(out / '输入界面.png'), full_page=True)
            await page.locator('#start').click()
            await expect(page.locator('#error')).to_contain_text('请选择文件')
            file = out / '示例名单.txt'
            file.write_text(HEAD + '\n'.join(ROW.replace('测试学院', result['college']) for result in fixtures()), encoding='utf-8')
            await page.locator('#files').set_input_files(file)
            gate, first_done, finalize = threading.Event(), threading.Event(), threading.Event()
            async def delayed(args, on_progress, should_cancel):
                await asyncio.to_thread(gate.wait, 10)
                def progress(event, data):
                    on_progress(event, data)
                    if event == 'college_finished' and data['index'] == 1:
                        first_done.wait(20)
                return await simulated_run(args, progress, should_cancel)
            export_simple = server.store.export_simple
            def delayed_export(job):
                finalize.wait(20)
                return export_simple(job)
            with patch('app.run', delayed), patch.object(server.store, 'export_simple', delayed_export):
                await page.locator('#start').click()
                await expect(page.locator('#result')).to_be_visible()
                await expect(page.locator('#status-title')).to_have_text('正在核验…')
                await expect(page.locator('#start')).to_be_disabled()
                await expect(page.locator('#download')).to_be_hidden()
                await expect(page.locator('#progress-percent')).to_have_text('0%')
                await expect(page.locator('#progress-text')).to_have_text('已完成 0 / 7')
                gate.set()
                await expect(page.locator('#progress-percent')).to_have_text('14%')
                await expect(page.locator('#progress')).to_have_attribute('value', '14')
                await expect(page.locator('#progress-text')).to_have_text('已完成 1 / 7')
                await page.screenshot(path=str(out / '核验进度.png'), full_page=True)
                await page.reload()
                await expect(page.locator('#result')).to_be_visible()
                await expect(page.locator('#status-title')).to_have_text('正在核验…')
                await expect(page.locator('#start')).to_be_disabled()
                await expect(page.locator('#progress-percent')).to_have_text('14%')
                first_done.set()
                await expect(page.locator('#status-title')).to_have_text('正在整理结果…')
                await expect(page.locator('#progress-percent')).to_have_text('')
                assert await page.locator('#progress').get_attribute('value') is None
                await expect(page.locator('#download')).to_be_hidden()
                finalize.set()
                await expect(page.locator('#download')).to_be_visible(timeout=15000)
                await expect(page.locator('#progress-percent')).to_have_text('100%')
            # Exercise the alternate paste input as well, without a visible API field.
            await page.locator('summary').click()
            await page.locator('#text').fill(HEAD + ROW)
            previous_download = await page.locator('#download').get_attribute('href')
            with patch('app.run', simulated_run):
                await page.locator('#start').click()
                await page.wait_for_function("previous => document.querySelector('#download').getAttribute('href') && document.querySelector('#download').getAttribute('href') !== previous", arg=previous_download)
                await expect(page.locator('#download')).to_be_visible(timeout=15000)
                await asyncio.to_thread(server.store.worker.join, 10)
            await page.locator('summary').click()
        await expect(page.locator('#status-title')).to_have_text('核验完成')
        await expect(page.locator('#progress')).to_have_attribute('value', '100')
        await expect(page.locator('#progress-percent')).to_have_text('100%')
        await expect(page.locator('#status-text')).to_contain_text('人工查看原网页')
        for reason in ('分页不完整','学院网页无法访问','模型判断不一致','名单区域不明确'):
            assert reason not in await page.locator('body').inner_text()
        await page.screenshot(path=str(out / '结果界面.png'), full_page=True)
        async with page.expect_download() as event:
            await page.locator('#download').click()
        download = await event.value
        target = out / '核验结果_实际导出样例.zip'
        await download.save_as(target)
        counts = check_download(target)
        state = await page.request.get(base + '/api/state')
        state = await state.json()
        assert state['job']['counts'] == counts
        await page.set_viewport_size({'width':390, 'height':850})
        assert await page.evaluate('document.documentElement.scrollWidth <= innerWidth')
        await page.screenshot(path=str(out / '窄屏界面.png'), full_page=True)
        await page.reload()
        await expect(page.locator('#download')).to_be_visible()
        if server:
            successful = state['job']['id']
            assert restore(out / 'MissionB.log', out / '日志还原', successful) == successful
            async def fail(*a, **k):
                raise RuntimeError('simulated private provider reason ' + KEY)
            await page.locator('summary').click()
            await page.locator('#text').fill(HEAD + ROW)
            with patch('app.run', fail):
                await page.locator('#start').click()
                await expect(page.locator('#status-title')).to_have_text('没能完成核验', timeout=15000)
            await expect(page.locator('#download')).to_be_hidden()
            await expect(page.locator('#progress-area')).to_be_hidden()
            assert 'private provider' not in await page.locator('body').inner_text()
            assert KEY not in (out / 'MissionB.log').read_text(encoding='utf-8')
        await page.locator('#quit').click()
        await expect(page.locator('#error')).to_contain_text('工具已退出')
        assert not errors, errors
        await browser.close()
    return {'categories':counts, 'empty_notes':True, 'no_api_ui':True, 'no_reason_ui':True,
            'refresh_download':True, 'mobile_overflow':False, 'browser_errors':errors,
            'completion_progress':100, 'progress_lifecycle_checked':bool(server),
            'export_sha256':hashlib.sha256(target.read_bytes()).hexdigest()}


async def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--portable', type=Path)
    parser.add_argument('--send-test-mail', action='store_true', help='Explicitly send one synthetic support mail from the portable executable')
    parser.add_argument('--output-root', type=Path, default=APP_ROOT / 'validation/lite-20261022')
    args = parser.parse_args()
    if args.send_test_mail and not args.portable:
        parser.error('--send-test-mail requires --portable')
    root = args.output_root.resolve()
    root.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime('%H%M%S')
    if not args.portable:
        out = root / ('ui-' + stamp)
        out.mkdir()
        candidates = sorted(APP_ROOT.glob('.lite-stage-*/MissionB*/.browsers'))
        os.environ['PLAYWRIGHT_BROWSERS_PATH'] = str(candidates[-1])
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY':''}):
            server = make_server(out / 'lite-data')
        server.store.api_key = KEY
        server.store.journal.protect(KEY)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            value = await exercise_ui(f'http://127.0.0.1:{server.server_port}', server.token, out, server)
        finally:
            server.shutdown()
            server.server_close()
            thread.join(5)
        value['input_and_paste'] = value['failure_log'] = value['log_restore'] = True
    else:
        out = root / ('portable-' + stamp)
        out.mkdir()
        extracted = out / '中文路径 软件解压'
        extracted.mkdir()
        with zipfile.ZipFile(args.portable) as archive:
            names = archive.namelist()
            assert not any('/lite-data/' in n or n.endswith('MissionB.log') or '/web/' in n for n in names)
            archive.extractall(extracted)
        folder = next(extracted.iterdir())
        # Seed only isolated result data: no real API invocation in executable checks.
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY':''}):
            store = LiteStore(folder / 'lite-data', Journal(folder / 'MissionB.log'))
        test_job_id = None
        if args.send_test_mail:
            with patch('app.run', simulated_run):
                store.api_key = KEY
                store.journal.protect(KEY)
                test_job_id = store.submit({'text': HEAD + ROW})['id']
                store.worker.join(10)
            assert store.view(test_job_id)['download_ready']
            config = load_config(folder / '_internal/mail_config.json')
            store.journal.protect(config['password'])
            with patch('lite_mail.send_mail', side_effect=ConnectionError('离线准备测试邮件，待便携程序启动后发送')):
                reporter = MailReporter(folder / 'lite-data/mail-outbox', store.journal, config,
                                        LITE_VERSION + '（自动回传测试，合成数据）')
                reporter.submit(test_job_id, 5)
                reporter.pending.join()
                reporter.close()
        else:
            seed(store)
        store.close()
        env = {k:v for k,v in os.environ.items() if not k.upper().startswith(('PYTHON','DEEPSEEK'))}
        env['PATH'] = str(Path(os.environ['SystemRoot']) / 'System32')
        executable = folder / '名单核验.exe'
        process = subprocess.Popen([str(executable), '--no-browser'], cwd=folder, env=env,
                                   creationflags=subprocess.CREATE_NO_WINDOW)
        try:
            info_file = folder / 'lite-data/instance.json'
            deadline = time.monotonic() + 45
            while not info_file.exists():
                assert process.poll() is None, 'Executable stopped before initialization'
                assert time.monotonic() < deadline, 'Executable startup timeout'
                await asyncio.sleep(.2)
            info = json.loads(info_file.read_text(encoding='utf-8'))
            base = f"http://127.0.0.1:{info['port']}"
            token = info['url'].split('#token=')[1]
            with httpx.Client(base_url=base, trust_env=False) as client:
                assert client.get('/health').json()['version'] == LITE_VERSION
                assert client.post('/api/session', json={'token':token}).status_code == 200
                assert client.get('/api/state').json()['configured'] is True
                assert client.get('/api/settings').status_code == 404
                assert client.get('/api/jobs').status_code == 404
                assert client.get('/lite_config.json').status_code == 404
                assert client.get('/api/template.xlsx').status_code == 200
                assert client.get('/mail_config.json').status_code == 404
                assert client.post('/api/start', json={'text':'invalid input'}).status_code == 400
                configured_mail = (folder / '_internal/mail_config.json').is_file()
                assert client.get('/api/state').json()['automatic_support_mail'] == configured_mail
            mail_receipt = None
            if test_job_id:
                mail_state = folder / 'lite-data/mail-outbox' / test_job_id / 'state.json'
                deadline = time.monotonic() + 45
                while True:
                    mail_receipt = json.loads(mail_state.read_text(encoding='utf-8'))
                    if mail_receipt['status'] in ('sent', 'failed', 'uncertain'):
                        break
                    assert time.monotonic() < deadline, 'Portable SMTP acceptance timeout'
                    await asyncio.sleep(.2)
                assert mail_receipt['status'] == 'sent', {k:mail_receipt[k] for k in ('status','attempts')}
            duplicate = subprocess.run([str(executable), '--no-browser'], cwd=folder, env=env,
                                       creationflags=subprocess.CREATE_NO_WINDOW, timeout=20)
            assert duplicate.returncode == 0
            os.environ['PLAYWRIGHT_BROWSERS_PATH'] = str(folder / '.browsers')
            value = await exercise_ui(base, token, out)
            assert process.wait(timeout=15) == 0
            key = json.loads((folder / '_internal/lite_config.json').read_text(encoding='utf-8'))['api_key']
            assert key and key not in (folder / 'MissionB.log').read_text(encoding='utf-8')
            if configured_mail:
                mail_config = load_config(folder / '_internal/mail_config.json')
                assert mail_config['password'] not in (folder / 'MissionB.log').read_text(encoding='utf-8')
            value.update(embedded_key_loaded=True, clean_path=True, chinese_path=True, duplicate_launch=True,
                         shutdown_exit=0, bundled_browser=True, real_api_calls=0)
            value['automatic_mail_configured'] = configured_mail
            if mail_receipt:
                value['mail_test'] = {k:mail_receipt[k] for k in ('job_id','status','attempts')}
                value['mail_test']['gmail_inbox_confirmed'] = False
                support = folder / 'lite-data/mail-outbox' / test_job_id / 'support.zip'
                with zipfile.ZipFile(support) as archive:
                    mailed_log = out / '回传日志_测试样例.log'
                    mailed_log.write_bytes(archive.read('MissionB.log'))
                assert restore(mailed_log, out / '回传日志还原') == test_job_id
                value['mail_test']['log_restore'] = True
                # Start a fresh executable process, not just a second browser tab.
                prior_mtime = info_file.stat().st_mtime_ns
                process = subprocess.Popen([str(executable), '--no-browser'], cwd=folder, env=env,
                                           creationflags=subprocess.CREATE_NO_WINDOW)
                deadline = time.monotonic() + 30
                while info_file.stat().st_mtime_ns == prior_mtime:
                    assert process.poll() is None and time.monotonic() < deadline
                    await asyncio.sleep(.2)
                new_info = json.loads(info_file.read_text(encoding='utf-8'))
                with httpx.Client(base_url=f"http://127.0.0.1:{new_info['port']}", trust_env=False) as client:
                    client.post('/api/session', json={'token': new_info['url'].split('#token=')[1]}).raise_for_status()
                    assert client.get('/api/state').json()['job']['download_ready']
                    await asyncio.sleep(.5)
                    client.post('/api/shutdown', json={}).raise_for_status()
                assert process.wait(timeout=15) == 0
                unchanged = json.loads(mail_state.read_text(encoding='utf-8'))
                assert unchanged == mail_receipt
                records = [json.loads(line) for line in (folder / 'MissionB.log').read_text(encoding='utf-8').splitlines()]
                assert sum(r['event'] == 'mail_sent' and r.get('job_id') == test_job_id for r in records) == 1
                value['mail_test']['restart_did_not_resend'] = True
        finally:
            if process.poll() is None:
                process.terminate()
                process.wait(timeout=10)
    (out / 'verification.json').write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'directory':str(out), 'checks':value}, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    asyncio.run(main())
