"""Presentation contract: fixed categories, uncertainty, and shared export/UI data."""
import copy
import json
from pathlib import Path
import tempfile
import unittest

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from ai_regions import prepare
from api_engine import qualification, check_coverage
from app import brief, job_view, Store
from audit import parse_roster, reconcile
from reporting import verification_summary, college_summary_rows, export, SUMMARY_HEADERS
from runtime_paths import APP_ROOT
from test_audit import SOURCE, PLAN, task


HTML = '<title>教师名单</title><div class="roster"><a href="/a">甲</a><a href="/b">乙</a></div>'


def sample_results():
    """Synthetic inputs processed by the real parser/comparator, never live claims."""
    cases = [('正确学院', None), ('差异学院', None), ('网页访问异常学院', '学院网页无法访问'),
             ('分页待核学院', '分页不完整'), ('名单待核学院', '名单区域不明确'),
             ('模型分歧学院', '两次模型判断不一致'), ('原因待补学院', '')]
    results = []
    for i, (college, issue) in enumerate(cases):
        t = task('https://example.edu.cn/a', 'https://example.edu.cn/c' if i == 1 else 'https://example.edu.cn/b')
        t.update(id=f'{i + 1:012x}', school='示例大学', college=college)
        parsed = parse_roster(HTML, SOURCE, PLAN)
        if issue is not None:
            parsed.update(complete=False, issues=[issue] if issue else [])
        results.append({**{k: v for k, v in t.items() if k != 'rows'}, **reconcile(t, parsed),
                        'source': parsed, 'attempts': [{'method': '隔离测试样本', 'captured_at': '2026-10-16T00:00:00Z'}]})
    return results


def sample_manifest():
    return {'scope': '输出格式验证：隔离测试样本，不代表真实学院核验', 'configuration_note': '使用固定名单和可控异常测试输出',
            'mode': '隔离测试样本', 'model_api_calls': 0, 'explicit_profile_fetches': 0,
            'finished_at': '2026-10-16T00:00:00Z', 'needs_human': 5, 'run_id': 'sample20261016'}


class ResultSummaryTests(unittest.TestCase):
    def test_three_categories_and_empty_reason(self):
        results = sample_results()
        self.assertEqual([verification_summary(r)['verification_result'] for r in results],
                         ['正确', '错误'] + ['特殊情况'] * 5)
        self.assertEqual(verification_summary(results[-1]), {'verification_result': '特殊情况', 'special_note': ''})
        # Even apparent counts/differences cannot authorize a decision if unfinished.
        uncertain = {**results[-1], 'stats': {'missing': 5, 'extra': 1}}
        self.assertEqual(verification_summary(uncertain)['verification_result'], '特殊情况')
        for r in ({}, {'complete': True, 'stats': {}}, {'complete': True, 'stats': {'missing': None, 'extra': None}}):
            self.assertEqual(verification_summary(r)['verification_result'], '特殊情况')

    def test_equal_counts_different_urls_and_repaired_output_still_error(self):
        r = sample_results()[1]
        self.assertEqual(r['stats']['input_unique'], r['stats']['source_unique'])
        self.assertEqual(set(r['corrected_urls']), set(r['source']['links']))
        self.assertEqual(verification_summary(r)['verification_result'], '错误')
        r['verification_result'] = '正确'  # Raw evidence takes precedence over a stale label.
        self.assertEqual(verification_summary(r)['verification_result'], '错误')

    def test_actual_model_disagreement_and_missing_page_are_special(self):
        soup, nodes, anchors, _ = prepare(HTML, SOURCE)
        ids = [i for i, n in nodes.items() if n.name == 'a']
        plans = [{'status': 'ready', 'roster_regions': [i], 'pagination_regions': [], 'exclude_regions': [], 'reason': ''} for i in ids]
        parsed, _ = qualification(soup, nodes, anchors, plans, SOURCE)
        r = reconcile(task('https://example.edu.cn/a'), parsed)
        self.assertEqual(verification_summary(r)['verification_result'], '特殊情况')
        self.assertTrue(verification_summary(r)['special_note'])
        sources = [{'roster_url': SOURCE, 'parsed': {'complete': True, 'issues': [], 'pager_counters': [(1, 3)]}}]
        check_coverage(sources)
        parsed = parse_roster(HTML, SOURCE, PLAN)
        parsed.update(sources[0]['parsed'])
        r = reconcile(task('https://example.edu.cn/a'), parsed)
        self.assertEqual(verification_summary(r)['verification_result'], '特殊情况')

    def test_one_row_per_college_and_existing_confirmed_error_precedence(self):
        good, bad, special = copy.deepcopy(sample_results()[:3])
        for r in (good, bad, special):
            r['college'] = '同学院'
        self.assertEqual(college_summary_rows([good, special]), [['示例大学', '同学院', '特殊情况', '学院网页无法访问']])
        self.assertEqual(college_summary_rows([bad, special]), [['示例大学', '同学院', '错误', '学院网页无法访问']])
        self.assertEqual(college_summary_rows([good, good]), [['示例大学', '同学院', '正确', '']])
        special['reason'] = ''
        self.assertEqual(college_summary_rows([good, special]), [['示例大学', '同学院', '特殊情况', '']])

    def test_app_brief_roundtrip_retains_set_decision_and_notes(self):
        results = sample_results()
        # Protocol-only duplicate historically yields counts that differ from set classification.
        t = task('http://example.edu.cn/a', 'https://example.edu.cn/a', 'https://example.edu.cn/b')
        parsed = parse_roster(HTML, SOURCE, PLAN)
        results.append({**t, **reconcile(t, parsed), 'source': parsed})
        for r in results:
            self.assertEqual(verification_summary(r), verification_summary(brief(brief(r))))
        self.assertEqual(verification_summary(results[-1])['verification_result'], '正确')

    def test_export_static_report_and_app_share_exact_four_columns(self):
        with tempfile.TemporaryDirectory(dir=APP_ROOT) as tmp:
            root = Path(tmp)
            results, manifest = sample_results(), sample_manifest()
            before = copy.deepcopy(results)
            export(root, results, manifest)
            expected = [SUMMARY_HEADERS] + college_summary_rows(results)
            for filename in (manifest['summary_filename'], manifest['report_filename']):
                book = load_workbook(root / filename, read_only=True, data_only=True)
                actual = [[v if v is not None else '' for v in row] for row in book['核验汇总'].values]
                self.assertEqual(actual, expected)
                self.assertEqual(book['核验汇总'].max_column, 4)
                book.close()
            soup = BeautifulSoup((root / 'report.html').read_text('utf-8'), 'html.parser')
            rows = [[c.get_text() for c in row.select('th,td')] for row in soup.select('#verification-summary tr')]
            self.assertEqual(rows, expected)
            view = job_view({'results': results})
            self.assertEqual(view['summary_headers'], SUMMARY_HEADERS)
            self.assertEqual([[r['school'], r['college'], r['verification_result'], r['special_note']] for r in view['college_summaries']], expected[1:])
            for old, new in zip(before, results):
                self.assertEqual(old, {k: v for k, v in new.items() if k != 'corrected_filename'})

    def test_completed_history_is_read_without_rewriting_files(self):
        with tempfile.TemporaryDirectory(dir=APP_ROOT) as tmp:
            root = Path(tmp)
            identifier = '1' * 32
            path = root / 'runs' / identifier / 'job.json'
            path.parent.mkdir(parents=True)
            old = sample_results()
            for r in old:
                r['original_error'] = 'No' if r['college'] == '正确学院' else 'Yes' if r['college'] == '差异学院' else '未完成核验'
                r.pop('source')
                r.pop('diffs')
            path.write_text(json.dumps({'id': identifier, 'status': 'completed', 'results': old}, ensure_ascii=False), 'utf-8')
            before = path.read_bytes()
            store = Store(root)
            view = store.get(identifier)
            self.assertEqual([r['verification_result'] for r in view['college_summaries']], ['正确', '错误'] + ['特殊情况'] * 5)
            self.assertEqual(path.read_bytes(), before)


if __name__ == '__main__':
    unittest.main()
