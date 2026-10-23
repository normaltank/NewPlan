"""Behavioral checks for the URL-only contract; no network or main database."""
import asyncio
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from audit import ROOT, Fetcher, load_inputs, parse_roster, reconcile, sha, url_key
from reporting import export

SOURCE = 'https://example.edu.cn/faculty/list.htm'
PLAN = {'selector': '.roster', 'title_contains': '教师'}


def parse(body):
    return parse_roster('<title>教师名单</title>' + body, SOURCE, PLAN)


def task(*values):
    return {'id': 'test', 'school': '测试大学', 'college': '测试学院', 'roster_url': SOURCE,
            'rows': [{'file': 'input.xlsx', 'sheet': 'Sheet1', 'row': i + 2,
                      'value': v, 'url': url_key(v), 'values': ['测试大学', '测试学院', SOURCE, v],
                      'hyperlink_mismatch': False} for i, v in enumerate(values)]}


class ContractTests(unittest.TestCase):
    def test_same_count_different_links_is_not_a_pass(self):
        result = reconcile(task('https://example.edu.cn/a', 'https://example.edu.cn/c'),
                           parse('<div class="roster"><a href="/a">甲</a><a href="/b">乙</a></div>'))
        self.assertEqual((result['stats']['missing'], result['stats']['extra']), (1, 1))
        self.assertEqual(set(result['corrected_urls']), {'https://example.edu.cn/a', 'https://example.edu.cn/b'})

    def test_unreadable_source_never_deletes_input(self):
        result = reconcile(task('https://example.edu.cn/a', '无详细页' + SOURCE), parse('<p>Access denied</p>'))
        self.assertTrue(result['needs_human'])
        self.assertIsNone(result['stats']['missing'])
        self.assertEqual([d['new_url'] for d in result['diffs']], ['https://example.edu.cn/a', '无详细页' + SOURCE])

    def test_duplicate_ids_and_overlapping_selectors_do_not_lose_or_double_count_links(self):
        body = '<title>教师</title><div id="names"><a href="/a">甲</a></div><div id="names"><a href="/b">乙</a><div class="inner"><a href="/c">丙</a></div></div>'
        result = parse_roster(body, SOURCE, {'selector': '#names,.inner', 'title_contains': '教师'})
        self.assertEqual(len(result['links']), 3)
        self.assertEqual(sum(x['occurrences'] for x in result['links'].values()), 3)

    def test_protocol_difference_is_explicit_replacement(self):
        result = reconcile(task('http://example.edu.cn/a'), parse('<div class="roster"><a href="/a">甲</a></div>'))
        self.assertEqual(result['stats']['protocol_changes'], 1)
        self.assertEqual(result['stats']['missing'], 0)
        self.assertEqual(result['stats']['extra'], 0)
        self.assertEqual(result['stats']['literal_missing'], 1)
        self.assertEqual(result['diffs'][0]['action'], '按源修正协议')

    def test_different_query_and_fragment_are_not_merged(self):
        result = parse('<div class="roster"><a href="/people?id=1#x">甲</a><a href="/people?id=2#x">乙</a><a href="/people?id=1#y">丙</a></div>')
        self.assertEqual(len(result['links']), 3)

    def test_placeholders_and_empty_anchor_have_explicit_evidence(self):
        result = parse('<div class="roster"><a href="#">甲</a><a href="">乙</a><a href="/bogus"></a><a href="/ok">丙</a></div>')
        self.assertEqual(len(result['anchors']), 4)
        self.assertEqual(list(result['links']), ['https://example.edu.cn/ok'])
        self.assertTrue(result['complete'])

    def test_target_status_not_used_and_home_link_retained(self):
        result = parse('<div class="roster"><a href="https://unreachable.invalid/missing">甲</a><a href="/index.htm">乙</a></div>')
        self.assertTrue(result['complete'])
        self.assertIn('https://unreachable.invalid/missing', result['links'])
        self.assertIn('https://example.edu.cn/index.htm', result['links'])
        self.assertEqual(result['notes'][0]['kind'], '官网目标疑似首页')

    def test_names_are_not_filtered_by_language_or_length(self):
        result = parse('<div class="roster"><a href="/a">W. C. Hoffmann</a><a href="/b">𠮷某某*</a><a href="/c">赵  娜</a></div>')
        self.assertEqual(len(result['links']), 3)

    def test_same_url_multiple_names_does_not_become_multiple_urls(self):
        result = parse('<div class="roster"><a href="/a">甲</a><a href="/a">乙</a></div>')
        self.assertEqual(len(result['links']), 1)
        self.assertEqual(result['links']['https://example.edu.cn/a']['occurrences'], 2)
        self.assertEqual(result['notes'][0]['kind'], '同一目标对应不同姓名')

    def test_input_duplicates_and_markers_are_separate(self):
        result = reconcile(task('https://example.edu.cn/a', 'https://example.edu.cn/a', '无详细页' + SOURCE),
                           parse('<div class="roster"><a href="/a">甲</a></div>'))
        self.assertEqual(result['stats']['input_duplicates'], 1)
        self.assertEqual(result['stats']['input_placeholders'], 1)
        self.assertEqual(result['stats']['corrected'], 1)

    def test_paging_or_script_target_requires_attention(self):
        for extra in ('<a href="list2.htm">下一页</a>', '<a href="javascript:openPerson()">乙</a>'):
            result = parse('<div class="roster"><a href="/a">甲</a>' + extra + '</div>')
            self.assertFalse(result['complete'])

    def test_base_element_used(self):
        result = parse('<base href="https://other.edu.cn/people/"><div class="roster"><a href="person.htm">甲</a></div>')
        self.assertIn('https://other.edu.cn/people/person.htm', result['links'])

    def test_url_key_preserves_meaningful_differences(self):
        self.assertIsNone(url_key('无详细页' + SOURCE))
        self.assertEqual(url_key('../a', SOURCE), 'https://example.edu.cn/a')
        self.assertNotEqual(url_key('http://EXAMPLE.edu.cn/a'), url_key('https://example.edu.cn/a'))

    def test_hyperlink_mismatch_not_silently_resolved(self):
        t = task('https://example.edu.cn/a')
        t['rows'][0]['hyperlink_mismatch'] = True
        result = reconcile(t, parse('<div class="roster"><a href="/a">甲</a></div>'))
        self.assertTrue(result['needs_human'])

    def test_empty_or_unknown_region_cannot_be_a_zero_link_pass(self):
        self.assertFalse(parse('<div class="roster"></div>')['complete'])
        self.assertFalse(parse_roster('<title>教师</title><a href="/a">甲</a>', SOURCE, None)['complete'])

    def test_unresolved_export_does_not_create_an_empty_corrected_workbook(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as directory:
            out = Path(directory)
            t = task('https://example.edu.cn/old', '无详细页' + SOURCE)
            parsed = parse('<title>拒绝访问</title>')
            result = {**t, **reconcile(t, parsed), 'source': parsed,
                      'attempts': [{'method': '测试', 'status': 412, 'captured_at': 'test', 'html_path': 'test.html'}]}
            manifest = {'scope': '测试', 'configuration_note': '测试', 'mode': '测试',
                        'model_api_calls': 0, 'explicit_profile_fetches': 0, 'finished_at': 'test', 'needs_human': 1}
            export(out, [result], manifest)
            self.assertEqual(list((out / 'corrected').iterdir()), [])
            book = load_workbook(next(out.glob('*_核验结果_*.xlsx')), read_only=True)
            self.assertEqual(book['待人工处理原始行'].max_row, 3)
            self.assertEqual(book['学院结论']['D2'].value, '是')
            book.close()

    def test_serialized_workbook_preserves_four_columns_and_unresolved_rows(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as directory:
            out = Path(directory)
            t = task('https://example.edu.cn/old')
            parsed = parse('<div class="roster"><a href="/new">甲</a></div>')
            result = {**t, **reconcile(t, parsed), 'source': parsed,
                      'attempts': [{'method': '测试快照', 'status': 200, 'captured_at': 'test', 'html_path': 'test.html'}]}
            manifest = {'scope': '测试', 'configuration_note': '测试', 'mode': '测试',
                        'model_api_calls': 0, 'explicit_profile_fetches': 0, 'finished_at': 'test', 'needs_human': 0}
            export(out, [result], manifest)
            book = load_workbook(next((out / 'corrected').glob('测试大学_测试学院_修正后链接_*.xlsx')), read_only=True)
            self.assertEqual(book.active.max_column, 4)
            self.assertEqual(book.active['D2'].value, 'https://example.edu.cn/new')
            book.close()

    def test_live_fetcher_requests_roster_not_teacher_targets(self):
        class FakeResponse:
            content = b'<title>Faculty</title><a href="https://target.invalid/person">Teacher</a>'
            status_code = 200
            url = SOURCE
            history = []
        class FakeClient:
            calls = []
            async def get(self, url):
                self.calls.append(url)
                return FakeResponse()
        with tempfile.TemporaryDirectory(dir=ROOT) as directory:
            fetcher = Fetcher(Path(directory), browser=False)
            fetcher.client = FakeClient()
            asyncio.run(fetcher.http(task('https://target.invalid/person')))
            self.assertEqual(fetcher.client.calls, [SOURCE])

    def test_saved_evidence_preserves_mixed_newlines_and_digest(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as directory:
            out = Path(directory)
            html = '<title>教师</title>\r\n<div>名单</div>\n'
            meta = {'requested_url': SOURCE}
            Fetcher(out, browser=False).save('test', 'http', html, meta)
            raw = (out / 'evidence/test/http.html').read_bytes()
            self.assertEqual(raw, html.encode('utf-8'))
            self.assertEqual(sha(raw), meta['html_sha256'])


if __name__ == '__main__':
    unittest.main(verbosity=2)
