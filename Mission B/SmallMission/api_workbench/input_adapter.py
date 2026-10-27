"""Four-column and paired-sheet input adapters with source coordinates."""
from __future__ import annotations

import csv
import hashlib
import io
from pathlib import Path
import re
from urllib.parse import urlsplit, urlunsplit
import zipfile
from types import SimpleNamespace

from openpyxl import load_workbook

SUPPORTED = {'.xlsx', '.xlsm', '.xls', '.csv', '.tsv', '.txt'}
MAX_ROWS = 50000
MAX_FILE_BYTES = 10 * 1024 * 1024
ALIASES = [
    {'学校', '大学', '大学名', '学校名称', '大学名称', 'school', 'university'},
    {'学院', '学院名', '学院名称', '院系', 'college', 'department'},
    {'师资url', '师资链接', '师资列表链接', '师资列表url', '师资网址', '教师名录链接', '名单链接', '名单url', 'rosterurl', 'facultyurl'},
    {'教师url', '教师个人url', '教师链接', '教师个人链接', '个人url', '个人主页', '个人主页url', '主页', '主页url', 'teacherurl', 'profileurl'},
]


def key(value):
    return re.sub(r'[\s_\-：:（）()]', '', str(value or '')).lower().lstrip('\ufeff')


def plain_link(value):
    text = str(value or '').strip().strip('`').strip()
    m = re.fullmatch(r'\[[^\]]*\]\(\s*(https?://[^\s]+)\s*\)', text)
    if m:
        text = m[1]
    return text.strip('<>')


def absolute_url(value):
    text = plain_link(value)
    if not text or re.search(r'[\s<>]', text):
        return None
    try:
        p = urlsplit(text)
        if p.scheme.lower() not in ('http', 'https') or not p.hostname or p.username or p.password:
            return None
        port = p.port
    except ValueError:
        return None
    host = p.hostname.lower()
    if ':' in host:
        host = '[' + host + ']'
    if port and (p.scheme.lower(), port) not in [('http', 80), ('https', 443)]:
        host += ':' + str(port)
    return urlunsplit((p.scheme.lower(), host, p.path or '/', p.query, p.fragment))


def decode_text(raw):
    if raw.startswith((b'\xff\xfe', b'\xfe\xff')):
        return raw.decode('utf-16')
    if b'\0' in raw[:200]:
        raise ValueError('文本包含空字符，请另存为 UTF-8、带BOM的UTF-16或GBK文本')
    for encoding in ('utf-8-sig', 'gb18030'):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    raise ValueError('无法识别文本编码，请另存为UTF-8')


def text_rows(text, extension='.txt'):
    """Return physical start-line numbers, including quoted multiline CSV."""
    lines = text.lstrip('\ufeff').splitlines()
    nonblank = [line for line in lines if line.strip()]
    if not nonblank:
        return []
    first = nonblank[0].strip()
    if first.startswith('|') and first.count('|') >= 3:
        result = []
        for number, line in enumerate(lines, 1):
            if not line.strip():
                continue
            fields = [x.strip() for x in line.strip().strip('|').split('|')]
            if all(re.fullmatch(r':?-{2,}:?', x) for x in fields):
                continue
            result.append((number, fields, []))
        return result
    if '\t' in first or extension == '.tsv':
        delimiter = '\t'
    elif extension == '.csv' or ',' in first:
        delimiter = ','
    elif '，' in first:
        delimiter = '，'
    else:
        return [(i, re.split(r'\s+', line.strip()), []) for i, line in enumerate(lines, 1) if line.strip()]
    reader = csv.reader(io.StringIO(text.lstrip('\ufeff'), newline=''), delimiter=delimiter, strict=True)
    result = []
    while True:
        line = reader.line_num + 1
        try:
            row = next(reader)
        except StopIteration:
            break
        if any(str(x).strip() for x in row):
            result.append((line, row, []))
    return result


def excel_cell(cell, is_url):
    if cell is None:
        return '', False
    value = str(cell.value or '').strip()
    mismatch = False
    if is_url:
        # Only a literal HYPERLINK address is understood. Never execute formulas.
        if getattr(cell, 'data_type', None) == 'f':
            m = re.fullmatch(r'=\s*HYPERLINK\(\s*"((?:[^"]|"")+)"\s*[,;]\s*"(?:[^"]|"")*"\s*\)', value, re.I)
            if not m:
                raise ValueError('URL列含无法静态读取的公式；请粘贴为值或使用普通超链接')
            value = m[1].replace('""', '"')
        target = cell.hyperlink.target if getattr(cell, 'hyperlink', None) else None
        if target:
            if absolute_url(value):
                mismatch = absolute_url(value) != absolute_url(target)
            elif absolute_url(target):
                value = target
    return plain_link(value) if is_url else value, mismatch


def header_fields(values):
    found = {}
    for col, value in enumerate(values):
        k = key(value)
        for index, aliases in enumerate(ALIASES):
            if k in aliases:
                if index in found:
                    raise ValueError('同一种字段出现多个表头，请保留一列')
                found[index] = col
    return found


def header_map(values):
    found = header_fields(values)
    if len(found) == 4:
        return [found[i] for i in range(4)]
    if len(found) >= 2:
        raise ValueError('表头需要包含：学校、学院、师资URL、教师URL；支持这些字段的常用别名')
    return None


def paired_sheets(datasets, path):
    """Join source and teacher sheets by school/college, never by row position."""
    rows, errors, warnings = [], [], []
    roles = set()
    sources, teachers = set(), set()
    sheets = []
    for sheet, records in datasets:
        if not records:
            continue
        try:
            fields = header_fields(records[0][1])
            if set(fields) not in ({0, 1, 2}, {0, 1, 3}, {3}):
                raise ValueError('双表文件须为师资表和个人URL表；仅单学院文件的个人URL表可省略学校、学院；不要与四列表混放')
        except ValueError as error:
            errors.append(f'{path.name} / {sheet}：{error}')
            continue
        sheets.append((sheet, records, fields))
    # Read source identities first, even when the teacher sheet appears first.
    for sheet, records, fields in sorted(sheets, key=lambda item: 2 not in item[2]):
        field = 2 if 2 in fields else 3
        role = 'source' if field == 2 else 'teacher'
        roles.add(role)
        for number, values, cells in records[1:]:
            location = f'{path.name} / {sheet} / 第{number}行'
            try:
                if values == records[0][1]:
                    continue
                if max(fields.values()) >= len(values):
                    raise ValueError('缺少字段，请保留学校、学院和链接三列')
                selected = (3,) if set(fields) == {3} else (0, 1, field)
                decoded = [excel_cell(cells[fields[i]], i >= 2) if cells else
                           (str(values[fields[i]] or '').strip(), False) for i in selected]
                if selected == (3,):
                    if len(sources) != 1:
                        raise ValueError('个人URL表未注明学校和学院；只有师资表对应唯一学院时才能自动关联，请补齐学校、学院列')
                    school, college = next(iter(sources))
                    value = decoded[0][0]
                else:
                    school, college, value = [v for v, _ in decoded]
                if not school or not college:
                    raise ValueError('学校和学院不能为空')
                if any(mismatch for _, mismatch in decoded):
                    raise ValueError('显示的URL与实际超链接目标不同，请确认后统一')
                url = absolute_url(value)
                if role == 'source' and not url:
                    raise ValueError('师资URL须为完整的http://或https://地址')
                if role == 'teacher' and value.startswith(('http:', 'https:')) and not url:
                    raise ValueError('教师URL含空格或格式错误，请检查')
                pair = (school, college)
                (sources if role == 'source' else teachers).add(pair)
                rows.append({'school': school, 'college': college, 'roster_url': url if role == 'source' else '',
                             'file': path.name, 'sheet': sheet, 'row': number, 'input_mode': 'college_union',
                             'record_kind': role, 'value': value if role == 'teacher' else '',
                             'url': url if role == 'teacher' else None, 'hyperlink_mismatch': False,
                             'values': [school, college, value if role == 'source' else '', value if role == 'teacher' else '']})
            except ValueError as error:
                errors.append(location + '：' + str(error))
    if roles != {'source', 'teacher'}:
        errors.append(path.name + '：双表文件必须同时包含师资URL表和教师个人URL表（个人URL表可只有表头）')
    for school, college in sorted(teachers - sources):
        errors.append(f'{path.name}：{school} / {college} 有教师链接，却没有对应师资入口；请统一两张表的学校和学院名称')
    for school, college in sorted(sources - teachers):
        warnings.append(f'{school} / {college}：个人URL表为空，将按所给师资入口核对缺少链接')
    return rows, errors, warnings


def parse_file(path):
    errors, warnings, parsed = [], [], []
    if path.suffix.lower() not in SUPPORTED:
        return [], [f'{path.name}：不支持此格式'], []
    if path.stat().st_size > MAX_FILE_BYTES:
        return [], [f'{path.name}：单个文件不能超过10MB'], []
    datasets = []
    book = None
    try:
        ext = path.suffix.lower()
        if ext in ('.xlsx', '.xlsm'):
            with zipfile.ZipFile(path) as z:
                if sum(i.file_size for i in z.infolist()) > 100 * 1024 * 1024:
                    raise ValueError('工作簿解压后的内容过大，请拆分文件')
            book = load_workbook(path, read_only=False, data_only=False, keep_links=False)
            for ws in book.worksheets:
                if ws.max_row > MAX_ROWS + 100:
                    raise ValueError('单个工作表超过50000行，请拆分')
                if ws.max_row * ws.max_column > 2000000:
                    raise ValueError('工作表范围过大，请去除多余的空行空列后重试')
                merged = {}
                for area in ws.merged_cells.ranges:
                    if (area.max_col - area.min_col + 1) * (area.max_row - area.min_row + 1) > MAX_ROWS:
                        raise ValueError('合并区域过大，请取消合并后导入')
                    for row in range(area.min_row, area.max_row + 1):
                        for col in range(area.min_col, area.max_col + 1):
                            merged[(row, col)] = ws.cell(area.min_row, area.min_col)
                records = []
                for row in ws.iter_rows():
                    cells = [merged.get((c.row, c.column), c) if hasattr(c, 'column') else c for c in row]
                    values = [str(c.value or '').strip() for c in cells]
                    if any(values):
                        records.append((row[0].row, values, cells))
                datasets.append((ws.title, records))
        elif ext == '.xls':
            import xlrd
            xls = xlrd.open_workbook(str(path), formatting_info=True)
            for ws in xls.sheets():
                if ws.nrows > MAX_ROWS + 100:
                    raise ValueError('单个工作表超过50000行，请拆分')
                merged = {}
                for rlo, rhi, clo, chi in ws.merged_cells:
                    if (rhi-rlo)*(chi-clo) > MAX_ROWS:
                        raise ValueError('合并区域过大')
                    for r in range(rlo, rhi):
                        for c in range(clo, chi):
                            merged[(r, c)] = ws.cell_value(rlo, clo)
                records = []
                for r in range(ws.nrows):
                    values = [str(merged.get((r, c), ws.cell_value(r, c)) or '').strip() for c in range(ws.ncols)]
                    if any(values):
                        cells = []
                        for c, value in enumerate(values):
                            hyperlink = ws.hyperlink_map.get((r, c))
                            target = getattr(hyperlink, 'url_or_path', None)
                            if target and isinstance(target, str) and getattr(hyperlink, 'textmark', None):
                                target += '#' + hyperlink.textmark
                            cells.append(SimpleNamespace(value=value, data_type='s', hyperlink=SimpleNamespace(target=target) if target else None))
                        records.append((r + 1, values, cells))
                datasets.append((ws.name, records))
        else:
            datasets = [('文本', text_rows(decode_text(path.read_bytes()), ext))]
        if ext in ('.xlsx', '.xlsm', '.xls') and any(
                records and set(header_fields(records[0][1])) in ({0, 1, 2}, {0, 1, 3})
                for _, records in datasets):
            return paired_sheets(datasets, path)
        for sheet, records in datasets:
            columns, first, has_header = None, True, False
            for number, values, cells in records:
                location = f'{path.name} / {sheet} / 第{number}行'
                try:
                    recognized = {i for value in values for i, aliases in enumerate(ALIASES) if key(value) in aliases}
                    mapping = header_map(values) if first or len(recognized) == 4 else None
                    if mapping is not None:
                        columns, first, has_header = mapping, False, True
                        continue
                    if first:
                        columns, first = [0, 1, 2, 3], False
                    if not has_header and len(values) != 4:
                        raise ValueError(f'应有4列，当前{len(values)}列；含空格的名称请使用制表符或CSV分隔')
                    if max(columns) >= len(values):
                        raise ValueError('缺少字段；教师URL可以为空，但需要保留第四列')
                    row, mismatch = [], False
                    for field, col in enumerate(columns):
                        value, different = excel_cell(cells[col], field >= 2) if cells else (plain_link(values[col]) if field >= 2 else values[col].strip(), False)
                        row.append(value)
                        mismatch |= different
                    school, college, roster, teacher = row
                    if not school or not college:
                        raise ValueError('学校和学院不能为空')
                    roster = absolute_url(roster)
                    if not roster:
                        raise ValueError('师资URL须为完整的http://或https://地址')
                    if mismatch:
                        raise ValueError('显示的URL与实际超链接目标不同，请确认后统一')
                    teacher_url = absolute_url(teacher)
                    if teacher.startswith(('http:', 'https:')) and not teacher_url:
                        raise ValueError('教师URL含空格或格式错误，请检查')
                    parsed.append({'school': school, 'college': college, 'roster_url': roster,
                                   'file': path.name, 'sheet': sheet, 'row': number,
                                   'values': [school, college, roster, teacher], 'value': teacher,
                                   'url': teacher_url, 'hyperlink_mismatch': False})
                except ValueError as e:
                    errors.append(location + '：' + str(e))
        if not parsed and not errors:
            errors.append(path.name + '：没有数据行（只有表头或空表）')
    except Exception as e:
        errors.append(path.name + '：' + (str(e) if isinstance(e, (ValueError, csv.Error)) else '无法读取文件，请检查格式或另存后重试（' + type(e).__name__ + '）'))
    finally:
        if book:
            book.close()
    return parsed, errors, warnings


def read_directory(directory):
    rows, files, errors, warnings = [], [], [], []
    for path in sorted(directory.iterdir()):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED or path.name.startswith('~$'):
            continue
        found, issues, notes = parse_file(path)
        rows.extend(found); errors.extend(issues); warnings.extend(notes)
        raw = path.read_bytes()
        files.append({'name': path.name, 'sha256': hashlib.sha256(raw).hexdigest(), 'bytes': len(raw)})
    if not files:
        errors.append('请粘贴四列数据，或导入Excel、CSV、TSV、TXT文件')
    if len(rows) > MAX_ROWS:
        errors.append('本批超过50000行，请拆分')
    return rows, files, errors, warnings


def tasks_from_rows(rows):
    tasks = {}
    modes = {}
    for row in rows:
        pair = (row['school'], row['college'])
        mode = row.get('input_mode', 'per_source')
        if pair in modes and modes[pair] != mode:
            raise ValueError('同一学院不能混用四列表和双表：' + ' / '.join(pair) + '；请统一格式后提交')
        modes[pair] = mode
        group = (*pair, 'college_union' if mode == 'college_union' else row['roster_url'])
        task = tasks.setdefault(group, {'school': pair[0], 'college': pair[1], 'roster_url': '',
                                       'input_mode': mode, 'roster_urls': [], 'source_rows': [], 'rows': []})
        if mode == 'per_source' or row.get('record_kind') == 'source':
            if row['roster_url'] not in task['roster_urls']:
                task['roster_urls'].append(row['roster_url'])
            if row.get('record_kind') == 'source':
                task['source_rows'].append(row)
        if row.get('record_kind') != 'source':
            task['rows'].append(row)
    for group, task in tasks.items():
        if not task['roster_urls']:
            raise ValueError('学院缺少师资入口：' + task['school'] + ' / ' + task['college'])
        task['roster_url'] = task['roster_urls'][0]  # Legacy display field; full scope is roster_urls.
        identity = list(group) + (sorted(task['roster_urls']) if task['input_mode'] == 'college_union' else [])
        task['id'] = hashlib.sha256('\n'.join(identity).encode()).hexdigest()[:12]
    return list(tasks.values())
