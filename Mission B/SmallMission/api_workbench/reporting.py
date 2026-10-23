"""Static, portable reports; no server or main-workflow database required."""
from html import escape
from datetime import datetime
import hashlib
from pathlib import Path
import re
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SUMMARY_HEADERS = ['大学名', '学院名', '核验结果', '特殊情况说明']
SUMMARY_SCHEMA = 'verification_v2'


def sheet(book, title, headers, rows):
    ws = book.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='174F48')
        cell.font = Font(name='Arial', color='FFFFFF', bold=True, size=11)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 34
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name='Arial', size=11)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value, str):
                # Treat all website/input content as text, never Excel formulas.
                cell.data_type = 's'
                if cell.value.startswith(('https://', 'http://')) and '\n' not in cell.value:
                    cell.hyperlink = cell.value
                    cell.font = Font(name='Arial', size=11, color='146B94', underline='single')
            if cell.row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F0F6F4')
    for col, header in enumerate(headers, 1):
        width = 45 if any(w in header for w in ('URL', '地址', '原因', '说明', '证据', '备注')) else 20
        ws.column_dimensions[ws.cell(1, col).column_letter].width = width
    return ws


def new_book():
    book = Workbook()
    book.remove(book.active)
    return book


def safe_name(text):
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', text).rstrip(' .')


def input_assessment(result):
    """Describe the original input independently of successful automatic repair."""
    if not result['complete']:
        return {'title': '尚未判定', 'tone': 'pending', 'detail': '未取得可靠名单，不能判断缺少或多出；原数据保留。'}
    stats = result['stats']
    main = []
    if stats.get('missing'):
        main.append(f"缺少 {stats['missing']} 条")
    if stats.get('extra'):
        main.append(f"多出 {stats['extra']} 条")
    other = []
    for key, label, unit in [('input_duplicates', '重复', '行'), ('protocol_changes', '协议写法差异', '条'), ('input_placeholders', '非URL标记', '行')]:
        if stats.get(key):
            other.append(f'{label} {stats[key]} {unit}')
    if main:
        return {'title': '，'.join(main), 'tone': 'pending', 'detail': '已按当前师资页自动修正。' + ('另有：' + '，'.join(other) + '。' if other else '')}
    if other:
        return {'title': '无缺少或多出；有其他差异', 'tone': 'changed', 'detail': '，'.join(other) + '；已自动整理。'}
    return {'title': '无链接差异', 'tone': 'good', 'detail': '原表与当前师资页吻合，无需修正。'}


def verification_summary(result):
    """Present the existing URL-set decision; missing evidence stays undecided."""
    status = '特殊情况'
    if result.get('complete') is True:
        if 'source' in result and 'diffs' in result:
            def comparison_key(url):
                p = urlsplit(url)
                return p.netloc.lower(), p.path or '/', p.query, p.fragment
            expected = {comparison_key(url) for url in result['source']['links']}
            supplied = {comparison_key(row['url']) for row in result['diffs'] if row.get('url')}
            status = '错误' if supplied != expected else '正确'
        elif result.get('verification_result') in ('正确', '错误'):
            # A persisted app brief already carries the full result's decision.
            status = result['verification_result']
        elif result.get('original_error') in ('Yes', 'No'):
            # Read old app briefs without rewriting their files or downloads.
            status = {'Yes': '错误', 'No': '正确'}[result['original_error']]
        else:
            missing, extra = (result.get('stats', {}).get(k) for k in ('missing', 'extra'))
            if any(type(n) is int and n > 0 for n in (missing, extra)):
                status = '错误'
            elif missing == 0 and extra == 0:
                status = '正确'
    note = ''
    if status == '特殊情况':
        note = result.get('special_note') or result.get('reason') or ''
        if not note:
            note = '；'.join(dict.fromkeys(issue for source in result.get('sources', [])
                                         for issue in source.get('issues', []) if issue))
        if not note and any(s.get('failure_kind') == 'access' for s in result.get('sources', [])):
            note = '学院网页无法访问'
    return {'verification_result': status, 'special_note': note.strip()}


def college_summaries(results):
    groups = {}
    for result in results:
        groups.setdefault((result['school'], result['college']), []).append(result)
    summaries = []
    for (school, college), members in groups.items():
        assessments = [verification_summary(r) for r in members]
        statuses = [a['verification_result'] for a in assessments]
        # A confirmed difference proves an original-input error even if a different
        # independently paired source is unavailable. Correct needs all scopes verified.
        status = next((s for s in ('错误', '特殊情况') if s in statuses), '正确')
        summaries.append({'school': school, 'college': college, 'verification_result': status,
                          'special_note': '；'.join(dict.fromkeys(a['special_note'] for a in assessments if a['special_note'])),
                          'result_ids': [r['id'] for r in members if 'id' in r]})
    return summaries


def college_summary_rows(results):
    return [[r['school'], r['college'], r['verification_result'], r['special_note']]
            for r in college_summaries(results)]


def roster_label(result):
    return '\n'.join(result.get('roster_urls', [result['roster_url']]))


def export_names(results, manifest):
    """Readable school/college names plus a run token prevent download collisions."""
    stamp = manifest.get('finished_at', '')
    try:
        stamp = datetime.fromisoformat(stamp.replace('Z', '+00:00')).astimezone().strftime('%Y%m%d_%H%M%S')
    except ValueError:
        stamp = '核验'
    identity = manifest.get('run_id') or hashlib.sha256((str(manifest.get('finished_at')) + '|'.join(r['id'] for r in results)).encode()).hexdigest()
    tag = stamp + '_' + identity[:6]
    colleges = list(dict.fromkeys((r['school'], r['college']) for r in results))
    school, college = colleges[0]
    base = safe_name(school + '_' + college)[:80]
    if len(colleges) > 1:
        base += f'_等{len(colleges)}学院'
    report = f'{base}_核验结果_{tag}.xlsx'
    corrected = {}
    stems = [safe_name(r['school'] + '_' + r['college'])[:80] for r in results]
    for result, stem in zip(results, stems):
        suffix = '_' + result['id'][:12] if stems.count(stem) > 1 else ''
        corrected[result['id']] = f'{stem}{suffix}_修正后链接_{tag}.xlsx'
    return report, corrected


def stats_row(r):
    s = r['stats']
    return [r['school'], r['college'], roster_label(r), '是' if r['needs_human'] else '否',
            r['verdict'], s['input_rows'], s['input_url_rows'], s['input_unique'],
            s['input_placeholders'], s['input_duplicates'], s['source_occurrences'],
            s['source_unique'], s['source_placeholders'], s['matched'], s['missing'],
            s['extra'], s['protocol_changes'], s['literal_missing'], s['literal_extra'],
            s['corrected'], r['reason'], r['attempts'][-1].get('captured_at', '')]


def export(out, results, manifest):
    filename, names = export_names(results, manifest)
    manifest['report_filename'] = filename
    summary_name = filename.replace('_核验结果_', '_核验汇总_')
    manifest['summary_filename'] = summary_name
    manifest['summary_schema'] = SUMMARY_SCHEMA
    summary_rows = college_summary_rows(results)
    summary = new_book()
    sheet(summary, '核验汇总', SUMMARY_HEADERS, summary_rows)
    summary.save(out / summary_name)
    check_summary = load_workbook(out / summary_name, read_only=True, data_only=True)
    assert check_summary.sheetnames == ['核验汇总']
    assert [[value if value is not None else '' for value in row] for row in check_summary.active.values] == [SUMMARY_HEADERS] + summary_rows
    check_summary.close()
    for result in results:
        result['corrected_filename'] = names[result['id']]
    book = new_book()
    sheet(book, '核验汇总', SUMMARY_HEADERS, summary_rows)
    sheet(book, '学院结论', ['学校', '学院', '师资URL', '是否需要人工介入', '结论', '输入行数',
                           '输入URL行数', '输入不同URL数', '输入非URL标记数', '输入重复行数',
                           '官网链接出现次数', '官网不同URL数', '官网占位及空锚点数', '完全匹配URL数',
                           '缺少URL数（排除协议差异）', '范围外URL数（排除协议差异）', '协议写法修正数',
                           '严格地址差集_官网多于输入', '严格地址差集_输入多于官网',
                           '修正后URL数', '说明', '取证时间UTC'], [stats_row(r) for r in results])
    diffs, anchors, corrected, pending, notices, mappings = [], [], [], [], [], []
    (out / 'corrected').mkdir()
    for r in results:
        for d in r['diffs']:
            diffs.append([r['school'], r['college'], '\n'.join(d.get('source_urls', [])) or roster_label(r), d['file'], d['sheet'], d['row'],
                          d['name'], d['value'], d['action'], d['new_url'], d['reason']])
            if not r['complete']:
                pending.append([r['school'], r['college'], roster_label(r), d['value'],
                                d['file'], d['row'], r['reason']])
        for a in r['source']['anchors']:
            anchors.append([r['school'], r['college'], a.get('source_url', r['roster_url']), a['root_index'], a['name'], a['href'],
                            a['url'], '是' if a['included'] else '否', a['reason'],
                            a.get('html_path', r['attempts'][-1].get('html_path', ''))])
        for note in r['source']['notes']:
            notices.append([r['school'], r['college'], note['kind'], note['name'], note['url'], note['message']])
        if r['complete']:
            rows = []
            for url in r['corrected_urls']:
                sources = r['source']['links'][url].get('sources', [r['roster_url']])
                rows.append([r['school'], r['college'], sources[0], url])
                mappings.extend([r['school'], r['college'], source, url] for source in sources)
            corrected.extend(rows)
            single = new_book()
            sheet(single, '核验后链接', ['学校', '学院', '师资url', '教师url'], rows)
            single.save(out / 'corrected' / r['corrected_filename'])
    sheet(book, '逐项对账', ['学校', '学院', '师资URL', '原文件', '原Sheet', '原行号', '名单文字',
                          '原始教师URL或标记', '处理', '修正后URL', '原因'], diffs)
    sheet(book, '官网链接证据', ['学校', '学院', '师资URL', '名单分组序号', '链接文字', '原始href',
                               '解析后URL', '是否计入', '说明', '页面证据路径'], anchors)
    sheet(book, '修正后四列表', ['学校', '学院', '师资url', '教师url'], corrected)
    sheet(book, '教师链接来源映射', ['学校', '学院', '师资url', '教师url'], mappings)
    sheet(book, '待人工处理原始行', ['学校', '学院', '师资URL', '原始教师URL或标记', '原文件', '原行号', '待处理原因'], pending)
    sheet(book, '官网提示', ['学校', '学院', '提示类型', '链接文字', '源URL', '说明'], notices)
    instructions = [
        ['任务范围', manifest['scope']],
        ['判定对象', '链接集合；不是教师人数，不做同名身份合并，不核对邮箱、职称或个人页正文'],
        ['访问状态', '不访问教师链接，不因404、超时等移出链接；师资来源页必须取得后才能对账'],
        ['协议差异', 'http/https只作显式源地址替换，不假定两个页面相同；严格地址差集另列'],
        ['占位', '空href或单独#不提供目标；无文字也无图像的空锚点不计为可点击教师链接，证据均保留'],
        ['官网异常', '有明确HTTP目标的姓名链接按源保留，包括疑似首页；同一地址的多姓名只作提示'],
        ['范围限制', '读取给定师资页及同栏目的可识别分页；按学院合并去重，不扩展其他栏目。任一来源或分页读取不全时不自动删改原表'],
        ['简表口径', '核验结果仅为正确、错误、特殊情况。全部指定来源完成核验且URL集合符合现有规则为正确；已确认缺少或多出为错误，自动修正后仍为错误；无法可靠判定为特殊情况。特殊情况说明可以为空，原因不明也不能判为正确或错误。重复本身及仅http/https写法差异不单独判错。同学院独立核验分组已确认差异时汇总为错误，其他分组未完成的原因仍保留在说明列。'],
        ['人工介入', '配置已核实的名单区域可完成链接对账和自动修正则为否；来源取不到或规则不适用则为是'],
        ['待处理', '未完成学院不生成空的修正表，不判断多采或少采；原文件副本在inputs中保留'],
        ['历史准确性', '当前差异以本次快照为准，不据此断言原始采集当时有错误'],
        ['规则准备', manifest['configuration_note']],
        ['执行模式', manifest['mode']],
        ['模型API调用次数', manifest['model_api_calls']],
        ['个人主页显式请求次数', manifest['explicit_profile_fetches']],
        ['完成时间UTC', manifest['finished_at']],
        ['详细证据', 'results.json、manifest.json、source_plans.json、evidence/'],
    ]
    sheet(book, '口径说明', ['项目', '说明'], instructions)
    path = out / filename
    book.save(path)
    # Verify serialized output, not just in-memory counts.
    check = load_workbook(path, read_only=True, data_only=False)
    assert check['学院结论'].max_row - 1 == len(results)
    assert check['修正后四列表'].max_row - 1 == len(corrected)
    assert check['待人工处理原始行'].max_row - 1 == len(pending)
    assert all(c.data_type != 'f' for ws in check for row in ws for c in row)
    check.close()
    for r in results:
        if not r['complete']:
            continue
        p = out / 'corrected' / r['corrected_filename']
        b = load_workbook(p, read_only=True, data_only=True)
        values = [row[3] for row in b.active.iter_rows(min_row=2, values_only=True)]
        assert len(values) == len(set(values)) and set(values) == set(r['source']['links'])
        b.close()
    write_html(out, results, manifest)


def link(url, label=None):
    if not url:
        return ''
    return '<a href="' + escape(url, quote=True) + '" target="_blank" rel="noopener noreferrer">' + escape(label or url) + '</a>'


def html_table(headers, rows):
    return '<div class="scroll"><table><thead><tr>' + ''.join('<th>' + escape(h) + '</th>' for h in headers) + '</tr></thead><tbody>' + ''.join('<tr>' + ''.join('<td>' + str(v if v is not None else '未判定') + '</td>' for v in row) + '</tr>' for row in rows) + '</tbody></table></div>'


def write_html(out, results, manifest):
    summary_rows, details = [], []
    for r in results:
        s = r['stats']
        summary_rows.append([escape(r['school']), '<a href="#' + r['id'] + '">' + escape(r['college']) + '</a>',
                             escape(input_assessment(r)['title']),
                             '<span class="' + ('pending' if r['needs_human'] else 'ok') + '">' + escape(r['verdict']) + '</span>',
                             s['input_rows'], s['source_unique'], s['missing'], s['extra'], s['protocol_changes']])
        changes = [d for d in r['diffs'] if d['action'] not in ('保留', '未判定，原样保留')]
        details.append('<section id="' + r['id'] + '"><h2>' + escape(r['school'] + ' · ' + r['college']) + '</h2><p>' + escape(r['reason']) + '</p><p>师资页：' + '<br>'.join(link(u) for u in r.get('roster_urls', [r['roster_url']])) + '</p>')
        if r['complete']:
            download = 'corrected/' + r['corrected_filename']
            details.append('<p>' + link(download, '下载本学院修正表') + ' · ' + str(s['source_occurrences']) + ' 次链接展示 / ' + str(s['source_unique']) + ' 个不同 URL</p>')
            if changes:
                details.append(html_table(['处理', '名单文字', '原地址', '新地址'], [
                    [escape(d['action']), escape(d['name']), link(d['value']) if d['url'] else escape(d['value']), link(d['new_url'])]
                    for d in changes]))
            else:
                details.append('<p class="ok">逐项吻合，无需修正。</p>')
        else:
            details.append('<p class="pending">待办：取得该师资页的正常名单内容，核实名单区域后重跑。本学院未自动增删，原表副本已保留。</p>')
        if r['source']['notes']:
            details.append('<details><summary>官网提示（不阻断链接对账）</summary>' + html_table(['提示', '名单文字', 'URL'], [
                [escape(n['kind']), escape(n['name']), link(n['url'])] for n in r['source']['notes']]) + '</details>')
        details.append('<details><summary>取证记录</summary>' + html_table(['方式', '来源状态', '取证时间UTC', '保存页面'], [
            [escape(a['method']), a.get('status', '失败'), escape(a.get('captured_at', '')), link(a.get('html_path'), '查看快照')]
            for a in r['attempts']]) + '</details></section>')
    style = '''body{margin:0;background:#f4f7f5;color:#183c37;font:15px/1.6 Arial,"Microsoft YaHei",sans-serif}main{max-width:1400px;margin:36px auto;padding:0 28px}h1{font-size:30px;margin-bottom:8px}h2{font-size:20px}p{max-width:1100px}a{color:#006e88;word-break:break-all}section,.card{padding:24px;background:white;border:1px solid #d6e4df;border-radius:12px;margin:22px 0}.stats{display:flex;gap:24px;flex-wrap:wrap}.stat{font-size:30px;font-weight:bold}.ok{color:#137458}.pending{color:#aa6017}table{width:100%;border-collapse:collapse}th,td{padding:12px;text-align:left;border-bottom:1px solid #e2ebe7;vertical-align:top}th{background:#eaf1ed;white-space:nowrap}td{max-width:420px;overflow-wrap:anywhere}.scroll{overflow:auto}summary{cursor:pointer;margin:16px 0}small{color:#59756f}input{padding:10px;border:1px solid #a9c8bb;border-radius:5px;width:280px}'''
    content = '<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Mission B · 师资链接核验</title><style>' + style + '</style><main>'
    content += '<h1>Mission B · 师资链接核验</h1><p>' + escape(manifest['scope']) + '。只核对源页面提供的地址，不检查教师页面是否能打开。</p>'
    content += '<p>' + link(manifest['summary_filename'], '下载核验汇总（四列）') + ' · ' + link(manifest['report_filename'], '下载详细核验 Excel') + ' · ' + link('manifest.json', '运行记录') + '</p>'
    colleges = college_summary_rows(results)
    content += '<div class="stats">' + ''.join('<div class="card"><div class="stat">' + str(sum(r[2] == status for r in colleges)) + '</div>' + status + '</div>' for status in ('正确', '错误', '特殊情况')) + '</div>'
    content += '<div id="verification-summary">' + html_table(SUMMARY_HEADERS, [[escape(v) for v in row] for row in colleges]) + '</div>'
    content += '<p><small>“未判定”不代表零。协议写法差异单列；未取得名单的学院不自动增删。结论以当前快照为准，不代表历史采集当时的准确度。</small></p>'
    content += '<h2>逐组对账详情</h2>' + html_table(['学校', '学院', '原始输入差异', '处理后结论', '输入行数', '官网不同URL', '缺少', '范围外', '协议修正'], summary_rows)
    content += ''.join(details) + '<p><small>模型 API 调用：' + str(manifest['model_api_calls']) + '；个人主页显式访问：' + str(manifest['explicit_profile_fetches']) + '。</small></p></main></html>'
    (out / 'report.html').write_text(content, encoding='utf-8')
