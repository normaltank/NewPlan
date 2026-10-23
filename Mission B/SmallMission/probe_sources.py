"""Capture only supplied roster pages for the isolated Mission B assessment."""
import asyncio
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

import httpx
from bs4 import BeautifulSoup
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent

def inputs():
    result = []
    for path in sorted(ROOT.parent.glob('*.xlsx')):
        wb = load_workbook(path, read_only=True, data_only=True)
        rows = [list(row[:4]) for row in wb.worksheets[0].iter_rows(min_row=2, values_only=True) if any(row)]
        wb.close()
        result.append({'file': path.name, 'school': rows[0][0], 'college': rows[0][1],
                       'url': rows[0][2], 'rows': rows})
    return result

async def main():
    out = ROOT / 'evidence' / 'initial'
    out.mkdir(parents=True, exist_ok=True)
    targets = inputs()
    limit = asyncio.Semaphore(4)
    async with httpx.AsyncClient(follow_redirects=True, timeout=30,
                                headers={'User-Agent': 'Mozilla/5.0 (compatible; FacultyLinkAudit/1.0)'}) as client:
        async def capture(t):
            key = hashlib.sha256(t['url'].encode()).hexdigest()[:16]
            meta = {k: v for k, v in t.items() if k != 'rows'}
            meta.update(key=key, captured_at=datetime.now(timezone.utc).isoformat())
            try:
                async with limit:
                    response = await client.get(t['url'])
                raw = response.content
                (out / (key + '.bin')).write_bytes(raw)
                soup = BeautifulSoup(raw, 'lxml')
                html = str(soup)
                (out / (key + '.html')).write_text(html, encoding='utf-8')
                meta.update(status=response.status_code, final_url=str(response.url),
                            bytes=len(raw), sha256=hashlib.sha256(raw).hexdigest(),
                            title=soup.title.get_text(' ', strip=True) if soup.title else '')
                links = []
                for a in soup.select('a[href]'):
                    parent = a.parent
                    links.append({'text': a.get_text(' ', strip=True), 'href': a['href'],
                                  'title': a.get('title'),
                                  'context': parent.get_text(' ', strip=True)[:180],
                                  'ancestors': [p.name + ('#' + p['id'] if p.get('id') else '') +
                                                ('.' + '.'.join(p['class']) if p.get('class') else '')
                                                for p in list(a.parents)[:6] if p.name]})
                (out / (key + '.links.json')).write_text(json.dumps(links, ensure_ascii=False, indent=2), encoding='utf-8')
            except Exception as e:
                meta['error'] = type(e).__name__ + ': ' + str(e)
            (out / (key + '.json')).write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding='utf-8')
            print(json.dumps(meta, ensure_ascii=False), flush=True)
        await asyncio.gather(*(capture(t) for t in targets))

if __name__ == '__main__':
    asyncio.run(main())
