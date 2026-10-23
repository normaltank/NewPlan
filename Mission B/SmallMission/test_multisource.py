"""Paired-sheet input, explicit multi-page scope, and original-input summary."""
import base64
import io
import json
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

import httpx
from openpyxl import Workbook, load_workbook
from app import make_server
from audit import Fetcher, combine_sources, parse_roster, reconcile
from input_adapter import parse_file, tasks_from_rows
from reporting import college_summary_rows, original_error_status
from runtime_paths import APP_ROOT
from test_audit import SOURCE, PLAN, task


def paired_file(path, school, college, sources, teachers):
    book = Workbook()
    people = book.active
    people.title = '个人url'
    people.append(['学校', '学院', '教师个人url'])
    for url in teachers:
        people.append([school, college, url])
    roster = book.create_sheet('师资url')
    roster.append(['学校', '学院', '师资url'])
    for url in sources:
        roster.append([school, college, url])
    book.save(path)


class InputAndUnionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=APP_ROOT)
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_two_sheets_join_by_college_not_row_position(self):
        path = self.root / '双表.xlsx'
        paired_file(path, '甲大学', '甲学院', [SOURCE, SOURCE + '?page=2', SOURCE], [SOURCE + '?id=1', SOURCE + '?id=2'])
        rows, errors, warnings = parse_file(path)
        self.assertFalse(errors)
        tasks = tasks_from_rows(rows)
        self.assertEqual(len(tasks), 1)
        self.assertEqual(tasks[0]['roster_urls'], [SOURCE, SOURCE + '?page=2'])
        self.assertEqual(len(tasks[0]['rows']), 2)
        self.assertEqual(len(tasks[0]['source_rows']), 3)

    def test_empty_teacher_sheet_still_creates_a_task(self):
        path = self.root / '空链接.xlsx'
        paired_file(path, '甲大学', '甲学院', [SOURCE], [])
        rows, errors, warnings = parse_file(path)
        self.assertFalse(errors)
        self.assertTrue(warnings)
        self.assertEqual(tasks_from_rows(rows)[0]['rows'], [])

    def test_unmatched_college_and_missing_sheet_report_errors(self):
        path = self.root / '不匹配.xlsx'
        paired_file(path, '甲大学', '甲学院', [SOURCE], [SOURCE + '?id=1'])
        book = load_workbook(path)
        book['个人url']['B2'] = '乙学院'
        book.save(path)
        self.assertTrue(any('没有对应师资入口' in error for error in parse_file(path)[1]))
        book.remove(book['师资url'])
        book.save(path)
        self.assertTrue(any('同时包含' in error for error in parse_file(path)[1]))

    def test_reordered_headers_aliases_and_hyperlink_targets(self):
        path = self.root / '其他表名.xlsx'
        book = Workbook()
        book.active.title = '已有教师'
        book.active.append(['个人主页', '大学名', '学院名'])
        book.active.append(['张老师', '甲大学', '甲学院'])
        book.active['A2'].hyperlink = SOURCE + '?id=1'
        roster = book.create_sheet('来源列表')
        roster.append(['师资链接', '大学名', '学院名'])
        roster.append([SOURCE, '甲大学', '甲学院'])
        book.save(path)
        rows, errors, _ = parse_file(path)
        self.assertFalse(errors)
        self.assertEqual(tasks_from_rows(rows)[0]['rows'][0]['url'], SOURCE + '?id=1')

    def test_mixed_semantics_for_same_college_are_rejected(self):
        path = self.root / '双表.xlsx'
        paired_file(path, '甲大学', '甲学院', [SOURCE], [SOURCE + '?id=1'])
        rows, _, _ = parse_file(path)
        rows.append({'school': '甲大学', 'college': '甲学院', 'roster_url': SOURCE, 'url': SOURCE})
        with self.assertRaisesRegex(ValueError, '不能混用'):
            tasks_from_rows(rows)

    def test_explicit_pagination_is_in_scope_but_unlisted_page_is_not(self):
        body = '<title>教师</title><div class="roster"><a href="/a">甲</a></div><a href="?page=2">下一页</a>'
        self.assertFalse(parse_roster(body, SOURCE, PLAN)['complete'])
        self.assertTrue(parse_roster(body, SOURCE, PLAN, [SOURCE, SOURCE + '?page=2'])['complete'])

    def test_union_deduplicates_sources_without_inventing_provenance(self):
        sources = []
        for index, targets in enumerate((['a', 'b'], ['b', 'c']), 1):
            url = SOURCE + f'?page={index}'
            body = '<title>教师</title><div class="roster">' + ''.join(f'<a href="/{v}">{v}</a>' for v in targets) + '</div>'
            sources.append({'roster_url': url, 'parsed': parse_roster(body, url, PLAN),
                            'attempts': [{'html_path': f'page{index}.html'}]})
        combined = combine_sources(sources)
        self.assertEqual(len(combined['links']), 3)
        self.assertEqual(combined['links']['https://example.edu.cn/b']['occurrences'], 2)
        self.assertEqual(len(combined['links']['https://example.edu.cn/b']['sources']), 2)
        self.assertEqual(combined['anchors'][0]['html_path'], 'page1.html')
        result = reconcile(task('https://example.edu.cn/a', 'https://example.edu.cn/b', 'https://example.edu.cn/c'), combined)
        self.assertEqual(original_error_status(result), 'No')
        self.assertEqual(result['stats']['input_duplicates'], 0)

    def test_one_failed_source_prevents_deleting_other_page_input(self):
        good = parse_roster('<title>教师</title><div class="roster"><a href="/a">甲</a></div>', SOURCE, PLAN)
        failed = parse_roster('<p>failure</p>', SOURCE + '?page=2', PLAN)
        combined = combine_sources([{'roster_url': SOURCE, 'parsed': good, 'attempts': [{}]},
                                    {'roster_url': SOURCE + '?page=2', 'parsed': failed, 'attempts': [{}]}])
        result = reconcile(task('https://example.edu.cn/a', 'https://example.edu.cn/b'), combined)
        self.assertFalse(result['complete'])
        self.assertIsNone(result['stats']['extra'])
        self.assertEqual([d['new_url'] for d in result['diffs']], ['https://example.edu.cn/a', 'https://example.edu.cn/b'])

    def test_summary_error_status_ignores_auto_repair_and_access_is_not_yes(self):
        self.assertEqual(original_error_status({'complete': True, 'stats': {'missing': 1}, 'needs_human': False}), 'Yes')
        self.assertEqual(original_error_status({'complete': True, 'stats': {}, 'needs_human': False}), 'No')
        self.assertEqual(original_error_status({'complete': False, 'sources': [{'failure_kind': 'access'}]}), '学院网页无法访问')
        self.assertEqual(original_error_status({'complete': False, 'sources': [{'failure_kind': 'structure'}]}), '未完成核验')
        results = [{'school': '甲大学', 'college': '甲学院', 'complete': True, 'stats': {}}] * 2
        self.assertEqual(college_summary_rows(results), [['甲大学', '甲学院', 'No']])

    def test_summary_deduplicates_before_comparing_sets_and_ignores_protocol(self):
        parsed = parse_roster('<title>教师</title><div class="roster"><a href="/a">甲</a><a href="/b">乙</a></div>', SOURCE, PLAN)
        def status(*urls):
            result = reconcile(task(*urls), parsed)
            return original_error_status({**result, 'source': parsed})
        self.assertEqual(status('https://example.edu.cn/a', 'https://example.edu.cn/a'), 'Yes')
        self.assertEqual(status('https://example.edu.cn/a', 'https://example.edu.cn/b', 'https://example.edu.cn/b'), 'No')
        self.assertEqual(status('http://example.edu.cn/a', 'https://example.edu.cn/b'), 'No')
        self.assertEqual(status('http://example.edu.cn/a', 'https://example.edu.cn/a', 'https://example.edu.cn/b'), 'No')
        self.assertEqual(status('https://example.edu.cn/a', 'https://example.edu.cn/c'), 'Yes')
        self.assertEqual(status('https://example.edu.cn/a?person=2', 'https://example.edu.cn/b'), 'Yes')


class BatchFlowTests(unittest.TestCase):
    def test_four_uploaded_colleges_produce_exactly_one_three_column_summary(self):
        with tempfile.TemporaryDirectory(dir=APP_ROOT) as temp:
            root = Path(temp)
            server = make_server(root / 'data')
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                files, plans = [], {}
                for index in range(4):
                    sources = [f'https://example.edu.cn/{index}/{page}' for page in (1, 2)]
                    teachers = ['https://example.edu.cn/a', 'https://example.edu.cn/extra' if index == 1 else 'https://example.edu.cn/b']
                    path = root / f'学院{index}.xlsx'
                    paired_file(path, '测试大学', f'学院{index}', sources, teachers)
                    files.append({'name': path.name, 'content': base64.b64encode(path.read_bytes()).decode()})
                    plans.update({u: PLAN for u in sources})
                plans['https://example.edu.cn/3/2'] = {'selector': None}
                config = root / 'plans.json'
                config.write_text(json.dumps({'scope': 'test', 'plans': plans}), encoding='utf-8')
                calls = []
                async def fetch(fetcher, task):
                    url = task['roster_url']
                    calls.append(url)
                    status = 503 if url.endswith('/2/2') else 200
                    targets = ['a', 'b'] if url.endswith('/0/1') else ['a'] if url.endswith('/1') else ['b']
                    body = '<title>教师</title><div class="roster">' + ''.join(f'<a href="/{v}">{v}</a>' for v in targets) + '</div>'
                    meta = {'requested_url': url, 'final_url': url, 'status': status, 'captured_at': 'test', 'method': 'fixture'}
                    fetcher.save(task['id'], 'http', body, meta)
                    return body, meta
                with httpx.Client(base_url=f'http://127.0.0.1:{server.server_port}', trust_env=False) as client:
                    client.post('/api/session', json={'token': server.token}).raise_for_status()
                    with patch('app.plans_path', return_value=config), patch.object(Fetcher, 'http', fetch), patch.object(Fetcher, 'browser_fetch', fetch):
                        preview = client.post('/api/preview', json={'files': files}).json()
                        self.assertFalse(preview['errors'])
                        self.assertEqual((preview['task_count'], preview['source_count'], preview['row_count']), (4, 8, 8))
                        job = client.post('/api/jobs', json={'preview_id': preview['id']}).json()
                        server.store.worker.join(timeout=15)
                    job = client.get('/api/jobs/' + job['id']).json()
                    self.assertEqual(job['status'], 'completed', job)
                    response = client.get('/api/jobs/' + job['id'] + '/download/summary')
                    response.raise_for_status()
                    book = load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
                    self.assertEqual(len(book.sheetnames), 1)
                    self.assertEqual(list(book.active.values), [('大学名', '学院名', '是否出错'),
                        ('测试大学', '学院0', 'No'), ('测试大学', '学院1', 'Yes'),
                        ('测试大学', '学院2', '学院网页无法访问'), ('测试大学', '学院3', '未完成核验')])
                    book.close()
                    self.assertTrue(all('/a' != url.rsplit('/', 1)[-1] and url.count('/') == 4 for url in calls))
                    results = job['results']
                    self.assertEqual(results[0]['stats']['source_unique'], 2)
                    self.assertEqual(results[0]['stats']['source_occurrences'], 3)
                    self.assertIsNone(results[2]['stats']['extra'])
            finally:
                if server.store.worker:
                    server.store.worker.join(timeout=15)
                server.shutdown()
                thread.join()
                server.server_close()


if __name__ == '__main__':
    unittest.main()
