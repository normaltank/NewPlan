"""Compare complete downstream results using the same live roster/raw capture.

The candidate replaces only eligible personal-page text. It is not a Fetch-only
crawler. Missing/denied sources, discovery, identity rules, OCR and PDFs retain
their original semantics. All outputs stay in this experiment.
"""
from __future__ import annotations

import argparse
import asyncio
from collections import Counter
import hashlib
import io
import json
from pathlib import Path
import re
import shutil
import sys
import time
import zipfile

from bs4 import BeautifulSoup
from markdown_it import MarkdownIt

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "runtime"))
from workbench_store import Store
from quality_pipeline import build
from quality_fields import fields_from_text
from workbench_export import tables, university_zip, workbook_bytes, TEACHER_HEADERS, TEACHER_SHEETS
from fetch_client import FetchClient, save

EMAIL = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")


def read(path, default=None):
    try:
        return json.loads(Path(path).read_text(encoding="utf8"))
    except (FileNotFoundError, ValueError):
        return default


def norm(text):
    return re.sub(r"[^\w@]", "", text or "").lower()


def markdown_text(markdown):
    """Parse CommonMark escapes, rather than matching emails in source markup."""
    html = MarkdownIt("commonmark").enable("table").render(markdown)
    soup = BeautifulSoup(html, "html.parser")
    for node in soup.select("script,style"):
        node.decompose()
    links = []
    for node in soup.select("a[href]"):
        href = node.get("href", "")
        links.append(href)
        if href.startswith("mailto:"):
            address = href[7:].split("?", 1)[0]
            if address not in node.get_text():
                node.append(" " + address)
    return soup.get_text("\n", strip=True), links


def retained_fraction(original, candidate):
    original, candidate = norm(original), norm(candidate)
    if not original:
        return None
    if len(original) < 32:
        return float(original in candidate)
    chunks = [original[i:i+32] for i in range(0, len(original)-31, 16)]
    return round(sum(chunk in candidate for chunk in chunks) / len(chunks), 4)


def attachment_suffix(text):
    offsets = [text.find(marker) for marker in ("[图片文字 / OCR", "[个人页 PDF 附件全文]") if marker in text]
    return text[min(offsets):] if offsets else ""


def apply_candidate(root, page, result, names):
    """A Fetch failure never changes the page's origin status or delivery class."""
    folder = root / page["directory"]
    record = {"url": page["url"], "kind": page["kind"], "names": sorted(names),
              "fetch_status": result["status"], "fetch_http_status": result.get("http_status"),
              "call_index": result.get("call_index"), "applied": False,
              "origin_status": page.get("status"), "reason": "fetch_not_successful"}
    if result["status"] != "ok":
        return record
    markdown = result["content"]
    plain, links = markdown_text(markdown)
    (folder / "fetch-original.md").write_text(markdown, encoding="utf8")
    (folder / "fetch-text.txt").write_text(plain, encoding="utf8")
    record.update(fetch_title=result.get("title", ""), fetch_chars=len(plain), fetch_links=links,
                  fetch_sha256=hashlib.sha256(markdown.encode()).hexdigest(),
                  names_present=sorted(name for name in names if norm(name) in norm(plain)))
    old_path = root / (page.get("body_path") or f"{page['directory']}/body-text.txt")
    original = old_path.read_text(encoding="utf8") if old_path.is_file() else ""
    record.update(native_chars=len(original), native_text_retention=retained_fraction(original.split("[图片文字 / OCR",1)[0].split("[个人页 PDF 附件全文]",1)[0], plain),
                  native_emails=sorted(set(EMAIL.findall(original))), fetch_emails=sorted(set(EMAIL.findall(plain))))
    if page["kind"] != "profile":
        record["reason"] = "diagnostic_only_roster_and_attachment_rules_unchanged"
        return record
    if page.get("body_status") in ("permission_denied", "redirect_out_of_scope"):
        record["reason"] = "existing_access_or_scope_restriction"
        return record
    if len(names) != 1 or not any(norm(name) in norm(plain) for name in names):
        record["reason"] = "personal_identity_not_bound_to_fetch_text"
        return record
    text = plain
    suffix = attachment_suffix(original)
    if suffix:
        text += "\n\n" + suffix
    path = folder / "fetch-body.txt"
    path.write_text(text, encoding="utf8")
    page.update(body_path=str(path.relative_to(root)).replace("\\", "/"), body_chars=len(text),
                body_status=page.get("body_status") if suffix else "text_saved",
                email_candidates=sorted(set(EMAIL.findall(text))),
                fetch_provenance={"provider":"deepseek_fetch", "response":f"{page['directory']}/fetch-original.md",
                                  "capture_origin_status":page.get("status"), "fetch_origin_status":None,
                                  "title":result.get("title", ""), "sha256":record["fetch_sha256"]})
    fields = fields_from_text(next(iter(names)), text, page.get("branch", ""))
    page["title_candidates"] = list(dict.fromkeys(value["value"] for value in fields["title_candidates"]))
    record.update(applied=True, reason="fetch_text_with_existing_media_and_same_quality_policy", candidate_chars=len(text))
    return record


def catalog_summary(cat):
    data = tables([cat])
    statuses = Counter(person["review"]["status"] for person in cat["people"])
    return {"observed_names":len({p["name"] for p in cat["people"]}), "person_records":len(cat["people"]),
            "review_statuses":dict(statuses), "delivered":len(data["teachers"]),
            "with_email":sum(bool(row[3]) for row in data["teachers"]),
            "with_title":sum(bool(row[2]) for row in data["teachers"]),
            "with_biography":sum(bool(row[8]) for row in data["teachers"]),
            "unresolved_issues":sum(p["review"]["status"]=="pending" for p in cat["issues"])}


def compare_catalogs(left, right):
    old = {p["id"]:p for p in left["people"]}
    new = {p["id"]:p for p in right["people"]}
    rows = []
    for key in sorted(old.keys() | new.keys()):
        a, b = old.get(key, {}), new.get(key, {})
        ar, br = a.get("review", {}), b.get("review", {})
        af, bf = ar.get("fields", {}), br.get("fields", {})
        changed = [field for field in ("name", "title", "email", "profile_url") if af.get(field) != bf.get(field)]
        if changed or ar.get("status") != br.get("status") or af.get("biography") != bf.get("biography"):
            rows.append({"person_id":key, "name":b.get("name", a.get("name")), "url":b.get("profile_url",a.get("profile_url")),
                         "changed_fields":changed, "before_status":ar.get("status"), "after_status":br.get("status"),
                         "before":{f:af.get(f, "") for f in ("name", "title", "email")},
                         "after":{f:bf.get(f, "") for f in ("name", "title", "email")},
                         "before_bio_chars":len(af.get("biography", "")), "after_bio_chars":len(bf.get("biography", "")),
                         "before_anomalies":ar.get("anomalies", []), "after_anomalies":br.get("anomalies", [])})
    return rows


def validate_export(data, archive_bytes):
    from openpyxl import load_workbook
    archive=zipfile.ZipFile(io.BytesIO(archive_bytes))
    names=[name for name in archive.namelist() if name.endswith('/教师列表.xlsx')]
    if len(names)!=1:
        raise AssertionError('Each case must export exactly one university workbook')
    manifest=json.loads(archive.read('清单.json'))
    if len(manifest)!=1 or manifest[0]['teachers']!=len(data['teachers']):
        raise AssertionError('Archive manifest differs from delivered records')
    wb = load_workbook(io.BytesIO(archive.read(names[0])), read_only=True)
    if wb.sheetnames != list(TEACHER_SHEETS):
        raise AssertionError("Unexpected export sheets")
    rows = 0
    for sheet in wb:
        values = list(sheet.values)
        if list(values[0]) != TEACHER_HEADERS:
            raise AssertionError("Unexpected export columns")
        rows += len(values)-1
    if rows != len(data["teachers"]):
        raise AssertionError("Export row count differs")
    fulltext_count=0
    for row in data['teachers']:
        for column,value in enumerate(row,1):
            if not isinstance(value,str) or len(value)<=32767:
                continue
            candidates=[item for item in manifest[0]['long_fields']
                        if item['field']==TEACHER_HEADERS[column-1]
                        and item['path'].startswith(f'全文附件/{row[0]:05d}-')]
            if len(candidates)!=1:
                raise AssertionError('Long field attachment is missing or ambiguous')
            attached=archive.read(manifest[0]['folder']+'/'+candidates[0]['path']).decode('utf-8-sig')
            if attached!=value:
                raise AssertionError('Long field attachment changed the original text')
            fulltext_count+=1
    result={"sheets":wb.sheetnames, "teacher_rows":rows, "columns":TEACHER_HEADERS,
            "verified_archive":True,"fulltext_attachments_verified":fulltext_count}
    wb.close(); archive.close()
    return result


async def run_case(case, result, client, store, batch_a, batch_b):
    folder = ROOT / "runs" / case["case"]
    baseline = Path(result["run_dir"])
    candidate = folder / "fetch-candidate"
    if (folder / "comparison.json").exists():
        return read(folder / "comparison.json")
    if not candidate.exists():
        shutil.copytree(baseline, candidate, ignore=shutil.ignore_patterns("storage", "assessments", "auto-assessment.json", "review-summary.json"))
    pages = read(candidate / "pages.json", [])
    observations = read(candidate / "observations.json", [])
    records = []
    semaphore = asyncio.Semaphore(4)
    async def fetch_page(page):
        if page.get("body_status") in ("permission_denied", "redirect_out_of_scope"):
            records.append({"url":page["url"], "fetch_status":"not_requested_access_or_scope_restriction", "applied":False})
            return
        if page.get("status", 0) >= 400:
            records.append({"url":page["url"], "fetch_status":"not_requested_origin_error", "applied":False})
            return
        names = {row["name"] for row in observations if row.get("profile_url", "").split("#")[0].replace("http:","https:",1)
                 == page["url"].split("#")[0].replace("http:","https:",1)}
        if page["kind"] != "profile":
            names |= {row["name"] for row in observations if row["source_url"] == page["url"]}
        async with semaphore:
            fetched = await client.fetch(page["url"])
        record = apply_candidate(candidate, page, fetched, names)
        records.append(record)
        save(folder / "fetch-pages.json", records)
        save(ROOT / "fetch-progress.json", {"case":case["case"], "processed_in_case":len(records), "total_in_case":len(pages), "calls":len(client.calls)})
        if len(records) % 10 == 0:
            print(json.dumps({"case":case["case"], "fetch_pages":len(records), "total":len(pages)}, ensure_ascii=False), flush=True)
    await asyncio.gather(*(fetch_page(page) for page in pages))
    order = {page['url']:i for i,page in enumerate(pages)}
    records.sort(key=lambda r:order[r['url']])
    save(folder / "fetch-pages.json",records)
    save(candidate / "pages.json", pages)
    summary = read(candidate / "summary.json", {})
    summary.update(mode="paired_fetch_candidate", baseline_run=str(baseline),
                   fetch_pages_applied=sum(r.get("applied",False) for r in records),
                   whole_college_completeness="unverified", fetch_only_crawler=False)
    save(candidate / "summary.json", summary)
    catalogs = []
    for arm, batch, run_dir in (("baseline",batch_a,baseline), ("fetch",batch_b,candidate)):
        with store.db() as con:
            candidates = [dict(r) for r in con.execute("SELECT id,target FROM jobs WHERE batch_id=?",(batch,))]
        key = next(row["id"] for row in candidates if json.loads(row["target"])["college"] == case["college"])
        store.update(key, state="finished", started=result["started"], finished=result["finished"], run_dir=str(run_dir))
        build(store,key,force=True)
        cat = store.catalog(key)
        catalogs.append(cat)
        save(folder / f"{arm}-catalog.json",cat)
        data = tables([cat])
        archive=university_zip([cat])
        save(folder / f"{arm}-export-validation.json",validate_export(data,archive))
        (folder / f"{arm}-试验结果.zip").write_bytes(archive)
    left, right = catalogs
    changes = compare_catalogs(left,right)
    save(folder / "person-differences.json", changes)
    changed_source_files = [str(path.relative_to(baseline)) for path in baseline.rglob("source.*")
                            if (candidate / path.relative_to(baseline)).is_file()
                            and path.read_bytes() != (candidate / path.relative_to(baseline)).read_bytes()]
    result = {"case":case["case"], "university":case["university"], "college":case["college"],
              "baseline":catalog_summary(left), "fetch":catalog_summary(right), "changed_records":len(changes),
              "field_or_status_changes":sum(bool(row["changed_fields"]) or row["before_status"]!=row["after_status"] for row in changes),
              "fetch_statuses":dict(Counter(r["fetch_status"] for r in records)),
              "fetch_applied":sum(r.get("applied",False) for r in records),
              "raw_sources_modified":changed_source_files, "whole_college_completeness":"unverified"}
    save(folder / "comparison.json",result)
    print(json.dumps(result,ensure_ascii=False),flush=True)
    return result


async def main(selected=None):
    cases = read(ROOT / "cases.json")
    if selected:
        cases = [case for case in cases if case["case"] in selected]
    store = Store(ROOT / "experiment.db")
    targets = [{k:v for k,v in c.items() if k in ("university","college","url","entries")} for c in cases]
    batches = read(ROOT / "batches.json")
    if not batches:
        batches = {"baseline":store.create_batch("DS Fetch 对照：现有方案",targets),
                   "fetch":store.create_batch("DS Fetch 对照：实验候选",targets)}
        save(ROOT / "batches.json",batches)
    client = FetchClient(ROOT / "fetch-api",max_calls=1500,timeout=60)
    results = []
    try:
        # Small attachment case first. Long college capture continues independently.
        for case in sorted(cases,key=lambda c:{"ntu":0,"fynu":1,"seu":2}[c["case"]]):
            folder = ROOT / "runs" / case["case"]
            started = time.monotonic()
            while not (folder / "capture-result.json").exists():
                if time.monotonic()-started > 2500:
                    raise RuntimeError("Capture not finished within experiment wait budget")
                await asyncio.sleep(5)
            result = read(folder / "capture-result.json")
            if not (Path(result.get("run_dir","")) / "summary.json").is_file():
                results.append({"case":case["case"], "state":"no_completed_capture", "capture":result})
                continue
            results.append(await run_case(case,result,client,store,batches["baseline"],batches["fetch"]))
            save(ROOT / "comparison.json",results)
    finally:
        await client.close()
        save(ROOT / "fetch-usage.json",{"attempts":len(client.calls), "successful":sum(r.get("billable_success",False) for r in client.calls),
             "estimated_usd":round(sum(r.get("estimated_fetch_usd",0) for r in client.calls),6), "account_bill_verified":False})


if __name__ == "__main__":
    parser=argparse.ArgumentParser()
    parser.add_argument("--case",action="append")
    args=parser.parse_args()
    asyncio.run(main(args.case))
