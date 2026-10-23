import io,unittest,unicodedata
from openpyxl import load_workbook
from workbench_export import tables,workbook_bytes,homepage_checks,teacher_sheet,TEACHER_SHEETS
from identity_export import consolidate
import test_bulk_identity as identities
from cross_college_identity import appointment_bridge
from api_pilot import title_name
from quality_fields import fields_from_text,profile_label_name

class LiveDiscoveryRegressionTests(unittest.TestCase):
    def test_biography_heading_prefix_is_person_not_document(self):
        for label,name in [('导师简介-王琦','王琦'),('导师简介—郭海鹏','郭海鹏'),
                           ('教师简介：何修传','何修传'),('个人简介 - 张三（教授）','张三')]:
            self.assertEqual(title_name(label),name)
        for label in ['创新设计与产业发展研究中心导师组介绍','导师简介-2026年通知','导师简介-张三与李四的研究项目']:
            self.assertIsNone(title_name(label))
        # The roster identity check uses the same normalization as acquisition.
        self.assertEqual(profile_label_name('导师简介-何修传'),'何修传')

    def test_explicit_rank_survives_inline_layout_newlines(self):
        for value,expected in [('副\n教授','副教授'),('副主任医师，副\n教授','副主任医师、副教授'),
                               ('副\n研究员','副研究员')]:
            selected=fields_from_text('刘苏','姓名：刘苏\n职称：'+value+'\n研究方向：神经科学')['title']
            self.assertEqual(selected['value'],expected)
            self.assertEqual(selected['rule'],'explicit_title')
            self.assertIn(unicodedata.normalize('NFKC',value),selected['evidence'])
        self.assertEqual(fields_from_text('张三','职称：讲师\n职务：教授委员会委员')['title']['value'],'讲师')
        self.assertEqual(fields_from_text('陈晨','职称：助教\n学历简介：本科')['title']['value'],'助教')

class DeliverySheetsTests(unittest.TestCase):
    def test_three_sheets_partition_and_access_failure_priority(self):
        records=[]
        for name,email,checks in [('有邮箱','a@test.edu.cn',[{'url':'https://t/a','state':'available'}]),
                                  ('无邮箱','',[]),('失效有邮箱','b@test.edu.cn',[{'url':'https://t/b','state':'failed'}]),
                                  ('尚未读','',[{'url':'https://t/c','state':'unknown'}]),
                                  ('另有正常页','c@test.edu.cn',[{'url':'https://t/d','state':'failed'},{'url':'https://t/e','state':'available'}])]:
            member={'homepage_checks':checks,'delivery_notes':[]}
            records.append({'name':name,'email':email,'members':[member],'sources':[]})
        rows=[[i+1,p['name'],'',p['email'],'大学','学院','自动采集','2026-10-14','全文',''] for i,p in enumerate(records)]
        wb=load_workbook(io.BytesIO(workbook_bytes({'teachers':rows,'records':records,'statuses':[]})))
        self.assertEqual(wb.sheetnames,list(TEACHER_SHEETS))
        names=[[r[1] for r in s.iter_rows(min_row=2,values_only=True)] for s in wb]
        self.assertEqual(names,[['有邮箱','另有正常页'],['无邮箱','尚未读'],['失效有邮箱']])
        self.assertEqual(sum(s.max_row-1 for s in wb),len(records))
        self.assertIn('失效链接',wb[TEACHER_SHEETS[2]]['J2'].comment.text)

    def test_access_status_from_actual_pages_not_blank_fields_or_attachment(self):
        person={'review':{'fields':{'profile_url':'https://t/a'}}}
        cat={'pages':[],'issues':[{'id':'x','url':'https://t/a','issue_kind':'profile_pdf_failed'}]}
        self.assertEqual(homepage_checks(cat,person)[0]['state'],'unknown')
        cat['issues'][0]['issue_kind']='request_failed'
        self.assertEqual(homepage_checks(cat,person)[0]['state'],'failed')
        cat['pages']=[{'url':'https://t/a','status':200}]
        self.assertEqual(homepage_checks(cat,person)[0]['state'],'available')
        cat['pages'][0]['body_status']='permission_denied'
        self.assertEqual(homepage_checks(cat,person)[0]['state'],'failed')

class IdentityBridgeTests(unittest.TestCase):
    record=identities.IdentityTests.record
    def test_official_appointment_plus_specific_shared_fact_links_colleges(self):
        a=self.record(email='',profile_url='https://test.edu/a',biography='任职经历：江苏省力学学会理事。研究智能材料。')
        b=self.record(college='乙学院',email='',profile_url='https://test.edu/b',biography='张三，教授，曾任测试大学甲学院副院长，江苏省力学学会理事。主讲课程：材料力学。')
        merged=consolidate([a,b]);self.assertEqual(len(merged),1)
        self.assertEqual(merged[0]['merge_evidence'][0]['rule'],'named_cross_college_appointment_with_corroboration')
        b['biography']='张三，教授，毕业于测试大学甲学院。江苏省力学学会理事。'
        self.assertIsNone(appointment_bridge(a,b))
        b['biography']='张三，教授，导师为测试大学甲学院副院长。江苏省力学学会理事。'
        self.assertIsNone(appointment_bridge(a,b))
        b['biography']='张三，教授，曾任测试大学甲学院副院长。'
        self.assertIsNone(appointment_bridge(a,b))

if __name__=='__main__':unittest.main()
