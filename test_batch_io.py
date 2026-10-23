import asyncio
import csv
import io
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch, AsyncMock

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from batch_input import preview_input
from workbench import make_app
from workbench_store import Store
from workbench_runner import run_job
from workbench_export import collected_date, source_links, TEACHER_HEADERS, SOURCE_HEADERS
from test_workbench import fixture, accept


class InputTests(unittest.TestCase):
    def test_provided_entry_equal_to_home_is_read_and_failure_does_not_skip_discovery(self):
        from api_pilot import Pilot
        from bs4 import BeautifulSoup
        from contextlib import redirect_stdout
        args=SimpleNamespace(url='https://test.edu.cn/',university='大学',college='学院',max_pages=10,entries=['https://test.edu.cn/'])
        with tempfile.TemporaryDirectory() as temp:
            pilot=Pilot(args,Path(temp),SimpleNamespace(calls=[]))
            pilot.persist=lambda:None
            pilot.roster=AsyncMock(side_effect=ValueError('isolated failure'));pilot.discover=AsyncMock()
            page={'url':args.url,'loaded_url':args.url,'title':'师资队伍','directory':'home'}
            context=SimpleNamespace(soup=BeautifulSoup('<h1>教师</h1>','lxml'),add_requests=AsyncMock())
            with patch('api_pilot.save_page',new=AsyncMock(return_value=page)),redirect_stdout(io.StringIO()):
                asyncio.run(pilot.consume(context))
            pilot.roster.assert_awaited_once();pilot.discover.assert_awaited_once()
            self.assertEqual(pilot.issues[0]['stage'],'provided_roster')

    def test_home_redirect_normalizes_subsequent_profile_links(self):
        from api_pilot import Pilot
        from bs4 import BeautifulSoup
        with tempfile.TemporaryDirectory() as temp:
            args=SimpleNamespace(url='http://test.edu.cn/',entries=[],max_pages=50)
            pilot=Pilot(args,Path(temp),SimpleNamespace(calls=[]))
            pilot.persist=lambda:None
            pilot.discover=AsyncMock()
            context=SimpleNamespace(soup=BeautifulSoup('<html></html>','html.parser'),add_requests=AsyncMock())
            saved={'url':args.url,'loaded_url':'https://test.edu.cn/','directory':'home','title':'学院首页'}
            with patch('api_pilot.save_page',new=AsyncMock(return_value=saved)):
                asyncio.run(pilot.consume(context))
            pilot.register('http://test.edu.cn/person','profile','师资')
            pilot.register('https://test.edu.cn/person','profile','师资')
            self.assertEqual(len(pilot.targets),2)
            self.assertEqual(args.url,'http://test.edu.cn/')

    def test_excel_three_and_four_columns_with_optional_header(self):
        result=preview_input('大学\t学院名\t学院链接\t师资列表链接\n测试大学\t甲学院\thttps://a.edu.cn/\t\n测试大学\t乙学院\thttps://b.edu.cn/\thttps://b.edu.cn/list.htm')
        self.assertEqual(result['errors'],[])
        self.assertEqual(result['count'],2)
        self.assertEqual(result['without_entries'],1)
        self.assertEqual(result['targets'][1]['entries'],['https://b.edu.cn/list.htm'])

    def test_markdown_links_and_separator(self):
        result=preview_input('| 大学 | 学院名 | 学院链接 |\n| -- | -- | -- |\n| 测试大学 | 甲学院 | [官网](https://a.edu.cn/) |')
        self.assertEqual(result['errors'],[])
        self.assertEqual(result['targets'][0]['url'],'https://a.edu.cn/')

    def test_spaces_csv_and_quoted_college_name(self):
        self.assertEqual(preview_input('测试大学 甲学院 https://a.edu.cn/')['count'],1)
        result=preview_input('Test University,"Arts, Sciences",https://a.edu.cn/,https://a.edu.cn/list;https://a.edu.cn/team')
        self.assertEqual(result['errors'],[])
        self.assertEqual(result['targets'][0]['college'],'Arts, Sciences')
        self.assertEqual(len(result['targets'][0]['entries']),2)

    def test_duplicate_college_merges_entries(self):
        result=preview_input('测试大学 甲学院 https://a.edu.cn/\n测试大学 甲学院 https://a.edu.cn/ https://a.edu.cn/list')
        self.assertEqual(result['duplicates_merged'],1)
        self.assertEqual(result['targets'][0]['entries'],['https://a.edu.cn/list'])

    def test_bad_fourth_column_or_extra_column_is_not_silently_discarded(self):
        result=preview_input('测试大学 甲学院 https://a.edu.cn/ https://b.edu.cn/list\n测试大学 乙学院 https://b.edu.cn/ https://b.edu.cn/list extra')
        self.assertEqual([e['row'] for e in result['errors']],[2,1])
        self.assertEqual(result['count'],0)
        self.assertTrue(preview_input(targets=[{'university':'大学','college':'学院','url':'https://a.edu.cn/','entries':'https://a.edu.cn/'}])['errors'])

    def test_batch_preview_has_no_side_effect_and_invalid_submission_is_atomic(self):
        with tempfile.TemporaryDirectory() as temp:
            store=Store(Path(temp)/'db');client=TestClient(make_app(store))
            text='测试大学 甲学院 https://a.edu.cn/\n测试大学 乙学院 invalid'
            self.assertEqual(client.post('/api/batches/preview',json={'text':text}).status_code,200)
            self.assertEqual(client.post('/api/batches',json={'name':'验证','text':text}).status_code,422)
            self.assertEqual(store.listing()['jobs'],[])

    def test_create_is_idempotent_and_keeps_fourth_column(self):
        with tempfile.TemporaryDirectory() as temp:
            store=Store(Path(temp)/'db');client=TestClient(make_app(store))
            payload={'name':'验证','text':'测试大学 甲学院 https://a.edu.cn/ https://a.edu.cn/list','request_id':'isolated-submit-1'}
            first=client.post('/api/batches',json=payload)
            self.assertEqual(first.status_code,200)
            self.assertEqual(first.json(),client.post('/api/batches',json=payload).json())
            self.assertEqual(len(store.listing()['jobs']),1)
            self.assertEqual(store.listing()['jobs'][0]['target']['entries'],['https://a.edu.cn/list'])
            self.assertEqual(client.post('/api/batches',json={**payload,'name':'changed'}).status_code,409)

    def test_worker_receives_user_entry_and_applies_assessment(self):
        with tempfile.TemporaryDirectory() as temp:
            store,key=fixture(Path(temp));job=store.job(key)
            job['target']['entries']=['https://test.edu.cn/list']
            calls=[]
            async def fake_process(*args,**kwargs):
                calls.append(args)
                pointer=Path(args[args.index('--pointer')+1])
                pointer.write_text(json.dumps({'run_dir':job['run_dir']}))
                (Path(job['run_dir'])/'summary.json').write_text('{"done":true}')
                return SimpleNamespace(returncode=0)
            with patch('workbench_runner.asyncio.create_subprocess_exec',side_effect=fake_process):
                asyncio.run(run_job(store,job))
            self.assertEqual(calls[0][calls[0].index('--entries')+1],'https://test.edu.cn/list')
            self.assertEqual(store.job(key)['state'],'finished')
            self.assertTrue(store.catalog(key)['automation'])


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.store,self.key=fixture(Path(self.temp.name));self.client=TestClient(make_app(self.store));self.addCleanup(self.client.close)
        self.person=self.store.catalog(self.key)['people'][0]['id']
        self.store.update(self.key,started='2026-10-13T18:00:00+00:00')

    def test_chinese_csv_schema_multiline_full_text_and_fixed_collector(self):
        data=accept();data['fields']['biography']='第一段，含逗号。\n第二段含"引号"。'
        self.store.decide(self.key,self.person,data)
        r=self.client.get(f'/api/jobs/{self.key}/export?format=csv')
        rows=list(csv.reader(io.StringIO(r.content.decode('utf-8-sig'))))
        self.assertEqual(rows[0],TEACHER_HEADERS)
        self.assertEqual(rows[1][6:9],['自动采集','2026-10-14',data['fields']['biography']])
        self.assertEqual(rows[1][9],data['fields']['profile_url'])

    def test_xlsx_tables_and_source_text_never_formulas(self):
        data=accept();data['fields']['name']='=HYPERLINK("https://bad")'
        self.store.decide(self.key,self.person,data)
        r=self.client.get(f'/api/jobs/{self.key}/export?format=xlsx')
        self.assertEqual(r.status_code,200)
        wb=load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames,['有邮箱教师','无邮箱教师','主页失效教师'])
        self.assertEqual([c.value for c in wb['无邮箱教师'][1]],TEACHER_HEADERS)
        self.assertEqual(wb['无邮箱教师']['B2'].data_type,'s')
        self.assertEqual(wb['无邮箱教师']['B2'].value,data['fields']['name'])
        self.assertIn('https://test.edu.cn/list',wb['无邮箱教师']['J2'].comment.text)
        self.assertEqual(wb['无邮箱教师'].freeze_panes,'A2')

    def test_only_latest_run_is_exported_and_sequence_is_continuous(self):
        self.store.decide(self.key,self.person,accept());batch=self.store.job(self.key)['batch_id']
        rows=list(csv.reader(io.StringIO(self.client.get(f'/api/batches/{batch}/export?format=csv').content.decode('utf-8-sig'))))
        self.assertEqual(len(rows),2);self.assertEqual(rows[1][0],'1')
        self.store.rerun(self.key,[],'测试','隔离新运行')
        rows=list(csv.reader(io.StringIO(self.client.get(f'/api/batches/{batch}/export?format=csv').content.decode('utf-8-sig'))))
        self.assertEqual(rows,[TEACHER_HEADERS])
        self.assertEqual(self.client.get('/api/batches/missing/export').status_code,404)

    def test_unknown_entry_keeps_empty_mapping_and_explicit_note(self):
        cat=self.store.catalog(self.key);cat['job']['target']['entries']=['https://test.edu.cn/guess']
        cat['discovery']={'entry_link_ids':[1],'url':'https://test.edu.cn/guess'}
        self.assertEqual(source_links(cat),[])
        r=self.client.get(f'/api/jobs/{self.key}/export?format=xlsx');wb=load_workbook(io.BytesIO(r.content))
        self.assertIn('尚未核实',wb['无邮箱教师']['A1'].comment.text)
        self.assertEqual(wb['无邮箱教师'].max_row,1)

    def test_multiple_qualified_branches_without_pagination_or_rejected_sources(self):
        cat=self.store.catalog(self.key)
        cat['rosters']=[{'url':'https://test.edu.cn/list','next':['https://test.edu.cn/list2'],'source_qualification':'faculty_roster'},
                        {'url':'https://test.edu.cn/list2','source_qualification':'faculty_roster'},
                        {'url':'https://test.edu.cn/team','source_qualification':'faculty_roster'},
                        {'url':'https://test.edu.cn/news','source_qualification':'faculty_roster'}]
        cat['issues']=[{'issue_kind':'source_not_faculty_roster','url':'https://test.edu.cn/news'}]
        self.assertEqual(source_links(cat),['https://test.edu.cn/list','https://test.edu.cn/team'])

    def test_long_biography_is_not_silently_truncated_by_excel(self):
        data=accept();data['fields']['biography']='全文'*17000
        self.store.decide(self.key,self.person,data)
        self.assertEqual(self.client.get(f'/api/jobs/{self.key}/export?format=xlsx').status_code,422)
        rows=list(csv.reader(io.StringIO(self.client.get(f'/api/jobs/{self.key}/export?format=csv').content.decode('utf-8-sig'))))
        self.assertEqual(rows[1][8],data['fields']['biography'])
        self.assertEqual(collected_date({'started':None}),'')


if __name__=='__main__':unittest.main()
